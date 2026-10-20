import openpyxl
import json
import re

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
with open('data/id_mappings.json', 'r', encoding='utf-8') as f:
    id_mappings = json.load(f)

# Let's inspect each sheet
for sname in wb.sheetnames:
    if sname in ['📊 خلاصه و داشبورد', '📖 راهنمای کاربری']:
        continue
    ws = wb[sname]
    print(f"\n==================== SHEET: {sname} ====================")
    current_brand = None
    current_section = None
    
    for r in range(1, ws.max_row + 1):
        vA = ws.cell(row=r, column=1).value
        vB = ws.cell(row=r, column=2).value
        vR = ws.cell(row=r, column=18).value
        
        if not vA:
            continue
        vA_str = str(vA).strip()
        
        if '📂' in vA_str:
            current_section = vA_str
            current_brand = vA_str.split('◀')[-1].strip() if '◀' in vA_str else vA_str
            continue
            
        if vA_str.startswith('↳'):
            continue
            
        if vA_str == 'نام محصول':
            continue
            
        # Analyze product
        issues = []
        
        # 1. Is it under 'متفرقه'?
        if current_brand == 'متفرقه':
            issues.append(f"Placed under 'متفرقه' (Actual brand might be inside title or spec)")
            
        # 2. Does title contain the section's brand?
        # Normalize brand names for comparison
        norm_brand = current_brand.replace('کلوِر', 'کلور').replace('هاورسون', 'هاورسن') if current_brand else ''
        norm_title = vA_str.replace('کلوِر', 'کلور').replace('هاورسون', 'هاورسن')
        
        if current_brand and current_brand != 'متفرقه':
            # Check if brand is in title
            if norm_brand.lower() not in norm_title.lower():
                issues.append(f"Brand '{current_brand}' NOT in title '{vA_str}'")
                
        # 3. Does title contain a DIFFERENT brand than current_brand?
        # Check against common brands
        all_brands = [
            'کلور', 'ایستکول', 'دیپوینت', 'بست', 'اسنوا', 'دوو', 'هیمالیا', 'الکترواستیل',
            'پاکشوما', 'ریتون', 'آراز', 'رانکو', 'ساپر', 'سران', 'سیماران', 'فریدولین', 'فیلیپس',
            'مونزا', 'هیتاچی', 'پاناسونیک', 'پاندا', 'پاکیزان', 'وگاترونیکس', 'دونالکس', 'گرنویک',
            'وستینگهاوس', 'پارس خزر', 'آلتون', 'لاوان', 'نیاک', 'ویلز', 'کنوود', 'تفال', 'کورکماز',
            'زامروتی', 'وگاتی', 'اریکسون', 'جنرال', 'کاستلو', 'وولگا', 'ژمبو', 'پروکس', 'اوریکس',
            'ونتو', 'سانی', 'هاورسن', 'پرومارون', 'پروامگا', 'امگا', 'الکتروسان',
            'تکنوسان', 'نیکسان', 'ریگور', 'هومگلد', 'هانو', 'وست واتر', 'نیکایی', 'افترون',
            'سی سی کا', 'سافت واتر', 'دلمونتی', 'بوتان', 'پلار', 'الکترا', 'اوجنرال', 'گلدن لوکس',
            'اویژ', 'اویز', 'اویزپلاس', 'آویز', 'آویزپلاس', 'دستی لوکس', 'دستی', 'کنرا', 'KENRA'
        ]
        
        detected_brands_in_title = [b for b in all_brands if b.lower() in norm_title.lower()]
        if current_brand and current_brand != 'متفرقه':
            for db in detected_brands_in_title:
                db_norm = db.replace('کلوِر', 'کلور').replace('هاورسون', 'هاورسن').replace('آویز', 'اویز')
                cb_norm = norm_brand.replace('آویز', 'اویز')
                if db_norm != cb_norm and db_norm not in cb_norm and cb_norm not in db_norm:
                    issues.append(f"Title has brand '{db}' but section is '{current_brand}'")
                    
        if issues:
            print(f"Row {r:3d} [Section: {current_brand}]: {vA_str}")
            for iss in issues:
                print(f"   --> {iss}")
            print(f"   Status: {vR}, Price: {vB}")
