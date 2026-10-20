import openpyxl
import json
import os
import glob
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
product_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

sheet_folder_map = {
    'یخچال و فریزر': 'refrigerators',
    'شستشو و نظافت': 'washing_and_cleaning',
    'لوازم برقی و پخت و پز': 'cooking_appliances',
    'نوشیدنی ساز': 'beverage_makers',
    'سرمایش و گرمایش': 'cooling_heating',
    'تصفیه کننده آب': 'water_purifiers',
    'صوتی و تصویری': 'audio_video',
    'زیبایی و سلامت فردی': 'beauty_health'
}

all_rows = []

for sname in product_sheets:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    current_parent = None
    current_parent_row = None
    
    for r in range(3, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        ptype = ws.cell(row=r, column=12).value
        brand = ws.cell(row=r, column=11).value
        status = ws.cell(row=r, column=18).value
        
        if not name:
            continue
        name_str = str(name).strip()
        if name_str.startswith('📂') or 'کل اقلام' in name_str:
            continue
            
        status_str = str(status).strip() if status else 'None'
        is_var = name_str.startswith('↳')
        
        if not is_var:
            current_parent = name_str
            current_parent_row = r
            parent_name = None
            parent_row = None
        else:
            parent_name = current_parent
            parent_row = current_parent_row
            
        all_rows.append({
            'sheet': sname,
            'folder': sheet_folder_map.get(sname, 'other'),
            'row': r,
            'name': name_str,
            'is_variation': is_var,
            'parent_name': parent_name,
            'parent_row': parent_row,
            'brand': str(brand).strip() if brand else '',
            'ptype': str(ptype).strip() if ptype else '',
            'status': status_str
        })

print(f"Total product rows extracted across all sheets: {len(all_rows)}")

from collections import Counter
status_counts = Counter([r['status'] for r in all_rows])
print("\nStatus distribution across product rows:")
for st, c in status_counts.items():
    print(f"  - {st}: {c}")

with open('scratch/catalog_current_rows.json', 'w', encoding='utf-8') as f:
    json.dump(all_rows, f, ensure_ascii=False, indent=2)
