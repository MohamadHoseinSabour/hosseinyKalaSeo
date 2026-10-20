import openpyxl
import json
import shutil
import sys

sys.stdout.reconfigure(encoding='utf-8')

with open('scratch/user_changes_audit.json', encoding='utf-8') as f:
    audit = json.load(f)

hold_candidates = audit['hold_candidates']

# Load workbook with data_only=False to preserve all formulas
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=False)

updated_count = 0
for h in hold_candidates:
    sheet_name = h['sheet']
    prow = h['row']
    ws = wb[sheet_name]
    
    # Update parent row
    old_val = ws.cell(row=prow, column=18).value
    ws.cell(row=prow, column=18).value = 'فعلا دست نگه دار'
    updated_count += 1
    print(f"Updated [{sheet_name} | Row {prow}] from '{old_val}' -> 'فعلا دست نگه دار'")
    
    # Update variations if any
    r = prow + 1
    while r <= ws.max_row:
        c1 = ws.cell(row=r, column=1).value
        if c1 and str(c1).strip().startswith('↳'):
            old_v_val = ws.cell(row=r, column=18).value
            ws.cell(row=r, column=18).value = 'فعلا دست نگه دار'
            updated_count += 1
            print(f"   -> Var Row {r}: from '{old_v_val}' -> 'فعلا دست نگه دار'")
            r += 1
        else:
            break

# Save workbook
wb.save('hosseiny_kala_catalog.xlsx')
print(f"\nSuccessfully updated {updated_count} rows in hosseiny_kala_catalog.xlsx!")
