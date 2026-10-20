import sys
import os
import json
import shutil
import datetime
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

print("=== شروع به‌روزرسانی مدل آسیاب قهوه مایر از 4155 به 4099 ===")

# ۱. پشتیبان‌گیری و اصلاح hosseiny_kala_catalog.xlsx
catalog_path = 'hosseiny_kala_catalog.xlsx'
ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
backup_path = f'hosseiny_kala_catalog.backup_model4099_{ts}.xlsx'
shutil.copy2(catalog_path, backup_path)
print(f"۱. پشتیبان از کاتالوگ اکسل گرفته شد: {backup_path}")

wb = openpyxl.load_workbook(catalog_path, data_only=False)
ws = wb['نوشیدنی ساز']

# ردیف 79 مربوط به شناسه 4249
row_target = 79
assert ws.cell(row_target, 14).value == 4249, f"شناسه ردیف {row_target} برابر ۴۲۴۹ نیست!"

print(f"   نام قبلی در اکسل: {ws.cell(row_target, 1).value}")
print(f"   لینک قبلی در اکسل: {ws.cell(row_target, 16).value}")

ws.cell(row_target, 1).value = 'آسیاب قهوه مایر مدل MR-4099'
ws.cell(row_target, 16).value = 'https://hosseinykala.ir/product/maier-mr4099-digital-conical-burr-coffee-grinder/'

wb.save(catalog_path)
print(f"   ✓ نام و لینک در ردیف {row_target} شیت نوشیدنی ساز با موفقیت به مدل MR-4099 اصلاح شد.")

# ۲. اصلاح data/products.json
products_path = 'data/products.json'
with open(products_path, 'r', encoding='utf-8') as f:
    products = json.load(f)

updated_products = 0
for p in products:
    if p.get('id') == 4249:
        p['name'] = 'آسیاب قهوه مایر مدل MR-4099'
        p['slug'] = 'maier-mr4099-digital-conical-burr-coffee-grinder'
        p['url'] = 'https://hosseinykala.ir/product/maier-mr4099-digital-conical-burr-coffee-grinder/'
        if 'description' in p:
            p['description'] = p['description'].replace('MR-4155', 'MR-4099').replace('۴۱۵۵', '۴۰۹۹')
        if 'short_description' in p:
            p['short_description'] = p['short_description'].replace('MR-4155', 'MR-4099').replace('۴۱۵۵', '۴۰۹۹')
        updated_products += 1

with open(products_path, 'w', encoding='utf-8') as f:
    json.dump(products, f, ensure_ascii=False, indent=2)
print(f"۲. فایل {products_path} با موفقیت به‌روزرسانی شد ({updated_products} محصول).")

# ۳. اصلاح data/id_mappings.json
mappings_path = 'data/id_mappings.json'
with open(mappings_path, 'r', encoding='utf-8') as f:
    id_map = json.load(f)

if '4249' in id_map:
    id_map['4249']['name'] = 'آسیاب قهوه مایر مدل MR-4099'
    id_map['4249']['slug'] = 'maier-mr4099-digital-conical-burr-coffee-grinder'
    print("   ✓ کلید 4249 در id_mappings به‌روزرسانی شد.")

if 'products' in id_map:
    for p in id_map['products']:
        if p.get('id') == 4249:
            p['name'] = 'آسیاب قهوه مایر مدل MR-4099'
            p['slug'] = 'maier-mr4099-digital-conical-burr-coffee-grinder'
            print("   ✓ آیتم 4249 در لیست products فایل id_mappings به‌روزرسانی شد.")

with open(mappings_path, 'w', encoding='utf-8') as f:
    json.dump(id_map, f, ensure_ascii=False, indent=2)
print(f"۳. فایل {mappings_path} با موفقیت ذخیره شد.")

# ۴. اصلاح فایل‌ها و پوشه‌های divar_organized
old_folder = 'divar_organized/P-25_maier-mr4155-digital-conical-burr-coffee-grinder'
new_folder = 'divar_organized/P-25_maier-mr4099-digital-conical-burr-coffee-grinder'

if os.path.exists(old_folder):
    if os.path.exists(new_folder):
        shutil.rmtree(new_folder)
    os.rename(old_folder, new_folder)
    print(f"۴. پوشه دیوار تغییر نام یافت به: {new_folder}")

# اصلاح description.txt در پوشه دیوار
desc_file = os.path.join(new_folder, 'description.txt')
if os.path.exists(desc_file):
    with open(desc_file, 'r', encoding='utf-8') as f:
        desc_text = f.read()
    desc_text = desc_text.replace('۴۱۵۵', '۴۰۹۹').replace('MR-4155', 'MR-4099').replace('4155', '4099')
    with open(desc_file, 'w', encoding='utf-8') as f:
        f.write(desc_text)
    print(f"   ✓ فایل {desc_file} با مشخصات مدل ۴۰۹۹ اصلاح شد.")

# اصلاح manifest.json
manifest_path = 'divar_organized/manifest.json'
if os.path.exists(manifest_path):
    with open(manifest_path, 'r', encoding='utf-8') as f:
        manifest = json.load(f)
    if 'P-25_maier-mr4155-digital-conical-burr-coffee-grinder' in manifest:
        manifest['P-25_maier-mr4099-digital-conical-burr-coffee-grinder'] = manifest.pop('P-25_maier-mr4155-digital-conical-burr-coffee-grinder')
        with open(manifest_path, 'w', encoding='utf-8') as f:
            json.dump(manifest, f, ensure_ascii=False, indent=2)
        print("   ✓ کلید پوشه در manifest.json دیوار اصلاح شد.")

# اصلاح divar_new_products_list.md
divar_md = 'divar_organized/divar_new_products_list.md'
if os.path.exists(divar_md):
    with open(divar_md, 'r', encoding='utf-8') as f:
        md_text = f.read()
    md_text = md_text.replace('۴۱۵۵', '۴۰۹۹').replace('MR-4155', 'MR-4099').replace('4155', '4099')
    md_text = md_text.replace('P-25_maier-mr4155-digital-conical-burr-coffee-grinder', 'P-25_maier-mr4099-digital-conical-burr-coffee-grinder')
    with open(divar_md, 'w', encoding='utf-8') as f:
        f.write(md_text)
    print(f"   ✓ فایل {divar_md} اصلاح شد.")

# ۵. اصلاح فایل‌های کمکی در scratch
for sfile in ['scratch/new_30_products_details.json', 'scratch/wc_live_dump.json']:
    if os.path.exists(sfile):
        try:
            with open(sfile, 'r', encoding='utf-8') as f:
                content = f.read()
            content = content.replace('MR-4155', 'MR-4099').replace('۴۱۵۵', '۴۰۹۹').replace('maier-mr4155', 'maier-mr4099')
            with open(sfile, 'w', encoding='utf-8') as f:
                f.write(content)
            print(f"   ✓ کش محلی {sfile} اصلاح شد.")
        except Exception as e:
            print(f"   خطا در اصلاح {sfile}: {e}")

print("\n=== پایان موفقیت‌آمیز به‌روزرسانی فایل‌های کاتالوگ و دیتابیس محلی ===")
