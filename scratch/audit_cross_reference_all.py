import openpyxl
import json
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

with open('data/products.json', 'r', encoding='utf-8') as f:
    site_products = json.load(f)

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

# List of all newly added items to audit
new_items_to_audit = [
    # صوتی و تصویری
    ('تلویزیون ۵۵ دوو 7000', ['7000', 'دوو'], 'صوتی و تصویری'),
    
    # شستشو و نظافت
    ('لباسشویی کلور 8vw1', ['8vw1', '8vw', 'vw1'], 'شستشو و نظافت'),
    ('اتو بخار پاناسونیک 650', ['650', 'پاناسونیک'], 'شستشو و نظافت'),
    ('جاروبرقی جنرال فیت 4500', ['4500', 'جنرال فیت'], 'شستشو و نظافت'),
    ('جاروبرقی جنرال فیت 5300', ['5300', 'جنرال فیت'], 'شستشو و نظافت'),
    ('جاروبرقی جنرال فیت 5500', ['5500', 'جنرال فیت'], 'شستشو و نظافت'),
    ('جاروبرقی جنرال فیت 4800', ['4800', 'جنرال فیت'], 'شستشو و نظافت'),
    ('جاروبرقی جنرال فیت 4300', ['4300', 'جنرال فیت'], 'شستشو و نظافت'),
    ('جارو شارژی باخ 1454', ['1454', 'باخ'], 'شستشو و نظافت'),
    ('بخارگر BUVCK 1266', ['1266', 'buvck', 'بخارگر'], 'شستشو و نظافت'),
    
    # سرمایش و گرمایش
    ('بخاری پلار سنگی', ['سنگ', 'سنگی'], 'سرمایش و گرمایش'),
    ('آبگرمکن اتومات ایران شرق', ['ایران شرق', 'اتومات'], 'سرمایش و گرمایش'),
    ('آبگرمکن 3020 ایران شرق', ['3020', 'ایران شرق'], 'سرمایش و گرمایش'),
    ('بخاری آبسال 405', ['405', 'آبسال'], 'سرمایش و گرمایش'),
    
    # نوشیدنی ساز
    ('آب مرکبات گیری کنوود', ['مرکبات', 'کنوود'], 'نوشیدنی ساز'),
    ('کتری تفال 2.2 لیتر', ['2.2', 'تفال'], 'نوشیدنی ساز'),
    ('اسپرسوساز وستینگهاوس 1663', ['1663', 'وستینگ'], 'نوشیدنی ساز'),
    ('کتری وستینگهاوس قرمز', ['wkke210', '210'], 'نوشیدنی ساز'),
    ('آسیاب قهوه مایر', ['مایر', 'آسیاب'], 'نوشیدنی ساز'),
    ('آسیاب زیکو 6501', ['6501', 'زیکو'], 'نوشیدنی ساز'),
    ('آسیاب زیکو 6502', ['6502', 'زیکو'], 'نوشیدنی ساز'),
    ('آبمیوه گیری زیکو 4038', ['4038', 'زیکو'], 'نوشیدنی ساز'),
    ('اسپرسوساز مباشی', ['مباشی'], 'نوشیدنی ساز'),
    
    # لوازم برقی و پخت و پز
    ('سایز ۲۸ قابلمه اویز', ['28', '۲۸', 'اویز'], 'لوازم برقی و پخت و پز'),
    ('خردکن ویلز 6033', ['6033', 'ویلز'], 'لوازم برقی و پخت و پز'),
    ('هود آلتون 900', ['900', 'آلتون'], 'لوازم برقی و پخت و پز'),
    ('گوشت کوب وستینگهاوس قرمز', ['گوشت کوب', 'وستینگ'], 'لوازم برقی و پخت و پز'),
    ('خردکن زیکو 70028', ['70028', 'زیکو'], 'لوازم برقی و پخت و پز'),
    ('خردکن بوش', ['بوش', 'خردکن'], 'لوازم برقی و پخت و پز'),
    ('میوه خشک کن مایر 5959', ['5959', 'مایر'], 'لوازم برقی و پخت و پز'),
    
    # زیبایی و سلامت
    ('سشوار زیکو 5020', ['5020', 'زیکو'], 'زیبایی و سلامت فردی'),
    
    # یخچال و فریزر
    ('دیپوینت دوقلو اکسپلور', ['اکسپلور', 'explore'], 'یخچال و فریزر'),
    ('دیپوینت کمبی دیسنت', ['دیسنت', 'decent'], 'یخچال و فریزر'),
    ('الکترواستیل دوقلو سری 10', ['سری 10', 'سری ۱۰', '10 چرم'], 'یخچال و فریزر'),
    ('الکترواستیل 14 فوت', ['14 فوت', '۱۴ فوت'], 'یخچال و فریزر'),
    ('الکترواستیل کمپکت', ['کمپکت', 'compact'], 'یخچال و فریزر'),
    ('یخچال 27 فوت تیتانیوم', ['27 فوت تیتانیوم', '۲۷ فوت تیتانیوم'], 'یخچال و فریزر'),
    ('یخچال 35 فوت هوم بار', ['35 فوت', '۳۵ فوت'], 'یخچال و فریزر'),
]

print("=== CHECKING EACH NEW ITEM AGAINST ALL EXISTING CATALOG ENTRIES & SITE PRODUCTS ===")
for title, keywords, target_sheet in new_items_to_audit:
    matches_catalog = []
    # Search catalog
    for sname in wb.sheetnames:
        if sname.startswith('📊') or sname.startswith('📖'): continue
        ws = wb[sname]
        for r in range(2, ws.max_row + 1):
            name = str(ws.cell(r, 1).value or '')
            status = str(ws.cell(r, 18).value or '')
            brand = str(ws.cell(r, 11).value or '')
            
            # Skip if this is the new item itself
            if status == 'در سایت تعریف نشده':
                continue
                
            # Check if all keywords are in name/brand
            full_text = f"{name} {brand}".lower()
            if all(kw.lower() in full_text for kw in keywords):
                matches_catalog.append((sname, r, name, brand, status))
                
    # Search site products
    matches_site = []
    for p in site_products:
        p_name = p.get('name', '').lower()
        if all(kw.lower() in p_name for kw in keywords):
            matches_site.append((p.get('id'), p.get('name')))
            
    if matches_catalog or matches_site:
        print(f"\n⚠️ POSSIBLE EXISTING MATCH FOR: {title}")
        if matches_catalog:
            print("   In Catalog (EXISTING rows):")
            for m in matches_catalog:
                print(f"     [{m[0]}] R{m[1]}: {m[2]} ({m[3]}) - {m[4]}")
        if matches_site:
            print("   On Site (products.json):")
            for m in matches_site:
                print(f"     ID {m[0]}: {m[1]}")
    else:
        # Perfectly unique
        pass

print("\nAudit complete!")
