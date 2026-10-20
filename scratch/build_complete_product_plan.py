import os
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
category_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

catalog_rows = []
for sheet in category_sheets:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    current_parent = None
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or '📂' in str(name):
            continue
        name_str = str(name).strip()
        status = str(ws.cell(row=r, column=18).value or '').strip()
        stock = str(ws.cell(row=r, column=8).value or '').strip()
        pid = ws.cell(row=r, column=14).value
        ptype = str(ws.cell(row=r, column=12).value or 'ساده').strip()
        var_attr = str(ws.cell(row=r, column=13).value or '').strip()

        if not name_str.startswith('↳'):
            current_parent = name_str
            parent_name = None
        else:
            parent_name = current_parent

        catalog_rows.append({
            'sheet': sheet,
            'row': r,
            'name': name_str,
            'parent_name': parent_name,
            'status': status,
            'stock': stock,
            'price_rial': ws.cell(row=r, column=2).value,
            'price_toman': ws.cell(row=r, column=3).value,
            'main_cat': str(ws.cell(row=r, column=9).value or sheet).strip(),
            'subcat': str(ws.cell(row=r, column=10).value or '').strip(),
            'brand': str(ws.cell(row=r, column=11).value or '').strip(),
            'ptype': ptype,
            'variant_attr': var_attr,
            'pid': pid,
            'vid': ws.cell(row=r, column=15).value,
            'link': ws.cell(row=r, column=16).value,
            'sku': ws.cell(row=r, column=17).value
        })

unregistered = [r for r in catalog_rows if r['status'] == 'در سایت تعریف نشده']

products = []
i = 0
while i < len(unregistered):
    item = unregistered[i]
    if item['ptype'] == 'متغیر' and not item['parent_name']:
        variations = []
        j = i + 1
        while j < len(unregistered) and unregistered[j]['parent_name'] == item['name']:
            variations.append(unregistered[j])
            j += 1
        products.append({
            'type': 'variable',
            'parent': item,
            'variations': variations
        })
        i = j
    elif not item['parent_name']:
        products.append({
            'type': 'simple',
            'parent': item,
            'variations': []
        })
        i += 1
    else:
        products.append({
            'type': 'simple',
            'parent': item,
            'variations': []
        })
        i += 1

pi_dir = 'product_images_2026-10-05'

slug_mapping = {
    'یخچال و فریزر دوقلو هیمالیا مدل آیس پول 60MW': 'himalia-ice-pool-60mw-twin-refrigerator',
    'یخچال و فریزر دوقلو 27 فوت کلور مدل اکونومی': 'clover-economy-27ft-twin-refrigerator',
    'یخچال و فریزر دوقلو کلور عرض 70 مدل روسو آبریزدار': 'clover-rosso-dispenser-70cm-twin-refrigerator',
    'یخچال و فریزر دوقلو کلور عرض 70 مدل روسو هوم بار': 'clover-rosso-homebar-70cm-twin-refrigerator',
    'مینی یخچال 10 فوت کلور مدل طرح دار فانتزی': 'clover-fantasy-patterned-10ft-mini-refrigerator',
    'فریزر صندوقی تکنوسان مدل سفید بزرگ': 'technosan-large-chest-freezer-white',
    'یخچال و فریزر کمبی نیکسان مدل RF84102E2': 'niksan-rf84102e2-combi-refrigerator',
    'ماشین لباسشویی 10 کیلوگرم اینورتر هیمالیا مدل تسلا پرو': 'himalia-tesla-pro-10kg-inverter-washing-machine',
    'ماشین لباسشویی دوقلو نیمه اتوماتیک پاکشوما': 'pakshoma-twin-tub-semi-automatic-washing-machine',
    'جاروبرقی سطلی کلور': 'clover-bucket-vacuum-cleaner',
    'جاروبرقی سطلی آراز 40 لیتر': 'araz-40l-bucket-vacuum-cleaner',
    'اتوبخار رانکو': 'runko-steam-iron',
    'لباسشویی دوقلو 17 کیلویی ریتون': 'reyton-17kg-twin-tub-washing-machine',
    'اتوبخار ساپر کف سرامیک مدل 4025': 'saper-4025-ceramic-steam-iron',
    'اتوبخار ساپر کف سرامیک مدل 4035': 'saper-4035-ceramic-steam-iron',
    'جاروبرقی سطلی 21 لیتری سران مدل آب و خاک': 'soran-21l-wet-dry-vacuum-cleaner',
    'اتوبخار سیماران مدل 8002 زرد': 'simaran-8002-steam-iron-yellow',
    'مینی واش فریدولین 3.5 کیلوگرم مدل تک درب': 'feridolin-sw35-single-door-mini-washer',
    'اتوبخار مونزا مدل VR3000W': 'monza-vr3000w-steam-iron',
    'جاروبرقی سطلی صنعتی آب و خاک پاندا 80 لیتر': 'panda-80l-wet-dry-industrial-vacuum-cleaner',
    'پتوشور صنعتی و خانگی پاکیزان': 'pakizan-heavy-duty-blanket-washer',
    'تابه گرانیتی مشکی اویز (سایزهای 22 تا 32)': 'aviz-granite-frying-pan-black',
    'گاز صفحه ای استیل آلتون مدل SG518': 'alton-sg518-steel-built-in-gas-stove',
    'فر توکار آلتون مدل 7303': 'alton-7303-built-in-oven',
    'فر گاز فردار آلتون مدل MDS': 'alton-mds-freestanding-gas-cooker',
    'هود آشپزخانه آلتون مدل H312': 'alton-h312-kitchen-range-hood',
    'هود آشپزخانه آلتون مدل H315B': 'alton-h315b-kitchen-range-hood',
    'سرویس قابلمه 11 پارچه اویزپلاس مدل السا سبز': 'avizplus-elsa-11pcs-cookware-set-green',
    'قابلمه مرغ خوری اویزپلاس مشکی': 'avizplus-poultry-pot-black',
    'هواپز ساپر 7.2 لیتر مدل 1710': 'saper-saf1710-air-fryer-7-2l',
    'قابلمه تک سایز 14 لاوان': 'lavan-single-pot-size-14',
    'سبزی خردکن KENRA': 'kenra-vegetable-chopper',
    'سرویس قابلمه استیل 10 پارچه': 'stainless-steel-10pcs-cookware-set',
    'سبزی خردکن خانگی نیاک 3 کیلوگرم': 'niak-3kg-home-vegetable-chopper',
    'خردکن برقی کاسه دار ویلز': 'willez-electric-bowl-chopper',
    'سرخ کن بدون روغن ویلز مدل 5062': 'willez-5062-air-fryer',
    'چرخ گوشت برقی ویلز': 'willez-electric-meat-grinder',
    'پلوپز پارس خزر مدل 361': 'parskhazar-ts361-rice-cooker',
    'کتری برقی تفال 2 لیتر': 'tefal-2l-electric-kettle',
    'آبمیوه گیری چندکاره جنرال مدل 607 غذاساز': 'general-607-multi-function-juicer-food-processor',
    'اسپرسوساز حرفه ای رانکو': 'runko-professional-espresso-maker',
    'کتری برقی ساپر مدل 1031': 'saper-1031-electric-kettle',
    'کتری برقی ساپر استیل مدل 1011': 'saper-1011-steel-electric-kettle',
    'کتری برقی سیماران استیل مدل 750': 'simaran-750-steel-electric-kettle',
    'آسیاب برقی ونتو': 'vento-electric-grinder',
    'کتری برقی وولگا': 'volga-electric-kettle',
    'چای ساز ویلز مدل 2015': 'willez-2015-tea-maker',
    'آسیاب برقی کاستلو مدل 470': 'castello-470-electric-grinder',
    'همزن برقی کنوود مدل HM430': 'kenwood-hm430-hand-mixer',
    'چای ساز کورکماز مدل A331': 'korkmaz-a331-tea-maker',
    'آبگرمکن مخزندار دیواری الکترواستیل': 'electrosteel-wall-mounted-water-heater',
    'کولر پنجره ای اوجنرال 18 هزار': 'ogeneral-18000-window-air-conditioner',
    'بخاری پارس خزر فن دار مدل FH-2000P': 'parskhazar-fh2000p-fan-heater',
    'پنکه پایه بلند پارس خزر': 'parskhazar-standing-pedestal-fan',
    'پنکه پاناسونیک 2 حالته': 'panasonic-2-in-1-fan',
    'بخاری پلار مدل 3PN-NEW ماشینی': 'polar-3pn-new-gas-heater',
    'بخاری گلدن لوکس': 'golden-luxe-gas-heater',
    'اتو مو مسافرتی پاناسونیک مدل HV11': 'panasonic-eh-hv11-travel-hair-straightener',
    'ریش تراش چرخشی پرومارون مدل 706': 'promaron-706-rotary-shaver',
    'دستگاه پیرایش سیمی سر و صورت پرومارون مدل 731': 'promaron-731-corded-hair-clipper',
    'فر کننده مو پرومارون مدل 9909': 'promaron-9909-hair-curler',
}

CAT_CANONICAL_URLS = {
    92: ("یخچال فریزر", "https://hosseinykala.ir/product-category/%db%8c%d8%ae%da%86%d8%a7%d9%84-%d9%81%d8%b1%db%8c%d8%b2%d8%b1/"),
    153: ("یخچال فریزر کلور", "https://hosseinykala.ir/product-category/%db%8c%d8%ae%da%86%d8%a7%d9%84-%d9%81%d8%b1%db%8c%d8%b2%d8%b1/clover-refrigerators/"),
    154: ("یخچال فریزر دیپوینت", "https://hosseinykala.ir/product-category/%db%8c%d8%ae%da%86%d8%a7%d9%84-%d9%81%d8%b1%db%8c%d8%b2%d8%b1/depoint-refrigerators/"),
    155: ("یخچال فریزر ایستکول", "https://hosseinykala.ir/product-category/%db%8c%d8%ae%da%86%d8%a7%d9%84-%d9%81%d8%b1%db%8c%d8%b2%d8%b1/eastcool-refrigerators/"),
    156: ("یخچال فریزر الکترواستیل", "https://hosseinykala.ir/product-category/%db%8c%d8%ae%da%86%d8%a7%d9%84-%d9%81%d8%b1%db%8c%d8%b2%d8%b1/electrosteel-refrigerators/"),
    239: ("شستشو و نظافت", "https://hosseinykala.ir/product-category/washing-cleaning/"),
    21: ("جاروبرقی", "https://hosseinykala.ir/product-category/washing-cleaning/%d8%ac%d8%a7%d8%b1%d9%88%d8%a8%d8%b1%d9%82%db%8c/"),
    93: ("ماشین لباسشویی", "https://hosseinykala.ir/product-category/washing-cleaning/%d9%85%d8%a7%d8%b4%db%8c%d9%86-%d9%84%d8%a8%d8%a7%d8%b3%d8%b4%d9%88%db%8c%db%8c/"),
    157: ("ماشین لباسشویی کلور", "https://hosseinykala.ir/product-category/washing-cleaning/%d9%85%d8%a7%d8%b4%db%8c%d9%86-%d9%84%d8%a8%d8%a7%d8%b3%d8%b4%d9%88%db%8c%db%8c/clover-washing-machines/"),
    97: ("ماشین ظرفشویی", "https://hosseinykala.ir/product-category/washing-cleaning/%d9%85%d8%a7%d8%b4%db%8c%d9%86-%d8%b8%d8%b1%d9%81%d8%b4%d9%88%db%8c%db%8c/"),
    118: ("اتو", "https://hosseinykala.ir/product-category/washing-cleaning/%d8%a7%d8%aa%d9%88/"),
    216: ("لوازم برقی و پخت و پز", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/"),
    121: ("سرخ کن", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/%d8%b3%d8%b1%d8%ae-%da%a9%d9%86/"),
    122: ("گوشت کوب برقی", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/%da%af%d9%88%d8%b4%d8%aa-%da%a9%d9%88%d8%a8-%d8%a8%d8%b1%d9%82%db%8c/"),
    209: ("اجاق گاز", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/%d8%a7%d8%ac%d8%a7%d9%82-%da%af%d8%a7%d8%b2/"),
    210: ("گاز صفحه ای", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/%d8%a7%d8%ac%d8%a7%d9%82-%da%af%d8%a7%d8%b2/%da%af%d8%a7%d8%b2-%d8%b5%d9%81%d8%ad%d9%87-%d8%a7%db%8c/"),
    215: ("ظروف و پخت و پز", "https://hosseinykala.ir/product-category/%d9%84%d9%88%d8%a7%d8%b2%d9%85-%d8%a8%d8%b1%d9%82%db%8c-%d8%a2%d8%b4%d9%be%d8%b2%d8%ae%d8%a7%d9%86%d9%87/%d8%b8%d8%b1%d9%88%d9%81-%d9%88-%d9%be%d8%ae%d8%aa-%d9%88-%d9%be%d8%b2/"),
    240: ("نوشیدنی ساز", "https://hosseinykala.ir/product-category/drink-makers/"),
    119: ("اسپرسو ساز", "https://hosseinykala.ir/product-category/drink-makers/%d8%a7%d8%b3%d9%be%d8%b1%d8%b3%d9%88-%d8%b3%d8%a7%d8%b2/"),
    120: ("آبمیوه گیری", "https://hosseinykala.ir/product-category/drink-makers/%d8%a2%d8%a8%d9%85%db%8c%d9%88%d9%87-%da%af%db%8c%d8%b1%db%8c/"),
    123: ("مخلوط کن", "https://hosseinykala.ir/product-category/drink-makers/%d9%85%d8%ae%d9%84%d9%88%d8%b7-%da%a9%d9%86/"),
    211: ("چای ساز", "https://hosseinykala.ir/product-category/drink-makers/%da%86%d8%a7%db%8c-%d8%b3%d8%a7%d8%b2/"),
    241: ("سرمایش و گرمایش", "https://hosseinykala.ir/product-category/heating-cooling/"),
    95: ("کولر آبی", "https://hosseinykala.ir/product-category/heating-cooling/%da%a9%d9%88%d9%84%d8%b1-%d8%a2%d8%a8%db%8c/"),
    206: ("کولر گازی", "https://hosseinykala.ir/product-category/heating-cooling/%da%a9%d9%88%d9%84%d8%b1-%da%af%d8%a7%d8%b2%db%8c/"),
    214: ("پنکه و گرمایشی", "https://hosseinykala.ir/product-category/heating-cooling/%d8%b5%d9%88%d8%aa%db%8c%d8%8c-%d8%aa%d8%b5%d9%88%db%8c%d8%b1%db%8c-%d9%88-%da%af%d8%b1%d9%85%d8%a7%db%8c%d8%b4%db%8c/"),
    212: ("زیبایی و سلامت فردی", "https://hosseinykala.ir/product-category/%d8%a7%d8%b5%d9%84%d8%a7%d8%ad-%d9%88-%d8%b2%db%8c%d8%a8%d8%a7%db%8c%db%8c/"),
    124: ("سشوار", "https://hosseinykala.ir/product-category/%d8%a7%d8%b5%d9%84%d8%a7%d8%ad-%d9%88-%d8%b2%db%8c%d8%a8%d8%a7%db%8c%db%8c/%d8%b3%d8%b4%d9%88%d8%a7%d8%b1/"),
    126: ("اصلاح موی سر و صورت", "https://hosseinykala.ir/product-category/%d8%a7%d8%b5%d9%84%d8%a7%d8%ad-%d9%88-%d8%b2%db%8c%d8%a8%d8%a7%db%8c%db%8c/%d8%a7%d8%b5%d9%84%d8%a7%d8%ad-%d9%85%d9%88%db%8c-%d8%b3%d8%b1/"),
    96: ("تصفیه کننده آب", "https://hosseinykala.ir/product-category/%d8%aa%d8%b5%d9%81%db%8c%d9%87-%da%a9%d9%86%d9%86%d8%af%d9%87-%d8%a2%d8%a8/")
}

def get_product_cats(sheet, subcat, brand, name):
    cats = []
    if sheet == 'یخچال و فریزر':
        cats.append(92)
        if 'کلور' in brand or 'کلور' in name:
            cats.append(153)
        elif 'دیپوینت' in brand or 'دیپوینت' in name:
            cats.append(154)
        elif 'ایستکول' in brand or 'ایستکول' in name:
            cats.append(155)
        elif 'الکترواستیل' in brand or 'الکترواستیل' in name:
            cats.append(156)
    elif sheet == 'شستشو و نظافت':
        cats.append(239)
        if 'جاروبرقی' in name or 'جاروبرقی' in subcat:
            cats.append(21)
        elif 'لباسشویی' in name or 'پتوشور' in name or 'مینی واش' in name or 'لباسشویی' in subcat:
            cats.append(93)
            if 'کلور' in brand or 'کلور' in name:
                cats.append(157)
        elif 'ظرفشویی' in name or 'ظرفشویی' in subcat:
            cats.append(97)
        elif 'اتو' in name or 'اتو' in subcat:
            cats.append(118)
    elif sheet == 'لوازم برقی و پخت و پز':
        cats.append(216)
        if 'سرخ کن' in name or 'هواپز' in name or 'سرخ کن' in subcat:
            cats.append(121)
        elif 'گوشت کوب' in name or 'گوشت کوب' in subcat:
            cats.append(122)
        elif 'گاز صفحه ای' in name:
            cats.append(209)
            cats.append(210)
        elif 'فر' in name:
            cats.append(216)
        elif 'قابلمه' in name or 'تابه' in name:
            cats.append(215)
    elif sheet == 'نوشیدنی ساز':
        cats.append(240)
        if 'اسپرسو' in name or 'اسپرسو' in subcat:
            cats.append(119)
        elif 'آبمیوه گیری' in name or 'آبمیوه گیری' in subcat:
            cats.append(120)
        elif 'مخلوط کن' in name or 'مخلوط کن' in subcat:
            cats.append(123)
        elif 'چای ساز' in name or 'چای ساز' in subcat:
            cats.append(211)
    elif sheet == 'سرمایش و گرمایش':
        cats.append(241)
        if 'کولر آبی' in name:
            cats.append(95)
        elif 'کولر' in name and 'گازی' in name:
            cats.append(206)
        elif 'بخاری' in name or 'پنکه' in name:
            cats.append(214)
    elif sheet == 'زیبایی و سلامت فردی':
        cats.append(212)
        if 'سشوار' in name or 'سشوار' in subcat:
            cats.append(124)
        elif 'ریش تراش' in name or 'پیرایش' in name or 'اصلاح' in name:
            cats.append(126)
            
    if not cats:
        cats = [239]
    return cats

def find_images_and_specs(pname, brand):
    found = []
    specs = ""
    # 1. First check if there is an exact folder named pname
    for cat in os.listdir(pi_dir):
        cat_p = os.path.join(pi_dir, cat)
        if not os.path.isdir(cat_p):
            continue
        sub_p = os.path.join(cat_p, pname)
        if os.path.isdir(sub_p):
            imgs = [os.path.join(sub_p, f) for f in sorted(os.listdir(sub_p)) if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))]
            if imgs:
                found.extend(imgs)
                sp_p = os.path.join(sub_p, 'مشخصات_محصول.txt')
                if os.path.exists(sp_p):
                    with open(sp_p, 'r', encoding='utf-8') as sf:
                        specs = sf.read()
                return found, specs

    # 2. Check standalone verified files in product_images_2026-10-05 with precise matching
    name_lower = pname.lower()
    for cat in os.listdir(pi_dir):
        cat_p = os.path.join(pi_dir, cat)
        if not os.path.isdir(cat_p):
            continue
        for f in os.listdir(cat_p):
            fp = os.path.join(cat_p, f)
            if os.path.isfile(fp) and fp.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')):
                f_base = os.path.splitext(f)[0].lower()
                if 'himalia' in f_base and '60mw' in f_base and 'هیمالیا' in pname and '60mw' in name_lower:
                    found.append(fp)
                elif 'niksan' in f_base and 'rf8410' in f_base and 'نیکسان' in pname and 'rf8410' in name_lower:
                    found.append(fp)
                elif 'clover' in f_base and 'rosso' in f_base and 'water' in f_base and 'کلور' in pname and 'روسو' in pname and 'آبریز' in pname:
                    found.append(fp)
                elif 'clover' in f_base and 'rosso' in f_base and 'homebar' in f_base and 'کلور' in pname and 'روسو' in pname and 'هوم بار' in pname:
                    found.append(fp)
                elif 'clover' in f_base and 'mini' in f_base and 'کلور' in pname and 'مینی' in pname:
                    found.append(fp)
                elif 'pakshoma' in f_base and 'twin-tub' in f_base and 'پاکشوما' in pname and 'دوقلو' in pname:
                    found.append(fp)
                elif 'saper' in f_base and '4025' in f_base and 'ساپر' in pname and '4025' in pname:
                    found.append(fp)
                elif 'saper' in f_base and 'iron' in f_base and 'ساپر' in pname and '4035' in pname:
                    found.append(fp)
                elif 'feridolin' in f_base and 'sw35' in f_base and 'فریدولین' in pname:
                    found.append(fp)
                elif 'panda' in f_base and '80l' in f_base and 'پاندا' in pname:
                    found.append(fp)
                elif 'pakizan' in f_base and 'blanket' in f_base and 'پاکیزان' in pname:
                    found.append(fp)
                elif 'aviz' in f_base and 'granite-pan' in f_base and 'اویز' in pname and 'تابه' in pname and 'اویزپلاس' not in pname:
                    found.append(fp)
                elif 'alton' in f_base and 'oven' in f_base and 'آلتون' in pname and '7303' in pname:
                    found.append(fp)
                elif 'willez' in f_base and 'chopper' in f_base and 'ویلز' in pname and 'خردکن' in pname:
                    found.append(fp)
                elif 'volga' in f_base and 'kettle' in f_base and 'وولگا' in pname:
                    found.append(fp)
                elif 'kenwood' in f_base and 'hm430' in f_base and 'کنوود' in pname:
                    found.append(fp)
                elif 'korkmaz' in f_base and 'a331' in f_base and 'کورکماز' in pname:
                    found.append(fp)
                elif 'parskhazar' in f_base and 'fh2000p' in f_base and 'پارس خزر' in pname and '2000' in pname:
                    found.append(fp)
                elif 'parskhazar' in f_base and '4060' in f_base and 'پارس خزر' in pname and 'پنکه' in pname:
                    found.append(fp)
                elif 'polar' in f_base and '3pn' in f_base and 'پلار' in pname:
                    found.append(fp)
                elif 'promaron' in f_base and 'پرومارون' in pname and '731' in pname:
                    found.append(fp)
                    
    # Deduplicate while preserving order
    dedup = []
    seen = set()
    for it in found:
        norm = os.path.normpath(it)
        if norm not in seen:
            seen.add(norm)
            dedup.append(norm)
    return dedup, specs

processed_products = []
for p in products:
    parent = p['parent']
    pname = parent['name']
    brand = parent['brand']
    sheet = parent['sheet']
    subcat = parent['subcat']
    
    slug = slug_mapping.get(pname, f"item-{parent['row']}")
    cats = get_product_cats(sheet, subcat, brand, pname)
    
    primary_cid = cats[-1] if cats[-1] in CAT_CANONICAL_URLS else cats[0]
    cat_name, cat_url = CAT_CANONICAL_URLS.get(primary_cid, ("دسته‌بندی محصولات", "https://hosseinykala.ir/"))
    
    imgs, specs = find_images_and_specs(pname, brand)
    
    prod_data = {
        'name': pname,
        'slug': slug,
        'type': p['type'],
        'sheet': sheet,
        'row': parent['row'],
        'price_rial': parent['price_rial'],
        'price_toman': parent['price_toman'],
        'brand': brand,
        'subcat': subcat,
        'categories': cats,
        'primary_cat_id': primary_cid,
        'primary_cat_name': cat_name,
        'primary_cat_url': cat_url,
        'images': imgs,
        'specs': specs,
        'variations': []
    }
    
    for v in p['variations']:
        prod_data['variations'].append({
            'row': v['row'],
            'name': v['name'],
            'variant_attr': v['variant_attr'],
            'price_rial': v['price_rial'],
            'price_toman': v['price_toman'],
            'sku': v['sku'],
        })
        
    processed_products.append(prod_data)

out_file = 'scratch/all_61_products.json'
with open(out_file, 'w', encoding='utf-8') as f:
    json.dump(processed_products, f, ensure_ascii=False, indent=2)

print(f"Refined plan saved to {out_file} with {len(processed_products)} products.")
