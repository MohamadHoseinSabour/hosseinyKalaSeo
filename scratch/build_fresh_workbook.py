import openpyxl
import pickle
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

COLOR_HEADER_BG = "1F4E79"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_TOMAN_HEADER_BG = "2E75B6"
COLOR_SUBCAT_BG = "EBF1F5"
COLOR_SUBCAT_TEXT = "375623"
COLOR_VAR_PARENT_BG = "D9E1F2"
COLOR_VARIATION_BG = "F2F2F2"

COLUMNS = [
    ("name", "نام محصول", 48),
    ("price_regular", "قیمت عادی (ریال)", 18),
    ("price_reg_toman", "قیمت عادی (تومان)", 18),
    ("price_base_d", "آخرین قیمت مشخص‌شده (ریال)", 18),
    ("price_diff_pct", "درصد تغییر قیمت", 15),
    ("price_sale", "قیمت حراج (ریال)", 18),
    ("price_sale_toman", "قیمت حراج (تومان)", 18),
    ("in_stock", "وضعیت موجودی", 14),
    ("main_cat", "دسته‌بندی اصلی سایت", 22),
    ("sub_cat", "زیردسته / گروه کالا", 22),
    ("brand", "برند", 16),
    ("type", "نوع محصول", 12),
    ("variant_attr", "مشخصه تنوع / متغیر", 25),
    ("product_id", "شناسه محصول (ID)", 14),
    ("variation_id", "شناسه تنوع", 14),
    ("link", "لینک مستقیم در سایت", 35),
    ("sku", "کد SKU", 14),
    ("status_site", "وضعیت در سایت", 18),
]

with open('scratch/target_data.pkl', 'rb') as f:
    target = pickle.load(f)

old_wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx')
new_wb = openpyxl.Workbook()
# remove default sheet
new_wb.remove(new_wb.active)

# 1. Copy dashboard sheet if exists
if '📊 خلاصه و داشبورد' in old_wb.sheetnames:
    ws_old = old_wb['📊 خلاصه و داشبورد']
    ws_new = new_wb.create_sheet('📊 خلاصه و داشبورد')
    ws_new.sheet_view.rightToLeft = True
    for r in range(1, ws_old.max_row + 1):
        for c in range(1, ws_old.max_column + 1):
            cell_old = ws_old.cell(row=r, column=c)
            cell_new = ws_new.cell(row=r, column=c, value=cell_old.value)
            if cell_old.has_style:
                cell_new.font = copy(cell_old.font)
                cell_new.fill = copy(cell_old.fill)
                cell_new.border = copy(cell_old.border)
                cell_new.alignment = copy(cell_old.alignment)
                cell_new.number_format = copy(cell_old.number_format)
    for m in ws_old.merged_cells.ranges:
        ws_new.merge_cells(str(m))

# 2. Build 8 category sheets
total_priced_written = 0

for sname, brands_dict in target.items():
    ws = new_wb.create_sheet(sname)
    ws.sheet_view.rightToLeft = True
    
    # Headers
    ws.row_dimensions[1].height = 28
    for col_idx, (_, h_text, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font = Font(bold=True, color=COLOR_HEADER_TEXT, size=10, name="Calibri")
        is_formula = col_idx in [3, 5, 7]
        h_bg = COLOR_TOMAN_HEADER_BG if is_formula else COLOR_HEADER_BG
        cell.fill = PatternFill(fill_type='solid', fgColor=h_bg)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    
    curr_r = 2
    for brand, blocks in brands_dict.items():
        # Banner row
        ws.row_dimensions[curr_r].height = 22
        ws.merge_cells(start_row=curr_r, start_column=1, end_row=curr_r, end_column=len(COLUMNS))
        sub_c = ws.cell(row=curr_r, column=1, value=f"  📂 {sname}  ◀  {brand}")
        sub_c.font = Font(bold=True, color=COLOR_SUBCAT_TEXT, size=10, name="Calibri")
        sub_c.fill = PatternFill(fill_type='solid', fgColor=COLOR_SUBCAT_BG)
        sub_c.alignment = Alignment(horizontal='right', vertical='center')
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=curr_r, column=c).border = border_thin
        curr_r += 1
        
        for b in blocks:
            is_var = len(b['variations']) > 0
            ws.row_dimensions[curr_r].height = 21
            p_fill = PatternFill(fill_type='solid', fgColor=COLOR_VAR_PARENT_BG) if is_var else None
            
            p_name = b['name']
            p_b = b['price_b'] if not is_var else None
            p_d = b['price_d'] if not is_var else None
            p_f = int(round(p_b * 0.95)) if p_b else None
            p_stk = "✓ موجود" if (p_b or is_var) else "ناموجود"
            p_status_r = b.get('status_r', 'ثبت‌شده در سایت')
            
            if p_b:
                total_priced_written += 1
                
            row_data = [
                p_name,
                p_b,
                f'=IF(OR(B{curr_r}="",B{curr_r}=0),"",B{curr_r}/10)',
                p_d,
                f'=IF(ISNUMBER(D{curr_r}),(B{curr_r}-D{curr_r})/D{curr_r},"")',
                p_f,
                f'=IF(OR(F{curr_r}="",F{curr_r}=0),"",F{curr_r}/10)',
                p_stk,
                sname,
                f"{sname} {brand}",
                brand,
                "متغیر" if is_var else "ساده",
                f"شامل {len(b['variations'])} تنوع" if is_var else "",
                "", # ID
                "", # var ID
                "", # link
                "", # SKU
                p_status_r
            ]
            
            for col_i, val in enumerate(row_data, start=1):
                c = ws.cell(row=curr_r, column=col_i, value=val)
                c.border = border_thin
                if p_fill: c.fill = p_fill
                if col_i == 1:
                    c.alignment = Alignment(horizontal='right', vertical='center')
                    c.font = Font(bold=is_var, size=9.5, name="Calibri")
                elif col_i in [2, 3, 4, 6, 7]:
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 5:
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '0.0%'
                else:
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.font = Font(size=9.5, name="Calibri")
            curr_r += 1
            
            for v in b['variations']:
                ws.row_dimensions[curr_r].height = 19
                v_fill = PatternFill(fill_type='solid', fgColor=COLOR_VARIATION_BG)
                v_b = v['price_b']
                v_d = v['price_d']
                v_f = int(round(v_b * 0.95)) if v_b else None
                v_stk = "✓ موجود" if v_b else "ناموجود"
                v_status_r = v.get('status_r', p_status_r)
                
                if v_b:
                    total_priced_written += 1
                    
                v_row_data = [
                    v['name'],
                    v_b,
                    f'=IF(OR(B{curr_r}="",B{curr_r}=0),"",B{curr_r}/10)',
                    v_d,
                    f'=IF(ISNUMBER(D{curr_r}),(B{curr_r}-D{curr_r})/D{curr_r},"")',
                    v_f,
                    f'=IF(OR(F{curr_r}="",F{curr_r}=0),"",F{curr_r}/10)',
                    v_stk,
                    sname,
                    f"{sname} {brand}",
                    brand,
                    "تنوع",
                    v['name'].replace('↳', '').strip(),
                    "", "", "", "",
                    v_status_r
                ]
                
                for col_i, val in enumerate(v_row_data, start=1):
                    c = ws.cell(row=curr_r, column=col_i, value=val)
                    c.border = border_thin
                    c.fill = v_fill
                    if col_i == 1:
                        c.alignment = Alignment(horizontal='right', vertical='center')
                        c.font = Font(size=9, color="333333", name="Calibri")
                    elif col_i in [2, 3, 4, 6, 7]:
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i == 5:
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '0.0%'
                    else:
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.font = Font(size=9, name="Calibri")
                curr_r += 1

# 3. Copy documentation sheet
if '📖 راهنمای کاربری' in old_wb.sheetnames:
    ws_old = old_wb['📖 راهنمای کاربری']
    ws_new = new_wb.create_sheet('📖 راهنمای کاربری')
    ws_new.sheet_view.rightToLeft = True
    for r in range(1, ws_old.max_row + 1):
        for c in range(1, ws_old.max_column + 1):
            cell_old = ws_old.cell(row=r, column=c)
            cell_new = ws_new.cell(row=r, column=c, value=cell_old.value)
            if cell_old.has_style:
                cell_new.font = copy(cell_old.font)
                cell_new.fill = copy(cell_old.fill)
                cell_new.border = copy(cell_old.border)
                cell_new.alignment = copy(cell_old.alignment)
                cell_new.number_format = copy(cell_old.number_format)
    for m in ws_old.merged_cells.ranges:
        ws_new.merge_cells(str(m))

print(f"Total priced rows written to fresh workbook: {total_priced_written}")
new_wb.save('scratch/fresh_reorganized_catalog.xlsx')
print("Saved scratch/fresh_reorganized_catalog.xlsx successfully!")
