import openpyxl
import sys
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

ids = []
links = []
skus = []

for sname in wb.sheetnames:
    if sname.startswith('📊') or sname.startswith('📖'): continue
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, 14).value
        link = ws.cell(r, 16).value
        sku = ws.cell(r, 17).value
        name = ws.cell(r, 1).value
        
        if pid:
            ids.append((str(pid).strip(), sname, r, name))
        if link:
            links.append((str(link).strip(), sname, r, name))
        if sku:
            skus.append((str(sku).strip(), sname, r, name))

id_counts = Counter([x[0] for x in ids])
link_counts = Counter([x[0] for x in links])
sku_counts = Counter([x[0] for x in skus])

print("=== CHECK DUPLICATE IDS (for products, excluding variations sharing parent ID) ===")
# Note: variations of a variable product share the same parent product_id in col 14, which is normal.
# But distinct parent products should NOT share the same ID.
parent_ids = []
for sname in wb.sheetnames:
    if sname.startswith('📊') or sname.startswith('📖'): continue
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or '')
        pid = ws.cell(r, 14).value
        if pid and not name.startswith('↳'):
            parent_ids.append((str(pid).strip(), sname, r, name))

p_id_counts = Counter([x[0] for x in parent_ids])
dup_p_ids = {k: v for k, v in p_id_counts.items() if v > 1}
if dup_p_ids:
    print(f"Warning: duplicate parent IDs found: {dup_p_ids}")
    for k in dup_p_ids:
        print(f"ID {k}:", [x for x in parent_ids if x[0] == k])
else:
    print("✓ هیچ شناسه محصول (ID) تکراری بین محصولات والد وجود ندارد.")

print("\n=== CHECK DUPLICATE SKUs ===")
dup_skus = {k: v for k, v in sku_counts.items() if v > 1}
if dup_skus:
    print(f"Warning: duplicate SKUs found: {dup_skus}")
else:
    print("✓ هیچ SKU تکراری در کل کاتالوگ وجود ندارد.")
