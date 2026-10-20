import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
all_rows = []
for sname in wb.sheetnames:
    if sname.startswith('📊') or sname.startswith('📖'):
        continue
    ws = wb[sname]
    curr_parent = None
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or '').strip()
        if name and not name.startswith('  📂'):
            item = {'sheet': sname, 'row': r, 'name': name, 'brand': ws.cell(r, 11).value, 'status': ws.cell(r, 18).value}
            if name.startswith('↳'):
                item['parent'] = curr_parent
            else:
                curr_parent = name
            all_rows.append(item)

exact_names = {}
for r in all_rows:
    if not r['name'].startswith('↳'):
        exact_names.setdefault(r['name'], []).append(r)

print('=== 1. EXACT DUPLICATE PRODUCT NAMES ===')
found_dup = False
for name, rows in exact_names.items():
    if len(rows) > 1:
        found_dup = True
        print(f"Duplicate product name: '{name}'")
        for x in rows:
            print(f"   [{x['sheet']}] R{x['row']} Brand={x['brand']} Status={x['status']}")
if not found_dup:
    print('✓ هیچ نام محصول تکراری در کل کاتالوگ وجود ندارد.')

print('\n=== 2. DUPLICATE VARIATIONS UNDER SAME PARENT ===')
var_by_parent = {}
for r in all_rows:
    if r['name'].startswith('↳') and r.get('parent'):
        var_by_parent.setdefault(r['parent'], []).append(r)

found_var_dup = False
for parent, vars in var_by_parent.items():
    var_names = {}
    for v in vars:
        var_names.setdefault(v['name'], []).append(v)
    for vname, vrows in var_names.items():
        if len(vrows) > 1:
            found_var_dup = True
            print(f"Parent: '{parent}' has duplicate '{vname}':")
            for x in vrows:
                print(f"   R{x['row']}")
if not found_var_dup:
    print('✓ هیچ تنوع تکراری زیر هیچ محصول والدی وجود ندارد.')
