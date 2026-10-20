import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
ws = wb['یخچال و فریزر']

print("=== ALL FRIDGES WITH 27 OR 35 OR تیتانیوم OR هوم بار ===")
for r in range(2, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or '')
    attr = str(ws.cell(r, 13).value or '')
    brand = str(ws.cell(r, 11).value or '')
    status = str(ws.cell(r, 18).value or '')
    
    if any(k in name or k in attr for k in ['27', '۲۷', '35', '۳۵', 'تیتانیوم', 'هوم بار', 'درب بار', 'بار']):
        print(f"Row {r:3d} [{brand:12s}] {name:65s} | attr: {attr:20s} | status: {status}")
