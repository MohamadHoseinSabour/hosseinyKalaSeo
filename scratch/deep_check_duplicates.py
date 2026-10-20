import openpyxl
import sys
import re
from difflib import SequenceMatcher

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

all_rows = []
for sname in wb.sheetnames:
    if sname.startswith('📊') or sname.startswith('📖'):
        continue
    ws = wb[sname]
    curr_parent = None
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name and not str(name).startswith('  📂'):
            item = {
                'sheet': sname,
                'row': r,
                'name': str(name).strip(),
                'brand': str(ws.cell(r, 11).value or '').strip(),
                'type': str(ws.cell(r, 12).value or '').strip(),
                'attr': str(ws.cell(r, 13).value or '').strip(),
                'status': str(ws.cell(r, 18).value or '').strip(),
                'is_new': (ws.cell(r, 18).value == 'در سایت تعریف نشده'),
                'parent': curr_parent if str(name).strip().startswith('↳') else None
            }
            if not str(name).strip().startswith('↳'):
                curr_parent = str(name).strip()
            all_rows.append(item)

print(f"Total rows (products + variations): {len(all_rows)}")
new_rows = [x for x in all_rows if x['is_new']]
existing_rows = [x for x in all_rows if not x['is_new']]
print(f"New rows (در سایت تعریف نشده): {len(new_rows)}")
print(f"Existing rows: {len(existing_rows)}")

# 1. Exact duplicate checks
print("\n--- 1. EXACT DUPLICATES AMONG PRODUCTS (non-variations) ---")
exact_names = {}
for r in all_rows:
    if not r['name'].startswith('↳'):
        exact_names.setdefault(r['name'], []).append(r)

for name, rows in exact_names.items():
    if len(rows) > 1:
        print(f"Duplicate product name: '{name}'")
        for x in rows:
            print(f"   [{x['sheet']}] R{x['row']} Brand={x['brand']} Status={x['status']}")

# 2. Duplicate variations within the same parent
print("\n--- 2. DUPLICATE VARIATIONS UNDER THE SAME PARENT ---")
var_by_parent = {}
for r in all_rows:
    if r['name'].startswith('↳') and r['parent']:
        var_by_parent.setdefault(r['parent'], []).append(r)

for parent, vars in var_by_parent.items():
    var_names = {}
    for v in vars:
        var_names.setdefault(v['name'], []).append(v)
    for vname, vrows in var_names.items():
        if len(vrows) > 1:
            print(f"Parent: '{parent}' has duplicate variation '{vname}':")
            for x in vrows:
                print(f"   R{x['row']} Status={x['status']}")

# 3. Fuzzy similarity between each NEW item and all EXISTING items
print("\n--- 3. HIGH SIMILARITY BETWEEN NEW ITEMS AND EXISTING ITEMS ---")
def norm(t):
    for f, e in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'): t = t.replace(f, e)
    t = t.replace('ي', 'ی').replace('ك', 'ک').replace('‌', ' ').replace('-', ' ').replace('_', ' ')
    t = re.sub(r'[\(\)\[\]\.\,\/]', ' ', t)
    return re.sub(r'\s+', ' ', t).strip().lower()

for n in new_rows:
    n_norm = norm(n['name'])
    n_tokens = set(n_norm.split())
    # Compare with existing rows
    for e in existing_rows:
        e_norm = norm(e['name'])
        # If it's a variation, also include parent name
        if e['parent']:
            e_norm = norm(f"{e['parent']} {e['name']}")
        
        ratio = SequenceMatcher(None, n_norm, e_norm).ratio()
        
        # Token overlap
        e_tokens = set(e_norm.split())
        common = n_tokens.intersection(e_tokens)
        meaningful_common = [w for w in common if w not in ['مدل', 'برقی', 'و', 'فوت', 'لیتر', 'سفید', 'مشکی', 'استیل', 'طرح', 'سری', 'ها', 'رنگ', '↳']]
        
        # Check if same brand and high similarity or specific model numbers match
        same_brand = (n['brand'] == e['brand']) and n['brand'] != '' and n['brand'] != 'متفرقه'
        
        # Extract numbers in n and e
        n_nums = re.findall(r'\d+', n_norm)
        e_nums = re.findall(r'\d+', e_norm)
        common_nums = set(n_nums).intersection(set(e_nums))
        
        is_suspicious = False
        reason = ""
        
        if ratio > 0.75:
            is_suspicious = True
            reason = f"High string ratio ({ratio:.2f})"
        elif same_brand and len(meaningful_common) >= 3:
            is_suspicious = True
            reason = f"Same brand '{n['brand']}' and 3+ common words: {meaningful_common}"
        elif same_brand and common_nums and len(meaningful_common) >= 2:
            is_suspicious = True
            reason = f"Same brand '{n['brand']}' and matching numbers {common_nums}"
        elif common_nums and ('مدل' in n_norm or 'مدل' in e_norm) and any(len(num) >= 3 for num in common_nums):
            # check if same category
            if n['sheet'] == e['sheet']:
                is_suspicious = True
                reason = f"Same sheet '{n['sheet']}' and common model number {common_nums}"

        if is_suspicious:
            print(f"POTENTIAL DUPLICATE:")
            print(f"   NEW: [{n['sheet']}] R{n['row']} Brand={n['brand']} '{n['name']}'")
            print(f"   EXISTING: [{e['sheet']}] R{e['row']} Brand={e['brand']} '{e['name']}' (Parent: {e['parent']})")
            print(f"   REASON: {reason}\n")
