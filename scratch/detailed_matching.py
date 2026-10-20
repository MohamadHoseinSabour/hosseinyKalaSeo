import os
import sys
import json
import re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
category_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

unregistered = []
for sheet in category_sheets:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    current_parent = None
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or '📂' in str(name):
            continue
        name_str = str(name).strip()
        status = str(ws.cell(row=r, column=18).value or '').strip()
        stock = str(ws.cell(row=r, column=8).value or '').strip()
        pid = ws.cell(row=r, column=14).value
        ptype = str(ws.cell(row=r, column=12).value or 'ساده').strip()
        
        # Track parent for variations
        if not name_str.startswith('↳'):
            current_parent = name_str
            parent_name = None
        else:
            parent_name = current_parent

        if status == 'در سایت تعریف نشده' or (not pid and status != 'فعلا دست نگه دار'):
            unregistered.append({
                'sheet': sheet,
                'row': r,
                'name': name_str,
                'parent_name': parent_name,
                'status': status,
                'stock': stock,
                'price_rial': ws.cell(row=r, column=2).value,
                'price_toman': ws.cell(row=r, column=3).value,
                'main_cat': str(ws.cell(row=r, column=9).value or sheet).strip(),
                'subcat': str(ws.cell(row=r, column=10).value or '').strip(),
                'brand': str(ws.cell(row=r, column=11).value or '').strip(),
                'ptype': ptype,
                'variant_attr': str(ws.cell(row=r, column=13).value or '').strip(),
                'sku': ws.cell(row=r, column=17).value,
            })

print(f"Total unregistered items in Excel: {len(unregistered)}")

# Check matching with product_images_2026-10-05 folders and files
pi_base = 'product_images_2026-10-05'
matched = []
unmatched = []

for item in unregistered:
    name = item['name']
    brand = item['brand']
    pname = item['parent_name']
    full_name = f"{pname} {name}" if pname else name
    
    # 1. Look for folder with exact or partial name in product_images_2026-10-05
    found_folder = None
    found_files = []
    
    for cat in os.listdir(pi_base):
        cat_p = os.path.join(pi_base, cat)
        if not os.path.isdir(cat_p):
            continue
        for sub in os.listdir(cat_p):
            sub_p = os.path.join(cat_p, sub)
            if os.path.isdir(sub_p):
                # Check if sub folder name matches
                clean_sub = sub.strip()
                clean_name = name.replace('↳', '').strip()
                if clean_sub in clean_name or clean_name in clean_sub or (pname and clean_sub in pname):
                    # Check images inside
                    imgs = [os.path.join(sub_p, f) for f in os.listdir(sub_p) if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))]
                    if imgs:
                        found_folder = sub_p
                        found_files = imgs
                        break
            elif os.path.isfile(sub_p) and sub_p.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')):
                # Check standalone file
                base_f = os.path.splitext(sub)[0].lower()
                # e.g., kenwood-hm430-hand-mixer vs 'همزن برقی کنوود مدل HM430'
                if brand and brand.lower() in base_f:
                    # check code
                    codes = re.findall(r'[a-zA-Z0-9]+', name)
                    for c in codes:
                        if len(c) >= 3 and c.lower() in base_f:
                            found_files = [sub_p]
                            break
        if found_files:
            break
            
    if found_files:
        item['images'] = found_files
        item['image_folder'] = found_folder
        matched.append(item)
    else:
        unmatched.append(item)

print(f"\nMatched with images in product_images_2026-10-05: {len(matched)}")
print(f"Unmatched: {len(unmatched)}")

for m in matched:
    print(f"  [Row {m['row']:2d}] {m['sheet']} | {m['name']} ({m['brand']}) -> {len(m['images'])} images: {[os.path.basename(x) for x in m['images']]}")

print("\n--- UNMATCHED ITEMS ---")
for u in unmatched:
    print(f"  [Row {u['row']:2d}] {u['sheet']} | {u['name']} (Brand: {u['brand']}, Parent: {u['parent_name']})")
