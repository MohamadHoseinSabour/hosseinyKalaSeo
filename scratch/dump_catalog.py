import openpyxl

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

with open('scratch/catalog_full_dump.txt', 'w', encoding='utf-8') as out:
    for sname in wb.sheetnames:
        if sname in ['📊 خلاصه و داشبورد', '📖 راهنمای کاربری']:
            continue
        ws = wb[sname]
        out.write(f"\n{'='*30} SHEET: {sname} {'='*30}\n")
        for r in range(1, ws.max_row + 1):
            vA = ws.cell(row=r, column=1).value
            vB = ws.cell(row=r, column=2).value
            vD = ws.cell(row=r, column=4).value
            vF = ws.cell(row=r, column=6).value
            vR = ws.cell(row=r, column=18).value
            if vA is not None:
                out.write(f"Row {r:3d} | A: {str(vA):<65} | B: {str(vB):<12} | D: {str(vD):<12} | R: {str(vR)}\n")

print("Dumped catalog to scratch/catalog_full_dump.txt")
