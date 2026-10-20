import openpyxl

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

print("=== ALL PRODUCTS UNDER 'متفرقه' IN ALL SHEETS ===")
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    current_brand = ""
    current_section = ""
    for r in range(1, ws.max_row + 1):
        vA = ws.cell(row=r, column=1).value
        vB = ws.cell(row=r, column=2).value
        vR = ws.cell(row=r, column=18).value
        if not vA:
            continue
        vA_str = str(vA).strip()
        if '📂' in vA_str:
            current_section = vA_str
            current_brand = vA_str.split('◀')[-1].strip() if '◀' in vA_str else vA_str
            continue
        if current_brand == 'متفرقه':
            print(f"[{sheetname}] Row {r:2d} | A: {vA_str} | Price: {vB} | Status: {vR}")
