import openpyxl

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
ws = wb['صوتی و تصویری']
for r in range(1, ws.max_row + 1):
    valA = ws.cell(row=r, column=1).value
    valB = ws.cell(row=r, column=2).value
    valR = ws.cell(row=r, column=18).value
    if valA:
        print(f"Row {r:2d}: A={str(valA):<55} | B={str(valB):<12} | R={valR}")
