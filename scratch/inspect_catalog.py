import openpyxl, os, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)

total_undefined = []
total_hold = []

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    headers = [ws.cell(1, c).value for c in range(1, 25)]
    
    undefined_in_sheet = []
    hold_in_sheet = []
    total_products = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        status_r = ws.cell(r, 18).value
        # If it's a category header
        if str(name).strip().startswith('📂'):
            continue
        total_products += 1
        pid = ws.cell(r, 15).value
        price = ws.cell(r, 2).value
        url = ws.cell(r, 16).value
        brand = ws.cell(r, 10).value
        cat = ws.cell(r, 12).value
        
        info = {
            'sheet': sheetname,
            'row': r,
            'name': name,
            'status_r': status_r,
            'pid': pid,
            'price': price,
            'url': url,
            'brand': brand,
            'category': cat,
            'stock': ws.cell(r, 8).value
        }
        
        if status_r == 'در سایت تعریف نشده':
            undefined_in_sheet.append(info)
            total_undefined.append(info)
        elif status_r == 'فعلا دست نگه دار':
            hold_in_sheet.append(info)
            total_hold.append(info)
            
    if undefined_in_sheet or hold_in_sheet:
        print(f"\n--- Sheet: {sheetname} (Total rows checked: {total_products}) ---")
        if undefined_in_sheet:
            print(f"  [در سایت تعریف نشده]: {len(undefined_in_sheet)}")
            for item in undefined_in_sheet:
                print(f"    Row {item['row']}: {item['name']} | PID: {item['pid']} | Brand: {item['brand']} | Price: {item['price']} | Stock: {item['stock']}")
        if hold_in_sheet:
            print(f"  [فعلا دست نگه دار]: {len(hold_in_sheet)}")
            for item in hold_in_sheet:
                print(f"    Row {item['row']}: {item['name']} | PID: {item['pid']} | Brand: {item['brand']}")

print(f"\n==========================================")
print(f"TOTAL UNDEFINED ('در سایت تعریف نشده'): {len(total_undefined)}")
print(f"TOTAL HOLD ('فعلا دست نگه دار'): {len(total_hold)}")
