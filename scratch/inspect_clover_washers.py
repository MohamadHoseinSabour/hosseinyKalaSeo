import json
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

print("=== PRODUCTS.JSON (SITE) ===")
with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

for p in products:
    name = p.get('name', '')
    if 'کلور' in name and ('لباسشویی' in name or 'wash' in str(p).lower() or '8' in name):
        print(f"ID: {p.get('id')} | Name: '{name}' | Type: {p.get('type')}")
        for v in p.get('variations', []):
            print(f"   Var: {v.get('id')} | {v.get('name')} | {v.get('attributes')}")

print("\n=== CURRENT CATALOG (شستشو و نظافت -> کلور) ===")
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
ws = wb['شستشو و نظافت']
for r in range(1, ws.max_row + 1):
    b = ws.cell(r, 11).value
    if b == 'کلور':
        print(f"Row {r:2d} | Col L: {ws.cell(r, 12).value} | Col M: {ws.cell(r, 13).value} | Name: '{ws.cell(r, 1).value}' | Status: {ws.cell(r, 18).value}")
