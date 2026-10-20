import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
category_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

stock_vals = set()
status_vals = set()

for sheet in category_sheets:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or '📂' in str(name):
            continue
        h_val = ws.cell(row=r, column=8).value
        r_val = ws.cell(row=r, column=18).value
        if h_val: stock_vals.add(str(h_val).strip())
        if r_val: status_vals.add(str(r_val).strip())

print("Unique values in Col H (وضعیت موجودی):")
for v in stock_vals:
    print(f"  '{v}'")

print("\nUnique values in Col R (وضعیت در سایت):")
for v in status_vals:
    print(f"  '{v}'")
