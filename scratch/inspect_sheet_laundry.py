import openpyxl

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
ws = wb['شستشو و نظافت']
for r in range(1, 100):
    valA = ws.cell(row=r, column=1).value
    valB = ws.cell(row=r, column=2).value
    valR = ws.cell(row=r, column=18).value
    if valA is not None:
        print(f"Row {r:2d}: A={str(valA):<50} | B={str(valB):<15} | R={str(valR)}")
