import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
file_path = r'C:/Users/mmhhs/.gemini/antigravity/brain/0354df7e-373f-404f-828c-335674433d4b/.user_uploaded/media_1791751795904_cfbfa0e4.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

print("Sheets:", wb.sheetnames)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\nSheet '{sname}': max_row={ws.max_row}, max_col={ws.max_column}")
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"R{r}: {row_vals}")
            
    # Check if there are non-empty cells in other columns
    print("Non-empty columns summary:")
    for c in range(1, ws.max_column + 1):
        non_empty = [ws.cell(r, c).value for r in range(1, ws.max_row + 1) if ws.cell(r, c).value is not None]
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): count={len(non_empty)}, first 3={non_empty[:3]}")
