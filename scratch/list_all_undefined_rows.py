import openpyxl, sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        status_r = ws.cell(r, 18).value
        name = ws.cell(r, 1).value
        if not name or str(name).strip().startswith('📂'):
            continue
        if status_r == 'در سایت تعریف نشده':
            print(f"Sheet: {sname:20s} | Row {r:3d} | PID: {str(ws.cell(r, 15).value):5s} | Name: {name}")
