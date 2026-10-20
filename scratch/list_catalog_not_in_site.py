import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
not_in_site = []

for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in range(5, ws.max_row + 1):
        status = ws.cell(row=row, column=18).value
        if status == 'در سایت تعریف نشده':
            name = ws.cell(row=row, column=1).value
            brand = ws.cell(row=row, column=8).value
            model = ws.cell(row=row, column=9).value
            price = ws.cell(row=row, column=3).value
            not_in_site.append({
                'sheet': sheet,
                'name': name,
                'brand': brand,
                'model': model,
                'price': price
            })

print(f"Total 'در سایت تعریف نشده' items: {len(not_in_site)}")
for idx, item in enumerate(not_in_site, 1):
    print(f"[{idx:02d}] [{item['sheet']}] {item['name']} | Brand: {item['brand']} | Model: {item['model']}")
