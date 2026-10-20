import sys
import json
import openpyxl
import re

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# 1. Load raw list
with open('data/raw_parsed_new_list.json', 'r', encoding='utf-8') as f:
    raw_list = json.load(f)

# 2. Load current catalog
wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
catalog = []
for sname in wb.sheetnames:
    if sname.startswith('📊') or sname.startswith('📖'):
        continue
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v:
            catalog.append({
                'sheet': sname,
                'row': r,
                'name': str(v).strip(),
                'brand': str(ws.cell(r, 11).value or '').strip(),
                'type': str(ws.cell(r, 12).value or '').strip(),
                'attr': str(ws.cell(r, 13).value or '').strip(),
                'status': str(ws.cell(r, 18).value or '').strip()
            })

def to_en(t):
    for f, e in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        t = t.replace(f, e)
    return t

def norm(t):
    if not t: return ''
    t = to_en(str(t).lower())
    t = t.replace('ي', 'ی').replace('ك', 'ک').replace('‌', ' ').replace('-', ' ').replace('_', ' ')
    t = re.sub(r'[\(\)\[\]\.\,\/]', ' ', t)
    return re.sub(r'\s+', ' ', t).strip()

# Let's inspect what items in raw_list exist in catalog
matched_items = []
unmatched_items = []

for idx, it in enumerate(raw_list, 1):
    desc = it['desc']
    qty = it['qty']
    if 'کسرواضافات' in desc:
        continue
    nd = norm(desc)
    
    match = None
    
    # 1. صوتی و تصویری
    if any(k in nd for k in ['aversen', 'هاورسن', 'هاورسون', 'ال ای دی']):
        for c in catalog:
            if c['sheet'] == 'صوتی و تصویری':
                for sz in ['85', '75', '65', '55', '32']:
                    if sz in nd and sz in norm(c['name']):
                        match = c
                        break
            if match: break
            
    # 2. تصفیه آب و آبسردکن
    if not match:
        if 'افترون' in nd and 'سردکن' in nd:
            match = next((c for c in catalog if c['sheet'] == 'تصفیه کننده آب' and 'افترون' in c['brand']), None)
        elif 'نیکائی' in nd or 'نیکایی' in nd:
            match = next((c for c in catalog if c['sheet'] == 'تصفیه کننده آب' and 'نیکایی' in c['brand']), None)
        elif 'وست واتر' in nd:
            match = next((c for c in catalog if c['sheet'] == 'تصفیه کننده آب' and 'وست واتر' in c['brand']), None)

    # 3. یخچال و فریزر
    if not match:
        if 'کلور' in nd and ('یخچال' in nd or 'فریزر' in nd):
            for c in catalog:
                if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'کلور':
                    cn = norm(c['name'])
                    ca = norm(c['attr'])
                    if 'اکونومی' in nd and 'اکونومی' in cn:
                        if ('استیل' in nd and 'استیل' in ca) or ('سفید' in nd and 'سفید' in ca):
                            match = c; break
                    elif '70' in nd and 'روسو' in nd and 'روسو' in cn:
                        if ('آبریز' in nd or 'ابریز' in nd) and 'آبریز' in cn:
                            if ('استیل' in nd and 'استیل' in ca) or ('سفید' in nd and 'سفید' in ca):
                                match = c; break
                        elif ('هوم' in nd or 'هوبا' in nd) and 'هوم' in cn:
                            if ('استیل' in nd and 'استیل' in ca) or ('سفید' in nd and 'سفید' in ca):
                                match = c; break
                    elif '60' in nd and 'گلوری' in nd and 'گلوری' in cn:
                        if 'پلاس' in nd and 'پلاس' in cn:
                            if 'سفید' in nd and 'سفید' in ca: match = c; break
                        elif 'کره' in nd and 'سفید' in ca: match = c; break
                        elif 'استیل' in nd and 'استیل' in ca: match = c; break
                    elif '60' in nd and 'روسو' in nd and 'روسو' in cn:
                        if 'مشکی' in nd and 'مشکی' in ca: match = c; break
                        elif 'استیل' in nd and 'استیل' in ca: match = c; break
                        elif 'سفید' in nd and 'سفید' in ca: match = c; break
                    elif 'مینی' in nd and 'مینی' in cn:
                        for col in ['آبی پاستیلی', 'سبز پاستیلی', 'سفید طرحدار', 'مشکی طرحدار', 'آبی طرحدار', 'سبز طرحدار', 'قرمز']:
                            if norm(col) in nd and norm(col) in ca:
                                match = c; break
                        if match: break
                    elif '101' in nd and '101' in cn:
                        if 'کمبی' in nd and 'کمبی' in cn: match = c; break
                        elif 'استیل' in nd and 'استیل' in ca: match = c; break
                        elif 'سفید' in nd and 'سفید' in ca: match = c; break
            # هیمالیا یخچال
        elif 'هیمالیا' in nd and ('یخچال' in nd or 'فریزر' in nd):
            for c in catalog:
                if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'هیمالیا':
                    cn = norm(c['name'])
                    ca = norm(c['attr'])
                    if 'آیس پول' in cn or '60mw' in cn:
                        if 'استیل' in nd and 'استیل' in ca: match = c; break
                        elif ('چرم' in nd or 'سفید' in nd) and 'چرم' in ca: match = c; break
                    elif '5 کاره' in cn and '5 کاره' in nd: match = c; break
                    elif 'امگا' in cn and 'omega' in nd: match = c; break
                    elif '70 30' in cn or '850' in cn:
                        if '850' in nd or '70' in nd: match = c; break
        elif 'ایستکول' in nd and '9' in nd:
            for c in catalog:
                if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'ایستکول' and '9' in norm(c['name']):
                    if 'سفید' in nd and 'سفید' in norm(c['attr']): match = c; break
                    elif ('مدادی' in nd or 'استیل' in nd) and 'استیل' in norm(c['attr']): match = c; break
        elif 'بست' in nd and ('یخچال' in nd or 'bost' in nd):
            match = next((c for c in catalog if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'بست'), None)
        elif 'دیپونت' in nd or 'دیپوینت' in nd:
            for c in catalog:
                if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'دیپوینت':
                    cn = norm(c['name'])
                    ca = norm(c['attr'])
                    if 't7' in nd and 't7' in cn:
                        if 'سفید' in nd and 'سفید' in ca: match = c; break
                        elif 'استیل' in nd and ('نقره' in ca or 'استیل' in ca): match = c; break
                    elif ('boss' in nd or 'boos' in nd) and 'boss' in cn:
                        if 'سفید' in nd and 'سفید' in ca: match = c; break
                        elif 'سیلور' in nd and ('استیل' in ca or 'سیلور' in ca): match = c; break
                    elif 'دیسنت' in nd and 'دیسنت' in cn:
                        match = c; break
                    elif 'استیل' in nd and 'استیل' in cn:
                        match = c; break
        elif 'نیکسان' in nd and 'rf84102e2' in nd:
            for c in catalog:
                if c['sheet'] == 'یخچال و فریزر' and c['brand'] == 'نیکسان':
                    ca = norm(c['attr'])
                    if 'استیل' in nd and 'استیل' in ca: match = c; break
                    elif 'چرم' in nd and 'چرم' in ca: match = c; break
        elif 'الکتروسان' in nd and 'صندوقی' in nd:
            match = next((c for c in catalog if c['sheet'] == 'یخچال و فریزر' and 'الکتروسان' in c['brand']), None)
        elif 'تکنوسان' in nd and 'صندوقی' in nd:
            match = next((c for c in catalog if c['sheet'] == 'یخچال و فریزر' and 'تکنوسان' in c['brand']), None)
        elif 'ریگور' in nd:
            match = next((c for c in catalog if c['sheet'] == 'یخچال و فریزر' and 'ریگور' in c['brand']), None)

    # 4. شستشو و نظافت
    if not match:
        if 'پاکشوما' in nd and ('لباسشویی' in nd or 'دوقلو' in nd):
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'پاکشوما':
                    if '21' in nd and '21' in norm(c['name'] + ' ' + c['attr']): match = c; break
                    elif '15' in nd and '15' in norm(c['name'] + ' ' + c['attr']): match = c; break
        elif 'پاکیزان' in nd and 'پتوشور' in nd:
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'پاکیزان':
                    if '30' in nd and '30' in norm(c['name'] + ' ' + c['attr']): match = c; break
                    elif '40' in nd and '40' in norm(c['name'] + ' ' + c['attr']): match = c; break
        elif 'فریدولین' in nd:
            match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'فریدولین' in c['brand']), None)
        elif 'دوو' in nd and 'لباسشویی' in nd:
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'دوو':
                    if '990s' in nd and '990s' in norm(c['name'] + ' ' + c['attr']): match = c; break
                    elif '991s' in nd and '991s' in norm(c['name'] + ' ' + c['attr']): match = c; break
                    elif '8' in nd and '8' in norm(c['name'] + ' ' + c['attr']): match = c; break
        elif 'هیمالیا' in nd and 'لباسشویی' in nd:
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'هیمالیا':
                    if 'استیل' in nd and 'استیل' in norm(c['attr']): match = c; break
                    elif 'سفید' in nd and 'سفید' in norm(c['attr']): match = c; break
        elif 'کلور' in nd and 'لباسشویی' in nd:
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'کلور':
                    if '9tw1' in nd and '9tw1' in norm(c['name']): match = c; break
                    elif '10' in nd and ('سبز' in nd and 'سبز' in norm(c['attr'])): match = c; break
                    elif '10' in nd and ('کرم' in nd and 'کرم' in norm(c['attr'])): match = c; break
                    elif '8' in nd and 'اینوکس' in nd and 'اینوکس' in norm(c['name']): match = c; break
        elif 'ظرفشویی' in nd:
            for c in catalog:
                if c['sheet'] == 'شستشو و نظافت':
                    if '15' in nd and '15' in norm(c['name']): match = c; break
                    elif '16' in nd and '16' in norm(c['name']): match = c; break
        elif 'جارو' in nd:
            if 'کلور' in nd and 'سطلی' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'کلور' and 'جاروبرقی' in c['name']), None)
            elif 'بست' in nd or 'bost' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and c['brand'] == 'بست'), None)
            elif 'پاندا' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'پاندا' in c['brand']), None)
            elif 'آراز' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'آراز' in c['brand']), None)
            elif 'سران' in nd:
                for c in catalog:
                    if c['sheet'] == 'شستشو و نظافت' and 'سران' in c['brand']:
                        if 'زرد' in nd and 'زرد' in norm(c['attr']): match = c; break
                        elif 'مشکی' in nd and 'مشکی' in norm(c['attr']): match = c; break
            elif 'هیتاچی' in nd and '1200' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'هیتاچی' in c['brand'] and '1200' in c['name']), None)
            elif 'دونالکس' in nd and '750' in nd:
                for c in catalog:
                    if c['sheet'] == 'شستشو و نظافت' and 'دونالکس' in c['brand'] and '750' in c['name']:
                        if ('سورمه' in nd or 'سرمه' in nd) and 'سرمه' in norm(c['attr']): match = c; break
                        elif 'طلایی' in nd and 'طلایی' in norm(c['attr']): match = c; break
        elif 'اتو' in nd:
            if 'دونالکس' in nd and '815' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'دونالکس' in c['brand'] and '815' in c['name']), None)
            elif 'فلیپس' in nd and '5040' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'فیلیپس' in c['brand'] and '5040' in c['name']), None)
            elif 'ساپر' in nd:
                for c in catalog:
                    if c['sheet'] == 'شستشو و نظافت' and 'ساپر' in c['brand']:
                        if '4025' in nd and '4025' in c['name']: match = c; break
                        elif '4035' in nd and '4035' in c['name']: match = c; break
            elif 'رانکو' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'رانکو' in c['brand']), None)
            elif 'سیماران' in nd and '8002' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'سیماران' in c['brand'] and '8002' in c['name']), None)
            elif 'مونزا' in nd and 'vr3000w' in nd:
                match = next((c for c in catalog if c['sheet'] == 'شستشو و نظافت' and 'مونزا' in c['brand']), None)

    # 5. اویز پخت و پز
    if not match and 'اویز' in nd:
        for c in catalog:
            if c['sheet'] == 'لوازم برقی و پخت و پز' and 'اویز' in c['brand']:
                cn = norm(c['name'])
                ca = norm(c['attr'])
                if 'تابه' in nd and 'تابه' in cn:
                    for sz in ['22', '24', '26', '28', '30', '32']:
                        if sz in nd and (sz in cn or sz in ca): match = c; break
                elif 'قابلمه' in nd and 'قابلمه' in cn:
                    for sz in ['16', '18', '20', '22', '24', '26', '28', '30', '32']:
                        if sz in nd and (sz in cn or sz in ca): match = c; break
                elif 'سرویس' in nd:
                    if 'السا' in nd and 'السا' in cn: match = c; break
                    elif 'ویولا' in nd and 'ویولا' in cn: match = c; break
                    elif 'اناری' in nd and ('اناری' in cn or 'اناری' in ca): match = c; break
                    elif '12' in nd and '12' in cn: match = c; break
                    elif '24' in nd and '24' in cn: match = c; break
                    elif '10' in nd and ('10' in cn or 'سرمه' in ca): match = c; break
                elif 'زودپز' in nd:
                    if 'تکی' in nd and 'تکی' in cn: match = c; break
                    elif 'دوقلو' in nd and 'دوقلو' in cn: match = c; break
                elif 'شیرجوش' in nd and 'شیرجوش' in cn: match = c; break
                elif 'مرغ خوری' in nd and 'مرغ خوری' in cn: match = c; break
            if match: break

    # 6. سایر اقلام پخت و پز و نوشیدنی ساز و زیبایی و سرمایش
    if not match:
        toks = [w for w in nd.split() if len(w) > 2]
        for c in catalog:
            cn = norm(c['name'])
            ca = norm(c['attr'])
            # models
            models = re.findall(r'[a-zA-Z0-9]{3,}', nd)
            if models and all(m in cn or m in ca for m in models):
                match = c; break
            elif toks and all(t in cn or t in ca for t in toks):
                match = c; break

    if match:
        matched_items.append((it, match))
    else:
        unmatched_items.append(it)

print(f'Total: {len(raw_list)}')
print(f'Matched: {len(matched_items)}')
print(f'Unmatched (New items to add): {len(unmatched_items)}')

with open('scratch/unmatched_items.json', 'w', encoding='utf-8') as f:
    json.dump(unmatched_items, f, ensure_ascii=False, indent=2)

print('Saved unmatched items to scratch/unmatched_items.json')
