import os
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
category_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

catalog_items = []
for sheet in category_sheets:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or '📂' in str(name):
            continue
        catalog_items.append({
            'sheet': sheet,
            'row': r,
            'name': str(name).strip(),
            'price_rial': ws.cell(row=r, column=2).value,
            'price_toman': ws.cell(row=r, column=3).value,
            'stock': str(ws.cell(row=r, column=8).value or '').strip(),
            'main_cat': str(ws.cell(row=r, column=9).value or '').strip(),
            'subcat': str(ws.cell(row=r, column=10).value or '').strip(),
            'brand': str(ws.cell(row=r, column=11).value or '').strip(),
            'ptype': str(ws.cell(row=r, column=12).value or 'ساده').strip(),
            'variant_attr': str(ws.cell(row=r, column=13).value or '').strip(),
            'pid': ws.cell(row=r, column=14).value,
            'vid': ws.cell(row=r, column=15).value,
            'link': ws.cell(row=r, column=16).value,
            'sku': ws.cell(row=r, column=17).value,
            'status': str(ws.cell(row=r, column=18).value or '').strip(),
        })

print(f"Total catalog items: {len(catalog_items)}")

# Check approved images in product_images_2026-10-05
img_dir = 'product_images_2026-10-05'
approved_files = []
for root, dirs, files in os.walk(img_dir):
    for f in files:
        if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')) and not f.startswith('.'):
            approved_files.append({
                'filename': f,
                'path': os.path.join(root, f),
                'folder': os.path.basename(root)
            })

print(f"Total image files in {img_dir}: {len(approved_files)}")

# Also check other image directories: new_products, extracted_products
new_prod_imgs = []
for root, dirs, files in os.walk('new_products'):
    for f in files:
        if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')):
            new_prod_imgs.append(os.path.join(root, f))
print(f"Total image files in new_products: {len(new_prod_imgs)}")

ext_prod_imgs = []
for root, dirs, files in os.walk('extracted_products'):
    for f in files:
        if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')):
            ext_prod_imgs.append(os.path.join(root, f))
print(f"Total image files in extracted_products: {len(ext_prod_imgs)}")

# Let's see how each image in approved_files matches catalog items that have status 'در سایت تعریف نشده'
unreg_items = [it for it in catalog_items if it['status'] == 'در سایت تعریف نشده']
print(f"Unregistered items count: {len(unreg_items)}")

