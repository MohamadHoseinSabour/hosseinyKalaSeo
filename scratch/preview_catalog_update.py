import openpyxl
import json
import shutil
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Backup first
shutil.copy2('hosseiny_kala_catalog.xlsx', 'scratch/hosseiny_kala_catalog_backup_before_hold_update.xlsx')
print("Created backup at scratch/hosseiny_kala_catalog_backup_before_hold_update.xlsx")

with open('scratch/user_changes_audit.json', encoding='utf-8') as f:
    audit = json.load(f)

hold_candidates = audit['hold_candidates']

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx')

rows_to_change = []

for h in hold_candidates:
    sheet_name = h['sheet']
    prow = h['row']
    ws = wb[sheet_name]
    
    # Check parent row
    pname = ws.cell(row=prow, column=1).value
    pstatus = ws.cell(row=prow, column=18).value
    rows_to_change.append({
        'sheet': sheet_name,
        'row': prow,
        'name': pname,
        'is_var': False,
        'old_status': pstatus
    })
    
    # Check variations if any
    r = prow + 1
    while r <= ws.max_row:
        c1 = ws.cell(row=r, column=1).value
        if c1 and str(c1).strip().startswith('↳'):
            rows_to_change.append({
                'sheet': sheet_name,
                'row': r,
                'name': c1,
                'is_var': True,
                'old_status': ws.cell(row=r, column=18).value
            })
            r += 1
        else:
            break

print(f"Total rows identified for update to 'فعلا دست نگه دار': {len(rows_to_change)}")
for i, item in enumerate(rows_to_change, 1):
    print(f"{i:2d}. [{item['sheet']} | Row {item['row']}] {item['name']} (Old: {item['old_status']})")
