import openpyxl
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import OrderedDict
import sys

# Import target building logic
from build_target_structure import target, wb, SOURCE_FILE

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

COLOR_HEADER_BG = "1F4E79"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_TOMAN_HEADER_BG = "2E75B6"
COLOR_SUBCAT_BG = "EBF1F5"
COLOR_SUBCAT_TEXT = "375623"
COLOR_VAR_PARENT_BG = "D9E1F2"
COLOR_VARIATION_BG = "F2F2F2"

COLUMNS = [
    ("name", "نام محصول", 48),
    ("price_regular", "قیمت عادی (ریال)", 18),
    ("price_reg_toman", "قیمت عادی (تومان)", 18),
    ("price_base_d", "آخرین قیمت مشخص‌شده (ریال)", 18),
    ("price_diff_pct", "درصد تغییر قیمت", 15),
    ("price_sale", "قیمت حراج (ریال)", 18),
    ("price_sale_toman", "قیمت حراج (تومان)", 18),
    ("in_stock", "وضعیت موجودی", 14),
    ("main_cat", "دسته‌بندی اصلی سایت", 22),
    ("sub_cat", "زیردسته / گروه کالا", 22),
    ("brand", "برند", 16),
    ("type", "نوع محصول", 12),
    ("variant_attr", "مشخصه تنوع / متغیر", 25),
    ("product_id", "شناسه محصول (ID)", 14),
    ("variation_id", "شناسه تنوع", 14),
    ("link", "لینک مستقیم در سایت", 35),
    ("sku", "کد SKU", 14),
    ("status_site", "وضعیت در سایت", 18),
]

for sname, brands_dict in target.items():
    ws = wb[sname]
    # Clear all rows from row 2 downwards
    ws.delete_rows(2, ws.max_row + 10)
    
    # Ensure header row 1 is clean and styled
    ws.row_dimensions[1].height = 28
    for col_idx, (_, h_text, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font = Font(bold=True, color=COLOR_HEADER_TEXT, size=10, name="Calibri")
        is_formula = col_idx in [3, 5, 7]
        h_bg = COLOR_TOMAN_HEADER_BG if is_formula else COLOR_HEADER_BG
        cell.fill = PatternFill(fill_type='solid', fgColor=h_bg)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    
    current_row = 2
    for brand, blocks in brands_dict.items():
        # Section header
        ws.row_dimensions[current_row].height = 22
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(COLUMNS))
        sub_cell = ws.cell(row=current_row, column=1, value=f"  📂 {sname}  ◀  {brand}")
        sub_cell.font = Font(bold=True, color=COLOR_SUBCAT_TEXT, size=10, name="Calibri")
        sub_cell.fill = PatternFill(fill_type='solid', fgColor=COLOR_SUBCAT_BG)
        sub_cell.alignment = Alignment(horizontal='right', vertical='center')
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=current_row, column=c).border = border_thin
        current_row += 1
        
        for b in blocks:
            # Parent row
            p_vals = b['parent']
            ws.row_dimensions[current_row].height = 21
            p_type = p_vals[11] # Col 12: type
            is_var_parent = p_type == 'متغیر' or len(b['variations']) > 0
            
            p_fill = PatternFill(fill_type='solid', fgColor=COLOR_VAR_PARENT_BG) if is_var_parent else PatternFill(fill_type=None)
            
            for col_i in range(1, 19):
                c = ws.cell(row=current_row, column=col_i)
                c.border = border_thin
                if p_fill.fill_type:
                    c.fill = p_fill
                    
                # Values & Formulas
                if col_i == 1:
                    c.value = p_vals[0]
                    c.alignment = Alignment(horizontal='right', vertical='center')
                    c.font = Font(bold=is_var_parent, size=9.5, name="Calibri")
                elif col_i == 2: # B
                    c.value = p_vals[1] if not is_var_parent else None
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 3: # C
                    c.value = f'=IF(OR(B{current_row}="",B{current_row}=0),"",B{current_row}/10)'
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 4: # D
                    c.value = p_vals[3] if not is_var_parent else None
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 5: # E
                    c.value = f'=IF(ISNUMBER(D{current_row}),(B{current_row}-D{current_row})/D{current_row},"")'
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '0.0%'
                elif col_i == 6: # F
                    c.value = p_vals[5] if not is_var_parent else None
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 7: # G
                    c.value = f'=IF(OR(F{current_row}="",F{current_row}=0),"",F{current_row}/10)'
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.number_format = '#,##0'
                elif col_i == 8: # H
                    c.value = p_vals[7]
                    c.alignment = Alignment(horizontal='center', vertical='center')
                elif col_i in [9, 10, 11, 12, 13]: # Category, subcat, brand, type, attr
                    c.value = p_vals[col_i - 1]
                    c.alignment = Alignment(horizontal='center', vertical='center')
                elif col_i in [14, 15]: # ID
                    c.value = p_vals[col_i - 1]
                    c.alignment = Alignment(horizontal='center', vertical='center')
                elif col_i == 16: # link
                    c.value = p_vals[15]
                    c.alignment = Alignment(horizontal='left', vertical='center')
                elif col_i == 17: # SKU
                    c.value = p_vals[16]
                    c.alignment = Alignment(horizontal='center', vertical='center')
                elif col_i == 18: # Status
                    c.value = p_vals[17]
                    c.alignment = Alignment(horizontal='center', vertical='center')
            current_row += 1
            
            # Variations
            for v_vals in b['variations']:
                ws.row_dimensions[current_row].height = 19
                v_fill = PatternFill(fill_type='solid', fgColor=COLOR_VARIATION_BG)
                
                for col_i in range(1, 19):
                    c = ws.cell(row=current_row, column=col_i)
                    c.border = border_thin
                    c.fill = v_fill
                    
                    if col_i == 1:
                        c.value = v_vals[0]
                        c.alignment = Alignment(horizontal='right', vertical='center')
                        c.font = Font(size=9, color="333333", name="Calibri")
                    elif col_i == 2:
                        c.value = v_vals[1]
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i == 3:
                        c.value = f'=IF(OR(B{current_row}="",B{current_row}=0),"",B{current_row}/10)'
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i == 4:
                        c.value = v_vals[3]
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i == 5:
                        c.value = f'=IF(ISNUMBER(D{current_row}),(B{current_row}-D{current_row})/D{current_row},"")'
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '0.0%'
                    elif col_i == 6:
                        c.value = v_vals[5]
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i == 7:
                        c.value = f'=IF(OR(F{current_row}="",F{current_row}=0),"",F{current_row}/10)'
                        c.alignment = Alignment(horizontal='center', vertical='center')
                        c.number_format = '#,##0'
                    elif col_i in range(8, 19):
                        c.value = v_vals[col_i - 1]
                        c.alignment = Alignment(horizontal='center' if col_i != 16 else 'left', vertical='center')
                current_row += 1

print("Saving modified workbook...")
wb.save(SOURCE_FILE)
print("Saved hosseiny_kala_catalog.xlsx successfully!")
