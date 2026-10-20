import json
import openpyxl

with open('data/products.json', 'r', encoding='utf-8') as f:
    products = json.load(f)

for p in products:
    name = str(p.get('name', ''))
    if 'ظرفشویی' in name:
        cats = [c.get('name') if isinstance(c, dict) else str(c) for c in p.get('categories', [])]
        print(f"ID: {p.get('id')}, Name: {name}, Slug: {p.get('slug')}, Cats: {cats}")

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        v_str = str(v) if v else ""
        if 'ظرفشویی' in v_str or '16 نفره' in v_str or '۱۶ نفره' in v_str:
            print(f"Sheet '{sheet}', Row {row}: A={v_str}, B={ws.cell(row=row, column=2).value}, R={ws.cell(row=row, column=18).value}")
