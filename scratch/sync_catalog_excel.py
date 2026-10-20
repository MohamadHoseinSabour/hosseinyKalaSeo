# -*- coding: utf-8 -*-
"""
sync_catalog_excel.py
=====================
بروزرسانی دقیق کاتالوگ اکسل (hosseiny_kala_catalog.xlsx) پس از اتمام ساخت محصولات:
- اصلاح وضعیت موجودی در ستون H به '✗ ناموجود'
- اصلاح وضعیت در سایت در ستون R به 'ثبت‌شده در سایت'
- درج شناسه محصول در ستون N
- درج شناسه تنوع در ستون O (برای تنوع‌ها)
- درج لینک مستقیم محصول در ستون P
- حفظ کامل فرمول‌ها (ستون E درصد تغییر، ستون C و G قیمت تومان) و استایل‌ها
"""

import os
import sys
import json
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

CATALOG_PATH = 'hosseiny_kala_catalog.xlsx'
PROGRESS_FILE = 'scratch/creation_progress.json'
PLAN_FILE = 'scratch/all_61_products.json'

def sync_catalog():
    if not os.path.exists(PROGRESS_FILE):
        print(f"فایل پیشرفت {PROGRESS_FILE} یافت نشد.")
        return
        
    with open(PROGRESS_FILE, 'r', encoding='utf-8') as f:
        progress = json.load(f)
        
    with open(PLAN_FILE, 'r', encoding='utf-8') as f:
        products = json.load(f)
        
    wb = openpyxl.load_workbook(CATALOG_PATH)
    
    updated_rows = 0
    updated_prods = 0
    
    for prod in products:
        name = prod['name']
        sheetname = prod['sheet']
        parent_row = prod['row']
        slug = prod['slug']
        
        if name not in progress:
            continue
            
        pinfo = progress[name]
        pid = pinfo.get('product_id')
        if not pid:
            continue
            
        if sheetname not in wb.sheetnames:
            continue
            
        ws = wb[sheetname]
        
        # ۱. بروزرسانی سطر والد یا محصول ساده
        # ستون H (8): وضعیت موجودی
        ws.cell(row=parent_row, column=8, value='✗ ناموجود')
        # ستون N (14): شناسه محصول
        ws.cell(row=parent_row, column=14, value=pid)
        # ستون P (16): لینک مستقیم
        direct_link = f"https://hosseinykala.ir/product/{slug}/"
        ws.cell(row=parent_row, column=16, value=direct_link)
        # ستون R (18): وضعیت در سایت
        ws.cell(row=parent_row, column=18, value='ثبت‌شده در سایت')
        
        updated_rows += 1
        updated_prods += 1
        
        # ۲. بروزرسانی تنوع‌ها در صورت وجود
        variations = prod.get('variations', [])
        created_vars = pinfo.get('variations', [])
        var_id_map = {cv['name']: cv['variation_id'] for cv in created_vars}
        
        for idx, v in enumerate(variations):
            v_row = v['row']
            v_name = v['name']
            
            # پیدا کردن variation_id
            vid = None
            if v_name in var_id_map:
                vid = var_id_map[v_name]
            elif idx < len(created_vars):
                vid = created_vars[idx].get('variation_id')
                
            ws.cell(row=v_row, column=8, value='✗ ناموجود')
            ws.cell(row=v_row, column=14, value=pid)
            if vid:
                ws.cell(row=v_row, column=15, value=vid)
            ws.cell(row=v_row, column=16, value=direct_link)
            ws.cell(row=v_row, column=18, value='ثبت‌شده در سایت')
            
            updated_rows += 1
            
    # ذخیره فایل اکسل
    wb.save(CATALOG_PATH)
    print(f"==================================================")
    print(f"همگام‌سازی اکسل با موفقیت انجام شد:")
    print(f"  ✓ تعداد محصولات بروزرسانی‌شده: {updated_prods}")
    print(f"  ✓ تعداد کل سطرهای اصلاح‌شده در کاتالوگ: {updated_rows}")
    print(f"  ✓ فایل ذخیره شد: {CATALOG_PATH}")
    print(f"==================================================")

if __name__ == '__main__':
    sync_catalog()
