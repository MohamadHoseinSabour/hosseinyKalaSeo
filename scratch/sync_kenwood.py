import openpyxl, sys, requests, time
from urllib3.util.ssl_ import create_urllib3_context
from requests.adapters import HTTPAdapter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

class TLSAdapter(HTTPAdapter):
    def init_poolmanager(self, *a, **k):
        ctx = create_urllib3_context()
        ctx.set_ciphers('DEFAULT@SECLEVEL=1')
        k['ssl_context'] = ctx
        return super().init_poolmanager(*a, **k)

s = requests.Session()
s.mount('https://', TLSAdapter())
auth = ('agent', 'Mlwg pHHX NQ64 drrz 7ca2 IVto')
headers = {'User-Agent': 'Mozilla/5.0', 'Content-Type': 'application/json'}

kenwood_short = """<ul>
<li><strong>مخزن جادار ۱ لیتری:</strong> پارچ شفاف مدرج با ظرفیت ۱۰۰۰ میلی‌لیتر و دستگیره ارگونومیک</li>
<li><strong>چرخش دوطرفه هوشمند:</strong> سیستم چرخش معکوس سری مخروطی برای استخراج حداکثری آب‌میوه</li>
<li><strong>صافی استیل ضدزنگ:</strong> شبکه فیلتر تمام استیل با دوام بالا برای جداسازی کامل هسته و پالپ</li>
<li><strong>مکانیزم فشاری خودکار:</strong> شروع به کار آسان تنها با فشردن میوه روی سری مخروطی</li>
<li><strong>شستشوی آسان:</strong> قطعات جداشونده با قابلیت شستشو در ماشین ظرفشویی</li>
<li><strong>تضمین اصالت:</strong> ضمانت اصالت ۱۰۰٪ کالا از فروشگاه حسینی کالا اهواز</li>
</ul>"""

kenwood_desc = """<h2>بررسی تخصصی و راهنمای خرید آب مرکبات‌گیری برقی کنوود مدل JE290</h2>
<p>نوشیدن یک لیوان آب پرتقال یا لیمو شیرین طبیعی و خنک، بهترین شیوه برای جذب ویتامین C و حفظ شادابی روزانه است. <strong>آب مرکبات گیری برقی کنوود مدل JE290</strong> از برند صاحب‌نام <strong>کنوود (Kenwood)</strong>، دستگاهی سبک، کارآمد و بادوام برای آبگیری سریع انواع مرکبات ریز و درشت در خانه است.</p>

<h3>ویژگی‌های کلیدی آب مرکبات‌گیری کنوود ۲۹۰</h3>
<ul>
<li><strong>موتور با چرخش دوطرفه:</strong> سری مخروطی دستگاه با هر بار فشردن میوه، جهت چرخش خود را تغییر می‌دهد. این چرخش دوطرفه (Two-way rotation) باعث می‌شود تا آخرین قطره آب‌میوه از بافت پرتقال، گریپ‌فروت یا لیمو ترش جدا شود.</li>
<li><strong>پارچ مدرج ۱۰۰۰ میلی‌لیتری:</strong> مخزن شفاف مدرج دستگاه به شما امکان می‌دهد حجم آبمیوه طبیعی تهیه شده را مشاهده کرده و به کمک لبه آبریز مخصوص، مستقیماً داخل لیوان سرو نمایید.</li>
<li><strong>فیلتر تمام استیل ضد زنگ:</strong> برخلاف مدل‌های ارزان‌قیمت پلاستیکی، صافی این دستگاه از استیل SUS ساخته شده که در برابر اسید میوه مقاوم بوده و طعم نوشیدنی را تغییر نمی‌دهد.</li>
<li><strong>سیستم ضد لغزش و جمع‌آوری کابل:</strong> پایه‌های مکنده لاستیکی از لغزش دستگاه جلوگیری کرده و سیم برق در پایه زیرین جمع‌آوری می‌شود.</li>
</ul>

<h2>جدول مشخصات فنی آب مرکبات‌گیری کنوود مدل JE290</h2>
<table style="width: 100%; border-collapse: collapse; margin-top: 15px;" border="1">
<tbody>
<tr><td style="padding: 8px; width: 35%;"><strong>برند</strong></td><td style="padding: 8px;">کنوود (Kenwood)</td></tr>
<tr><td style="padding: 8px;"><strong>مدل</strong></td><td style="padding: 8px;">JE290</td></tr>
<tr><td style="padding: 8px;"><strong>نوع دستگاه</strong></td><td style="padding: 8px;">آب مرکبات‌گیری برقی فشاری</td></tr>
<tr><td style="padding: 8px;"><strong>توان مصرفی</strong></td><td style="padding: 8px;">۶۰ وات بهینه</td></tr>
<tr><td style="padding: 8px;"><strong>ظرفیت مخزن</strong></td><td style="padding: 8px;">۱ لیتر (۱۰۰۰ میلی‌لیتر) مدرج</td></tr>
<tr><td style="padding: 8px;"><strong>جنس فیلتر</strong></td><td style="padding: 8px;">استیل ضدزنگ</td></tr>
<tr><td style="padding: 8px;"><strong>سیستم چرخش</strong></td><td style="padding: 8px;">دو جهته خودکار (راست‌گرد و چپ‌گرد)</td></tr>
<tr><td style="padding: 8px;"><strong>درپوش ضد گردوغبار</strong></td><td style="padding: 8px;">دارد</td></tr>
<tr><td style="padding: 8px;"><strong>گارانتی و اصالت</strong></td><td style="padding: 8px;">ضمانت اصالت و سلامت فیزیکی از حسینی کالا</td></tr>
</tbody>
</table>

<p>این محصول کاربردی را با بهترین قیمت و ضمانت اصالت فیزیکی از <a href="https://hosseinykala.ir/">فروشگاه لوازم خانگی حسینی</a> اهواز به صورت آنلاین یا حضوری خریداری فرمایید.</p>"""

payload = {
    'name': 'آب مرکبات گیری برقی کنوود مدل JE290',
    'slug': 'kenwood-je290-electric-citrus-juicer',
    'type': 'simple',
    'status': 'publish',
    'catalog_visibility': 'visible',
    'description': kenwood_desc,
    'short_description': kenwood_short,
    'categories': [{'id': 240}, {'id': 120}],
    'images': [{'id': 4247}],
    'manage_stock': False,
    'stock_status': 'outofstock',
    'attributes': [
        {'name': 'برند', 'visible': True, 'variation': False, 'options': ['کنوود (Kenwood)']},
        {'name': 'مدل', 'visible': True, 'variation': False, 'options': ['JE290']},
        {'name': 'توان مصرفی', 'visible': True, 'variation': False, 'options': ['۶۰ وات']},
        {'name': 'ظرفیت مخزن', 'visible': True, 'variation': False, 'options': ['۱ لیتر']},
        {'name': 'جنس صافی', 'visible': True, 'variation': False, 'options': ['استیل ضدزنگ']},
        {'name': 'عملکرد چرخش', 'visible': True, 'variation': False, 'options': ['چرخش دوطرفه خودکار']}
    ]
}

print("Creating Kenwood JE290...")
r = s.post('https://hosseinykala.ir/wp-json/wc/v3/products', auth=auth, headers=headers, json=payload, timeout=60)
if r.status_code == 201:
    res = r.json()
    pid = res['id']
    slug = res['slug']
    print(f"Created successfully: ID {pid} | Slug: {slug}")
    wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx')
    ws = wb['نوشیدنی ساز']
    ws.cell(68, 8).value = '✗ ناموجود'
    ws.cell(68, 15).value = pid
    ws.cell(68, 16).value = f"https://hosseinykala.ir/product/{slug}/"
    ws.cell(68, 18).value = 'ثبت‌شده در سایت'
    wb.save('hosseiny_kala_catalog.xlsx')
    print("Row 68 (آب مرکبات گیری برقی کنوود) synced in catalog sheet!")
else:
    print(f"Failed: {r.status_code} - {r.text[:300]}")
