import openpyxl

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=False)

for sname in wb.sheetnames:
    if sname in ['📊 خلاصه و داشبورد', '📖 راهنمای کاربری']:
        continue
    ws = wb[sname]
    blocks = []
    current_block = None
    current_section = None
    
    for r in range(2, ws.max_row + 1):
        vA = ws.cell(row=r, column=1).value
        if not vA:
            continue
        vA_str = str(vA).strip()
        if '📂' in vA_str:
            current_section = vA_str
            continue
        
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 19)]
        
        if vA_str.startswith('↳'):
            if current_block:
                current_block['variations'].append(row_vals)
        else:
            current_block = {
                'section': current_section,
                'parent': row_vals,
                'variations': []
            }
            blocks.append(current_block)
            
    print(f"Sheet '{sname}': extracted {len(blocks)} product blocks.")
