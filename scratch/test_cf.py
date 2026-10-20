import openpyxl, sys
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import FormulaRule

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Test'
ws.append(['نام', 'وضعیت'])
ws.append(['محصول ۱', 'ثبت‌شده در سایت'])
ws.append(['محصول ۲', 'فعلا دست نگه دار'])
ws.append(['محصول ۳', 'در سایت تعریف نشده'])

rule_orange = FormulaRule(formula=['$B2="فعلا دست نگه دار"'], fill=PatternFill(fill_type='solid', fgColor='FFFCE4D6'))
rule_yellow = FormulaRule(formula=['$B2="در سایت تعریف نشده"'], fill=PatternFill(fill_type='solid', fgColor='FFFFF2CC'))
ws.conditional_formatting.add('A2:B4', rule_orange)
ws.conditional_formatting.add('A2:B4', rule_yellow)

wb.save('scratch/test_cf.xlsx')
print('Test workbook saved successfully!')
