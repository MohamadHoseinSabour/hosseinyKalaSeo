import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)
sheets = ['یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز', 'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی']

items = []
for sname in sheets:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or str(name).strip().startswith(('📂', '•')):
            continue
        ptype = str(ws.cell(r, 12).value or '').strip()
        if ptype == 'متغیر':
            continue
        b_val = ws.cell(r, 2).value
        f_val = ws.cell(r, 6).value
        pid = ws.cell(r, 14).value
        vid = ws.cell(r, 15).value
        if b_val and isinstance(b_val, (int, float)) and b_val > 0:
            new_b = int(round(b_val * 1.30))
            new_f = int(round(f_val * 1.30)) if (f_val and isinstance(f_val, (int, float)) and f_val > 0) else None
            items.append({
                'sheet': sname, 'row': r, 'name': str(name).strip(), 'ptype': ptype,
                'pid': int(pid), 'vid': int(vid) if vid else None,
                'old_b': int(b_val), 'new_b': new_b,
                'old_f': int(f_val) if f_val else None, 'new_f': new_f
            })

print(f'Total items to increase by 30%: {len(items)}')
print('\nSample items:')
for it in items[:10]:
    f_str = f" | Sale: {it['old_f']:,} -> {it['new_f']:,}" if it['old_f'] else ""
    target_id = it['vid'] or it['pid']
    print(f"{it['sheet']} R{it['row']} [{it['ptype']} #{target_id}]: {it['name'][:30]} | Reg: {it['old_b']:,} -> {it['new_b']:,}{f_str}")
