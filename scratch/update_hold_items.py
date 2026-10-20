import os, sys, shutil, datetime, openpyxl
from openpyxl.styles import PatternFill

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')

now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = os.path.join(PROJECT_DIR, f'hosseiny_kala_catalog.backup_{now_str}.xlsx')
shutil.copy2(CATALOG_PATH, backup_path)
print(f"Backup created: {os.path.basename(backup_path)}")

wb = openpyxl.load_workbook(CATALOG_PATH, data_only=False)
hold_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

targets = [
    ('لوازم برقی و پخت و پز', 19, '↳ سایز (سانتی‌متر): سایز ۲۸'),
    ('لوازم برقی و پخت و پز', 66, 'هود شومینه ای آلتون مدل 900'),
    ('نوشیدنی ساز', 31, 'کتری برقی تفال 2.2 لیتر'),
    ('سرمایش و گرمایش', 30, 'بخاری گازی پلار طرح سنگ بدون دودکش'),
    ('سرمایش و گرمایش', 36, 'آبگرمکن دیواری اتوماتیک ایران شرق'),
]

for sname, row_idx, expected_name in targets:
    ws = wb[sname]
    actual_name = ws.cell(row_idx, 1).value
    print(f"Sheet: {sname} | Row {row_idx} | Expected: '{expected_name}' | Actual: '{actual_name}'")
    assert expected_name in str(actual_name), f"Mismatch in {sname} row {row_idx}"
    
    # Update status in column 18 (R)
    ws.cell(row_idx, 18).value = 'فعلا دست نگه دار'
    
    # Fill cells A..R with FCE4D6
    for c in range(1, 19):
        ws.cell(row_idx, c).fill = hold_fill
    print(f"  -> Successfully updated to 'فعلا دست نگه دار'")

wb.save(CATALOG_PATH)
print("Catalog saved successfully.")
