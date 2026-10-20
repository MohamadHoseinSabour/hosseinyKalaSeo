import openpyxl, os, sys
from openpyxl.styles import PatternFill

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')

def refine_catalog_dynamic_colors():
    print("=== REFINING CATALOG TO 100% PURE DYNAMIC COLORING ===")
    
    wb = openpyxl.load_workbook(CATALOG_PATH)
    no_fill = PatternFill(fill_type=None)

    total_hold_cleaned = 0
    total_reg_verified = 0

    for sheetname in wb.sheetnames:
        if sheetname in ['📖 راهنمای کاربری', '📊 خلاصه و داشبورد']:
            continue
            
        ws = wb[sheetname]
        
        # Verify conditional formatting range
        cf_ranges = [str(cf.sqref) for cf in ws.conditional_formatting if 'A2:R' in str(cf.sqref)]
        print(f"Sheet [{sheetname}]: max_row={ws.max_row} | CF range={cf_ranges}")

        for r in range(2, ws.max_row + 1):
            name = str(ws.cell(r, 1).value or '')
            if name.startswith('  📂'):
                continue
                
            st = ws.cell(r, 18).value
            
            # If status is 'فعلا دست نگه دار' and has manual orange fill:
            # Clear manual fill so it relies purely on Conditional Formatting
            # This ensures that when user changes status to 'ثبت‌شده در سایت', it immediately turns white!
            if st == 'فعلا دست نگه دار':
                f1 = ws.cell(r, 1).fill
                fg1 = f1.fgColor.rgb if f1 and f1.fgColor else 'None'
                if fg1 in ['FFFCE4D6', 'FCE4D6', '00FCE4D6']:
                    for c in range(1, 19):
                        ws.cell(r, c).fill = no_fill
                    total_hold_cleaned += 1
                    
            elif st == 'ثبت‌شده در سایت':
                # Ensure no yellow fill remains on any cell
                for c in range(1, 19):
                    f = ws.cell(r, c).fill
                    fg = f.fgColor.rgb if f and f.fgColor else 'None'
                    if fg in ['FFFFF2CC', 'FFF2CC']:
                        ws.cell(r, c).fill = no_fill
                total_reg_verified += 1

    wb.save(CATALOG_PATH)
    print(f"\n✓ Verification Complete:")
    print(f"  - Verified {total_reg_verified} 'ثبت‌شده در سایت' rows (all yellow removed).")
    print(f"  - Cleaned direct fill on {total_hold_cleaned} 'فعلا دست نگه دار' rows (now purely dynamic).")
    print(f"✓ Catalog saved successfully to {CATALOG_PATH}")

if __name__ == '__main__':
    refine_catalog_dynamic_colors()
