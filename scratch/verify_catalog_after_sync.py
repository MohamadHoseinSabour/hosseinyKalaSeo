import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=False)
category_sheets = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

unregistered_left = 0
registered_count = 0
on_hold_count = 0
stock_out_count = 0

for sheet in category_sheets:
    if sheet not in wb.sheetnames:
        continue
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or '📂' in str(name):
            continue
        status = str(ws.cell(row=r, column=18).value or '').strip()
        stock = str(ws.cell(row=r, column=8).value or '').strip()
        pid = ws.cell(row=r, column=14).value
        link = ws.cell(row=r, column=16).value
        
        if status == 'در سایت تعریف نشده':
            unregistered_left += 1
            print(f"  ⚠️ هنوز در سایت تعریف نشده: [{sheet}] سطر {r}: {name}")
        elif status == 'ثبت‌شده در سایت':
            registered_count += 1
        elif status == 'فعلا دست نگه دار':
            on_hold_count += 1
            
        if 'ناموجود' in stock:
            stock_out_count += 1

print("\n--- نتیجه بررسی وضعیت کاتالوگ اکسل ---")
print(f"تعداد اقلام ثبت‌شده در سایت: {registered_count}")
print(f"تعداد اقلام متوقف‌شده (فعلا دست نگه دار): {on_hold_count}")
print(f"تعداد اقلام باقیمانده در وضعیت «در سایت تعریف نشده»: {unregistered_left}")
print(f"تعداد کل اقلام با وضعیت موجودی ناموجود: {stock_out_count}")

# Check sample formulas
ws_y = wb['یخچال و فریزر']
print("\nبررسی فرمول‌ها در سطر 2 یخچال و فریزر:")
print(f"  فرمول درصد تغییر (ستون E): {ws_y.cell(row=2, column=5).value}")
print(f"  فرمول قیمت تومان (ستون C): {ws_y.cell(row=2, column=3).value}")
