import os
import sys
import shutil
import datetime
import json
import re
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.cell.cell import MergedCell

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
DUMP_PATH = os.path.join(PROJECT_DIR, 'scratch', 'wc_live_dump.json')

with open(DUMP_PATH, 'r', encoding='utf-8') as f:
    dump = json.load(f)

site_vars_by_parent = dump.get('variations', {})
site_products = {p['id']: p for p in dump.get('products', [])}

def clean_opt(text):
    if not text:
        return ''
    text = str(text).replace('↳', '').strip()
    text = re.sub(r'^(رنگ|سایز|ظرفیت|مدل|طول|مشخصه|طرح)\s*(\([^)]*\))?\s*:\s*', '', text).strip()
    return text.replace('ي', 'ی').replace('ك', 'ک').lower()

def safe_set_cell(ws, r, c, val, num_format=None, align=None):
    cell = ws.cell(r, c)
    if isinstance(cell, MergedCell):
        return
    cell.value = val
    if num_format:
        cell.number_format = num_format
    if align:
        cell.alignment = align

def update_catalog():
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(PROJECT_DIR, f'hosseiny_kala_catalog.backup_{now_str}.xlsx')
    shutil.copy2(CATALOG_PATH, backup_path)
    print(f"۱. نسخه پشتیبان تهیه شد: {os.path.basename(backup_path)}")

    wb = openpyxl.load_workbook(CATALOG_PATH, data_only=False)
    sheets = [
        'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
        'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
    ]

    fixed_simple_pids = 0
    matched_vids = 0
    applied_discounts = 0
    align_center = Alignment(horizontal='center', vertical='center')

    for sname in sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]

        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name:
                continue
            name_str = str(name).strip()
            if name_str.startswith('📂') or name_str.startswith('•'):
                continue

            ptype = str(ws.cell(r, 12).value or '').strip()
            reg_val = ws.cell(r, 2).value
            sale_val = ws.cell(r, 6).value
            last_val = ws.cell(r, 4).value
            pid_val = ws.cell(r, 14).value
            vid_val = ws.cell(r, 15).value

            # Format numbers standard safely
            for col_idx, fmt in [(2, '#,##0'), (3, '#,##0'), (4, '#,##0'), (5, '+0.0%;-0.0%;"0.0%"'), (6, '#,##0'), (7, '#,##0')]:
                c_obj = ws.cell(r, col_idx)
                if not isinstance(c_obj, MergedCell):
                    c_obj.number_format = fmt

            # 1. Parents must never have prices
            if ptype == 'متغیر':
                safe_set_cell(ws, r, 2, None)
                safe_set_cell(ws, r, 4, None)
                safe_set_cell(ws, r, 6, None)
                continue

            # 2. Fix simple products having ID in VID instead of PID
            if ptype == 'ساده' and (pid_val is None or str(pid_val).strip() == '') and vid_val is not None:
                safe_set_cell(ws, r, 14, int(vid_val), align=align_center)
                safe_set_cell(ws, r, 15, None)
                pid_val = int(vid_val)
                vid_val = None
                fixed_simple_pids += 1

            # 3. Match variations missing VID from WooCommerce dump
            if ptype in ['تنوع', 'متغیر-فرزند'] and pid_val and (vid_val is None or str(vid_val).strip() == ''):
                pid_str = str(int(pid_val))
                if pid_str in site_vars_by_parent:
                    row_opt = clean_opt(name_str)
                    found_vid = None
                    for sv in site_vars_by_parent[pid_str]:
                        for attr in sv.get('attributes', []):
                            attr_opt = clean_opt(attr.get('option', ''))
                            if row_opt == attr_opt or row_opt in attr_opt or attr_opt in row_opt:
                                found_vid = sv['id']
                                break
                        if found_vid:
                            break
                    if found_vid:
                        safe_set_cell(ws, r, 15, found_vid, align=align_center)
                        vid_val = found_vid
                        matched_vids += 1

            # 4. Apply 5% discount on products without discount
            if isinstance(reg_val, (int, float)) and reg_val > 0:
                # If sale price is empty or None or 0
                if sale_val is None or sale_val == '' or sale_val == 0:
                    sale_rial = int(round(reg_val * 0.95))
                    safe_set_cell(ws, r, 6, sale_rial, num_format='#,##0', align=align_center)
                    safe_set_cell(ws, r, 7, f'=IF(OR(ISBLANK(F{r}),F{r}=""),"",IF($L{r}="متغیر","",F{r}/10))', num_format='#,##0', align=align_center)

                    # If Column D (last price) is empty, populate it with regular price
                    if last_val is None or last_val == '' or last_val == 0:
                        safe_set_cell(ws, r, 4, reg_val, num_format='#,##0', align=align_center)

                    applied_discounts += 1

            # Ensure standard formulas for Columns C, E, G
            safe_set_cell(ws, r, 2, reg_val, num_format='#,##0')
            safe_set_cell(ws, r, 3, f'=IF(OR(ISBLANK(B{r}),B{r}=""),"",IF($L{r}="متغیر","⚠️ والد متغیر نباید قیمت داشته باشد",B{r}/10))', num_format='#,##0')
            safe_set_cell(ws, r, 5, f'=IF(OR(ISBLANK(B{r}),B{r}="",ISBLANK(D{r}),D{r}="",D{r}<=0,$L{r}="متغیر"),"",(B{r}-D{r})/D{r})', num_format='+0.0%;-0.0%;"0.0%"')
            safe_set_cell(ws, r, 7, f'=IF(OR(ISBLANK(F{r}),F{r}=""),"",IF($L{r}="متغیر","",F{r}/10))', num_format='#,##0')

    wb.save(CATALOG_PATH)
    print(f"\n۲. فایل کاتالوگ با موفقیت ذخیره شد: {os.path.basename(CATALOG_PATH)}")
    print(f"   ✓ تعداد شناسه‌های جابجا شده برای کالاهای ساده: {fixed_simple_pids}")
    print(f"   ✓ تعداد تنوع‌های شناسایی‌شده و ثبت VID: {matched_vids}")
    print(f"   ✓ تعداد کالاهای اعمال تخفیف ۵ درصدی: {applied_discounts}")

if __name__ == '__main__':
    update_catalog()
