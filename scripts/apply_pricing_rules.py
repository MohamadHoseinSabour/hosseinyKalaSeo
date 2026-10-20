"""
apply_pricing_rules.py
======================
اسکریپت اعمال قوانین قیمت‌گذاری روی کاتالوگ اکسل حسینی کالا:
۱. افزایش ۱۰ درصدی قیمت تمام محصولات غیر از برند کلور (+10%)
۲. حفظ ۱۰۰٪ قیمت اصلی محصولات کلور و درج قیمت اصلی با ۱۰٪ تخفیف به عنوان قیمت حراج (-10%)
۳. دست نزدن به وضعیت موجودی‌ها (حفظ کامل ستون موجودی)
۴. حفظ فرمول‌های ستون تومان و ساختار RTL و استایل‌ها
"""

import os
import sys
import shutil
import datetime
import openpyxl
from openpyxl.styles import Alignment

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
CATALOG_LATEST_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog_latest.xlsx')

def run_pricing_update():
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(PROJECT_DIR, f'hosseiny_kala_catalog.backup_{now_str}.xlsx')
    
    print(f"1. Creating backup: {os.path.basename(backup_path)} ...")
    shutil.copy2(CATALOG_PATH, backup_path)
    print("   Backup successfully created.")

    print(f"2. Loading workbook: {os.path.basename(CATALOG_PATH)} ...")
    wb = openpyxl.load_workbook(CATALOG_PATH, data_only=False)
    
    sheets_to_process = [
        'یخچال و فریزر',
        'شستشو و نظافت',
        'لوازم برقی و پخت و پز',
        'نوشیدنی ساز',
        'سرمایش و گرمایش',
        'تصفیه کننده آب',
        'صوتی و تصویری',
        'زیبایی و سلامت فردی'
    ]

    total_clover_updated = 0
    total_non_clover_updated = 0
    total_skipped_unpriced = 0

    for sname in sheets_to_process:
        if sname not in wb.sheetnames:
            print(f"   [WARN] Sheet '{sname}' not found, skipping.")
            continue
        
        ws = wb[sname]
        current_parent_name = ""
        current_parent_brand = ""
        current_parent_is_clover = False

        sheet_clover = 0
        sheet_non_clover = 0

        for r in range(2, ws.max_row + 1):
            name_cell = ws.cell(r, 1)
            name = name_cell.value
            if not name:
                continue
            
            name_str = str(name).strip()
            if name_str.startswith('📂') or name_str.startswith('•'):
                continue
            
            row_type = str(ws.cell(r, 10).value or '').strip()
            brand_val = str(ws.cell(r, 9).value or '').strip()
            subcat = str(ws.cell(r, 8).value or '').strip()
            
            is_variation = name_str.startswith('↳') or row_type == 'تنوع'

            if not is_variation:
                current_parent_name = name_str
                current_parent_brand = brand_val
                current_parent_is_clover = (
                    'کلور' in current_parent_name or 'clover' in current_parent_name.lower() or
                    'کلور' in current_parent_brand or 'clover' in current_parent_brand.lower() or
                    'کلور' in subcat or 'clover' in subcat.lower()
                )
                is_clover = current_parent_is_clover
            else:
                is_clover = current_parent_is_clover

            cell_p_reg = ws.cell(r, 2)
            cell_p_sale = ws.cell(r, 4)

            p_reg = cell_p_reg.value
            p_sale = cell_p_sale.value

            # If no regular price, skip
            if p_reg is None or not isinstance(p_reg, (int, float)):
                total_skipped_unpriced += 1
                continue

            if is_clover:
                # Clover: Regular price stays untouched
                # Sale price becomes 90% of regular price (10% discount)
                new_sale_price = int(round(p_reg * 0.90))
                cell_p_sale.value = new_sale_price
                cell_p_sale.number_format = '#,##0'
                cell_p_sale.alignment = Alignment(horizontal='center', vertical='center')
                total_clover_updated += 1
                sheet_clover += 1
            else:
                # Non-clover: Regular price increases by 10%
                new_reg_price = int(round(p_reg * 1.10))
                cell_p_reg.value = new_reg_price
                cell_p_reg.number_format = '#,##0'
                cell_p_reg.alignment = Alignment(horizontal='center', vertical='center')
                
                # If product already had a sale price, increase it by 10% too
                if p_sale is not None and isinstance(p_sale, (int, float)):
                    new_sale_price = int(round(p_sale * 1.10))
                    cell_p_sale.value = new_sale_price
                    cell_p_sale.number_format = '#,##0'
                    cell_p_sale.alignment = Alignment(horizontal='center', vertical='center')

                total_non_clover_updated += 1
                sheet_non_clover += 1

        print(f"   [{sname}] Updated: {sheet_clover} Clover items (discounted), {sheet_non_clover} Non-Clover items (+10%)")

    print(f"\n3. Saving updated catalog to: {os.path.basename(CATALOG_PATH)} ...")
    wb.save(CATALOG_PATH)
    print("   Saved successfully.")

    print(f"4. Syncing to latest catalog: {os.path.basename(CATALOG_LATEST_PATH)} ...")
    shutil.copy2(CATALOG_PATH, CATALOG_LATEST_PATH)
    print("   Synced successfully.")

    print("\nSummary of Updates:")
    print(f" - Total Clover items with 10% discount: {total_clover_updated}")
    print(f" - Total Non-Clover items increased by 10%: {total_non_clover_updated}")
    print(f" - Total unpriced/parent items kept as-is: {total_skipped_unpriced}")
    print(f" - Backup saved at: {backup_path}")

if __name__ == '__main__':
    run_pricing_update()
