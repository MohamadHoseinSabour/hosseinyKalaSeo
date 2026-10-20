#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
apply_proper_rial_and_toman_prices.py
=====================================
اصلاح بنیادین و صحیح قیمت‌های ریال و تومان در کل کاتالوگ حسینی کالا:
۱. مبالغ پایه موجود در کاتالوگ (مثلاً ۲۴۰,۰۰۰,۰۰۰ برای یخچال الکترواستیل فیوچر)
   همگی بر حسب «تومان» هستند (۲۴۰ میلیون تومان).
۲. در نتیجه:
   - ستون B (قیمت عادی ریال) = مبلغ تومان ضرب در ۱۰ (یک صفر بیشتر): ۲,۴۰۰,۰۰۰,۰۰۰ ریال
   - ستون C (قیمت عادی تومان) = فرمول B/10 با فرمت سه‌رقم‌سه‌رقم (#,##0): ۲۴۰,۰۰۰,۰۰۰ تومان
   - ستون D (آخرین قیمت مشخص‌شده ریال) = دقیقاً برابر با ستون B: ۲,۴۰۰,۰۰۰,۰۰۰ ریال
   - ستون E (درصد تغییر) = فرمول (B-D)/D با فرمت +0.0%;-0.0%;"0.0%": 0.0%
   - ستون F (قیمت حراج ریال) = در صورت وجود، ضرب در ۱۰ به ریال
   - ستون G (قیمت حراج تومان) = فرمول F/10 با فرمت سه‌رقم‌سه‌رقم (#,##0)
۳. محصولات والد متغیر: فاقد هرگونه قیمت (B, D, F کاملاً خالی).
۴. اعمال همزمان روی:
   - hosseiny_kala_catalog.xlsx
   - hosseiny_kala_catalog.xlsm
   - hosseiny_kala_catalog_latest.xlsx
"""

import os
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.backup_20261005_032828.xlsx')

TARGET_FILES = [
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx'), False),
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsm'), True),
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog_latest.xlsx'), False),
]

SHEETS = [
    'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
    'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
]

def load_toman_prices_from_backup():
    if not os.path.exists(BACKUP_PATH):
        raise FileNotFoundError(f"Backup file not found: {BACKUP_PATH}")
    wb = openpyxl.load_workbook(BACKUP_PATH, data_only=True)
    prices_map = {} # (sheet_name, row) -> (reg_toman, sale_toman, stock)
    for sname in SHEETS:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        for r in range(2, ws.max_row + 1):
            reg = ws.cell(r, 2).value
            sale = ws.cell(r, 6).value
            stock = ws.cell(r, 8).value
            prices_map[(sname, r)] = (reg, sale, stock)
    return prices_map

def apply_to_file(filepath, is_xlsm, prices_map):
    if not os.path.exists(filepath):
        print(f"Skipping missing: {filepath}")
        return 0

    print(f"\n📂 در حال اصلاح قیمت‌ها در: {os.path.basename(filepath)}")
    wb = openpyxl.load_workbook(filepath, keep_vba=is_xlsm)
    updated_items = 0

    for sname in SHEETS:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        sheet_count = 0

        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name or str(name).strip().startswith('📂') or str(name).strip().startswith('•'):
                continue

            ptype = str(ws.cell(r, 12).value or '').strip()

            # تنظیم فرمت‌های ۳ رقم ۳ رقم (#,##0)
            ws.cell(r, 2).number_format = '#,##0'
            ws.cell(r, 3).number_format = '#,##0'
            ws.cell(r, 4).number_format = '#,##0'
            ws.cell(r, 5).number_format = '+0.0%;-0.0%;"0.0%"'
            ws.cell(r, 6).number_format = '#,##0'
            ws.cell(r, 7).number_format = '#,##0'

            # فرمول‌های تومان و درصد تغییر
            ws.cell(r, 3).value = f'=IF(OR(ISBLANK(B{r}),B{r}=""),"",IF($L{r}="متغیر","⚠️ والد متغیر نباید قیمت داشته باشد",B{r}/10))'
            ws.cell(r, 5).value = f'=IF(OR(ISBLANK(B{r}),B{r}="",ISBLANK(D{r}),D{r}="",D{r}<=0,$L{r}="متغیر"),"",(B{r}-D{r})/D{r})'
            ws.cell(r, 7).value = f'=IF(OR(ISBLANK(F{r}),F{r}=""),"",IF($L{r}="متغیر","",F{r}/10))'

            # والد متغیر بدون قیمت
            if ptype == 'متغیر':
                ws.cell(r, 2).value = None
                ws.cell(r, 4).value = None
                ws.cell(r, 6).value = None
                continue

            # دریافت قیمت پایه (تومان) از بک‌آپ
            pair = prices_map.get((sname, r), (None, None, None))
            base_toman = pair[0]
            base_sale_toman = pair[1]

            if base_toman is not None and isinstance(base_toman, (int, float)) and base_toman > 0:
                # تبدیل تومان به ریال (یک صفر بیشتر: ضرب در ۱۰)
                rial_val = int(round(base_toman * 10))
                ws.cell(r, 2).value = rial_val # قیمت عادی (ریال)
                ws.cell(r, 4).value = rial_val # آخرین قیمت مشخص‌شده (ریال)
                sheet_count += 1
            else:
                ws.cell(r, 2).value = None
                ws.cell(r, 4).value = None

            if base_sale_toman is not None and isinstance(base_sale_toman, (int, float)) and base_sale_toman > 0:
                sale_rial_val = int(round(base_sale_toman * 10))
                ws.cell(r, 6).value = sale_rial_val
            else:
                ws.cell(r, 6).value = None

        updated_items += sheet_count
        print(f"  - شیت {sname}: {sheet_count} محصول به ریال صحیح (ضرب در ۱۰) تبدیل شد.")

    wb.save(filepath)
    print(f"✓ فایل {os.path.basename(filepath)} با موفقیت ذخیره شد. (تعداد محصولات: {updated_items})")
    return updated_items

def main():
    print("=" * 70)
    print(" تصحیح استاندارد قیمت‌های ریال (اضافه شدن صفر) و تفکیک ۳ رقم تومان")
    print("=" * 70)

    prices_map = load_toman_prices_from_backup()
    print(f"✓ قیمت‌های تومان از فایل مرجع استخراج شد.")

    for fpath, is_xlsm in TARGET_FILES:
        apply_to_file(fpath, is_xlsm, prices_map)

    print("\n" + "=" * 70)
    print("✅ تمامی فایل‌های اکسل کاتالوگ با موفقیت اصلاح شدند:")
    print("   - ستون B (قیمت عادی ریال) = مبلغ تومان ضرب در ۱۰ (یک صفر بیشتر)")
    print("   - ستون C (قیمت عادی تومان) = مبالغ دقیق تومان با جداکننده ۳ رقم ۳ رقم")
    print("   - ستون D (آخرین قیمت ریال) = معادل ستون B (درصد تغییر صفر)")
    print("=" * 70)

if __name__ == '__main__':
    main()
