import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog.xlsx', data_only=True)

all_missing = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    current_parent = None
    sheet_missing = []
    
    # Check headers
    headers = [ws.cell(1, c).value for c in range(1, 25)]
    name_col = 1
    status_col = 18
    link_col = 16
    
    # Try finding exact column index for 'وضعیت در سایت'
    for c_idx, h in enumerate(headers, 1):
        if h and 'وضعیت در سایت' in str(h):
            status_col = c_idx
        elif h and 'لینک مستقیم' in str(h):
            link_col = c_idx
            
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        status_site = ws.cell(r, status_col).value
        link = ws.cell(r, link_col).value
        brand = ws.cell(r, 11).value
        cat = ws.cell(r, 9).value
        subcat = ws.cell(r, 10).value
        sku = ws.cell(r, 17).value
        var_spec = ws.cell(r, 13).value
        
        if not name:
            continue
            
        name_str = str(name).strip()
        if name_str.startswith('📂'):
            continue
            
        is_var = name_str.startswith('↳')
        
        # Check if not on site
        if status_site == 'در سایت تعریف نشده' or (status_site and 'تعریف نشده' in str(status_site)):
            item = {
                'sheet': sheet_name,
                'row': r,
                'name': name_str,
                'is_variation': is_var,
                'parent_name': current_parent['name'] if is_var and current_parent else None,
                'brand': brand or (current_parent['brand'] if is_var and current_parent else None),
                'cat': cat or (current_parent['cat'] if is_var and current_parent else None),
                'subcat': subcat or (current_parent['subcat'] if is_var and current_parent else None),
                'status': status_site,
                'link': link,
                'sku': sku,
                'var_spec': var_spec
            }
            sheet_missing.append(item)
            
        if not is_var:
            current_parent = {
                'name': name_str,
                'brand': brand,
                'cat': cat,
                'subcat': subcat
            }
            
    if sheet_missing:
        print(f"Sheet '{sheet_name}': {len(sheet_missing)} missing products found")
        all_missing.extend(sheet_missing)

print(f"\nTOTAL missing items across all sheets: {len(all_missing)}")
with open('scratch/missing_products.json', 'w', encoding='utf-8') as f:
    json.dump(all_missing, f, ensure_ascii=False, indent=2)

