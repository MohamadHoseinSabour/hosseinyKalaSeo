#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
finalize_clean_catalog.py
=========================
۱. اعمال فرمت‌بندی شرطی پویا و استاندارد با اولویت‌بندی دقیق (Priority 1-8):
   - ستون H (اولویت ۱ و ۲):
     * قرمز برای ناموجود (FFF8D7DA با فونت FF721C24)
     * سبز برای موجود (FFD4EDDA با فونت FF155724)
   - ستون B و F (اولویت ۳ و ۴): هشدار قرمز برای قیمت‌دار شدن والد متغیر
   - ستون B و E (اولویت ۵ و ۶): هشدار نارنجی تغییر قیمت بالای ۵۰٪
   - سطرهای A تا R (اولویت ۷ و ۸): پس‌زمینه زرد برای «در سایت تعریف نشده» و نارنجی برای «فعلا دست نگه دار»
۲. ذخیره فایل اصلی hosseiny_kala_catalog.xlsx
۳. ایجاد نسخه پشتیبان یکتا و تمیز hosseiny_kala_catalog.backup.xlsx
۴. حذف تمامی فایل‌های اکسل اضافی و تکراری از ریشه پروژه
"""

import os
import sys
import glob
import shutil
from collections import OrderedDict
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAIN_CATALOG = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
BACKUP_CATALOG = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.backup.xlsx')

SHEETS = [
    'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
    'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
]

def reconfigure_catalog_conditional_formatting():
    print(f"📂 در حال باز کردن {MAIN_CATALOG}...")
    wb = openpyxl.load_workbook(MAIN_CATALOG)

    # رنگ‌های با آلفای کامل (FF) برای رندر بدون نقص در اکسل
    fill_stock_red = PatternFill(start_color="FFF8D7DA", end_color="FFF8D7DA", fill_type="solid")
    font_stock_red = Font(color="FF721C24", bold=True)

    fill_stock_green = PatternFill(start_color="FFD4EDDA", end_color="FFD4EDDA", fill_type="solid")
    font_stock_green = Font(color="FF155724", bold=True)

    fill_warning_orange = PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid")
    font_warning_orange = Font(color="FFC65911", bold=True)

    fill_parent_price_red = PatternFill(start_color="FFFCE8E6", end_color="FFFCE8E6", fill_type="solid")
    font_parent_price_red = Font(color="FFC00000", bold=True)

    fill_row_hold = PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid")
    fill_row_unreg = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")

    for sname in SHEETS:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        max_r = ws.max_row
        
        ws.conditional_formatting._cf_rules = OrderedDict()
        ws.conditional_formatting.max_priority = 0

        # اولویت ۱ و ۲: موجودی (ستون H)
        rule_stock_red = FormulaRule(
            formula=[f'ISNUMBER(SEARCH("ناموجود", H2))'],
            stopIfTrue=True,
            fill=fill_stock_red,
            font=font_stock_red
        )
        rule_stock_red.priority = 1

        rule_stock_green = FormulaRule(
            formula=[f'AND(ISNUMBER(SEARCH("موجود", H2)), NOT(ISNUMBER(SEARCH("ناموجود", H2))))'],
            stopIfTrue=True,
            fill=fill_stock_green,
            font=font_stock_green
        )
        rule_stock_green.priority = 2

        ws.conditional_formatting.add(f"H2:H{max_r}", rule_stock_red)
        ws.conditional_formatting.add(f"H2:H{max_r}", rule_stock_green)

        # اولویت ۳ و ۴: خطای قیمت والد متغیر (ستون B و F)
        rule_parent_b = FormulaRule(
            formula=[f'AND($L2="متغیر", OR(ISNUMBER(B2), ISNUMBER(F2)))'],
            stopIfTrue=True,
            fill=fill_parent_price_red,
            font=font_parent_price_red
        )
        rule_parent_b.priority = 3
        ws.conditional_formatting.add(f"B2:B{max_r}", rule_parent_b)

        rule_parent_f = FormulaRule(
            formula=[f'AND($L2="متغیر", OR(ISNUMBER(B2), ISNUMBER(F2)))'],
            stopIfTrue=True,
            fill=fill_parent_price_red,
            font=font_parent_price_red
        )
        rule_parent_f.priority = 4
        ws.conditional_formatting.add(f"F2:F{max_r}", rule_parent_f)

        # اولویت ۵ و ۶: هشدار تغییر قیمت بالای ۵۰٪ (ستون B و E)
        rule_price_b = FormulaRule(
            formula=[f'AND(ISNUMBER($B2), ISNUMBER($D2), $D2>0, ABS(($B2-$D2)/$D2)>=0.5, $L2<>"متغیر")'],
            stopIfTrue=True,
            fill=fill_warning_orange,
            font=font_warning_orange
        )
        rule_price_b.priority = 5
        ws.conditional_formatting.add(f"B2:B{max_r}", rule_price_b)

        rule_price_e = FormulaRule(
            formula=[f'AND(ISNUMBER($B2), ISNUMBER($D2), $D2>0, ABS(($B2-$D2)/$D2)>=0.5, $L2<>"متغیر")'],
            stopIfTrue=True,
            fill=fill_warning_orange,
            font=font_warning_orange
        )
        rule_price_e.priority = 6
        ws.conditional_formatting.add(f"E2:E{max_r}", rule_price_e)

        # اولویت ۷ و ۸: وضعیت سطر (A تا R)
        rule_row_hold = FormulaRule(
            formula=[f'$R2="فعلا دست نگه دار"'],
            stopIfTrue=False,
            fill=fill_row_hold
        )
        rule_row_hold.priority = 7

        rule_row_unreg = FormulaRule(
            formula=[f'$R2="در سایت تعریف نشده"'],
            stopIfTrue=False,
            fill=fill_row_unreg
        )
        rule_row_unreg.priority = 8

        ws.conditional_formatting.add(f"A2:R{max_r}", rule_row_hold)
        ws.conditional_formatting.add(f"A2:R{max_r}", rule_row_unreg)

        # اعتبارسنجی داده‌ها (Data Validation) فقط روی ستون‌های اختصاصی
        ws.data_validations.dataValidation.clear()
        
        dv_stock = DataValidation(
            type="list",
            formula1='"✓ موجود,✗ ناموجود"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="ورود نامعتبر",
            error="لطفاً یکی از گزینه‌های «✓ موجود» یا «✗ ناموجود» را انتخاب کنید."
        )
        dv_stock.add(f"H2:H{max_r}")
        ws.add_data_validation(dv_stock)

        dv_status = DataValidation(
            type="list",
            formula1='"ثبت‌شده در سایت,در سایت تعریف نشده,فعلا دست نگه دار"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="وضعیت نامعتبر",
            error="لطفاً یکی از گزینه‌های معتبر وضعیت در سایت را انتخاب کنید."
        )
        dv_status.add(f"R2:R{max_r}")
        ws.add_data_validation(dv_status)

        print(f"  ✓ شیت '{sname}': فرمت‌بندی شرطی داینامیک موجودی و قواعد کامل بازسازی شد.")

    wb.save(MAIN_CATALOG)
    print(f"✓ فایل اصلی کاتالوگ با موفقیت ذخیره شد.")

def cleanup_extra_excels():
    print("\n🧹 در حال پاکسازی فایل‌های اکسل اضافی...")
    shutil.copy2(MAIN_CATALOG, BACKUP_CATALOG)
    print(f"✓ نسخه پشتیبان به‌روزشده ایجاد شد: hosseiny_kala_catalog.backup.xlsx")

    all_excels = glob.glob(os.path.join(PROJECT_DIR, "*.xls*"))
    kept_files = {
        os.path.abspath(MAIN_CATALOG),
        os.path.abspath(BACKUP_CATALOG)
    }

    deleted_count = 0
    for f in all_excels:
        abs_f = os.path.abspath(f)
        if abs_f not in kept_files:
            try:
                os.remove(abs_f)
                print(f"  🗑️ حذف فایل اضافی: {os.path.basename(abs_f)}")
                deleted_count += 1
            except Exception as e:
                print(f"  ⚠️ خطا در حذف {os.path.basename(abs_f)}: {e}")

    print(f"✓ مجموع فایل‌های اضافی حذف‌شده: {deleted_count}")

def main():
    print("=" * 60)
    print(" بازسازی فرمت‌های شرطی داینامیک و نهایی‌سازی کاتالوگ اکسل")
    print("=" * 60)

    reconfigure_catalog_conditional_formatting()
    cleanup_extra_excels()

    print("\n" + "=" * 60)
    print("✅ کار با موفقیت پایان یافت!")
    print("=" * 60)

if __name__ == '__main__':
    main()
