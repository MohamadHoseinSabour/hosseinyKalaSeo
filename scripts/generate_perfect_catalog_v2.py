"""
generate_perfect_catalog_v2.py
==============================
ساخت کاتالوگ بی‌نقص با ساختار اصولی محصولات ساده و متغیر (والد + تنوع‌ها)
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_DIR = r"c:\Users\mmhhs\OneDrive\Documents\GitHub\hosseinyKalaSeo"
DATA_DIR = os.path.join(PROJECT_DIR, 'data')
RAW_FILE = os.path.join(DATA_DIR, 'products_raw.json')
EXCEL_OUTPUT = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
EXCEL_UPDATED = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog_updated.xlsx')

with open(RAW_FILE, encoding='utf-8') as f:
    site_products = json.load(f)

site_by_id = {p['product_id']: p for p in site_products}

# ساختار محصولات (ساده یا متغیر با سطرهای والد و تنوع)
# فرمت هر آیتم:
# برای محصول ساده:
#   {"name": "...", "type": "simple", "site_id": 1234, ...}
# برای محصول متغیر:
#   {"name": "...", "type": "variable", "site_id": 1234, "variants": [
#       {"name": "  ↳ رنگ: سفید", "attr": "سفید", "vid": 5678, "price": 1000, ...},
#       ...
#   ]}

STRUCTURE = {
    "یخچال و فریزر": {
        "کلور": [
            {
                "name": "یخچال ۱۰ فوت کلور مدل Mini Refrigerator",
                "type": "variable",
                "site_id": 2431,
                "variants": [
                    {"name": "    ↳ رنگ: سبز پاستیلی", "attr": "سبز پاستیلی", "vid": 2443, "price": 60000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: قرمز", "attr": "قرمز", "vid": 2439, "price": 60000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال و فریزر ۲۷ فوت کلور مدل کمبی FRNT101",
                "type": "variable",
                "site_id": 2408,
                "variants": [
                    {"name": "    ↳ رنگ: سفید ساده", "attr": "سفید ساده", "vid": 2418, "price": 115000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال و فریزر ۲۷ فوت کلور مدل کمبی هوم بار روسو",
                "type": "simple",
                "site_id": 2484,
                "price": 130000000,
                "in_stock": "موجود"
            },
            {
                "name": "یخچال فریزر دوقلو ۴۰ فوت کلور مدل گلوری عرض ۶۰",
                "type": "variable",
                "site_id": 2471,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2475, "price": 239000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2476, "price": 243000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال فریزر دوقلو ۴۰ فوت کلور مدل رسوپلاس عرض ۶۰",
                "type": "variable",
                "site_id": 2457,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2461, "price": 255000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2462, "price": 262000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل بدون آبریز", "attr": "استیل بدون آبریز", "vid": 2463, "price": 280000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال فریزر دوقلو ۴۴ فوت کلور مدل Glory Plus LM70",
                "type": "variable",
                "site_id": 2464,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2470, "price": 280000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2469, "price": 285000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال و فریزر ۲۷ فوت کلور مدل کمبی FRNT-70 Glory Plus",
                "type": "variable",
                "site_id": 2497,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2503, "price": 117000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2504, "price": 145000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال و فریزر دوقلو ۳۶ فوت کلور مدل RNT101-FNT101",
                "type": "variable",
                "site_id": 2637,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2646, "price": 210000000, "in_stock": "موجود"},
                ]
            },
        ],
        "دیپوینت": [
            {
                "name": "یخچال و فریزر ۲۸ فوت دیپوینت مدل T7",
                "type": "variable",
                "site_id": 2585,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2593, "price": 112000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: نقره‌ای / استیل", "attr": "نقره‌ای / استیل", "vid": 2592, "price": 115000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال فریزر ۳۰ فوت دیپوینت مدل Boss Pro",
                "type": "variable",
                "site_id": 2594,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2858, "price": 132000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2859, "price": 135000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال فریزر دوقلو ۳۴ فوت دیپوینت مدل MAX.DU مکس",
                "type": "variable",
                "site_id": 2745,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2748, "price": 195000000, "in_stock": "موجود"},
                ]
            },
        ],
        "ایستکول": [
            {
                "name": "یخچال ۵ فوت ایستکول مدل 2835",
                "type": "variable",
                "site_id": 2704,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2733, "price": 24500000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: نوک مدادی / مشکی", "attr": "نوک مدادی / مشکی", "vid": 2734, "price": 27127000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال فریزر ۹ فوت ایستکول مدل هتلی 2919",
                "type": "variable",
                "site_id": 2691,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2701, "price": 32926400, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: نوک مدادی", "attr": "نوک مدادی", "vid": 2702, "price": 36601900, "in_stock": "موجود"},
                ]
            },
        ],
        "الکترواستیل": [
            {
                "name": "یخچال فریزر دوقلو ۲۱ فوت الکترواستیل مدل فیوچر ES21",
                "type": "variable",
                "site_id": 2680,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2688, "price": 215000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: تیتانیوم / استیل", "attr": "تیتانیوم / استیل", "vid": 2689, "price": 225000000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "یخچال الکترواستیل یخساز تیتانیوم",
                "type": "simple",
                "site_id": None
            },
        ],
        "هیمالیا": [
            {
                "name": "یخچال هیمالیا مدل آیسپول 60MW",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ / طرح: استیل", "attr": "استیل", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ / طرح: چرم", "attr": "چرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {
                "name": "فریزر هیمالیا مدل آیسپول 60MW",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ / طرح: استیل", "attr": "استیل", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ / طرح: چرم", "attr": "چرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
        "بست (BOST)": [
            {"name": "یخچال فریزر BOST", "type": "simple", "site_id": None}
        ],
        "الکتروسان": [
            {"name": "فریزر صندوقی الکتروسان ۳۴۵", "type": "simple", "site_id": None}
        ],
        "متفرقه": [
            {"name": "یخچال فریزر R(re01)", "type": "simple", "site_id": None}
        ],
    },

    "ماشین لباسشویی": {
        "کلور": [
            {
                "name": "ماشین لباسشویی ۹ کیلوگرم اتوماتیک کلور مدل BMWT2",
                "type": "variable",
                "site_id": 2505,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2508, "price": 95800000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: نقره‌ای / اینوکس", "attr": "نقره‌ای / اینوکس", "vid": 2509, "price": 94900000, "in_stock": "موجود"},
                ]
            },
            {
                "name": "ماشین لباسشویی ۸ کیلوگرم اتوماتیک کلور مدل تاچ سری BMW T3",
                "type": "variable",
                "site_id": 2510,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2513, "price": 92000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: اینوکس", "attr": "اینوکس", "vid": 2514, "price": 92000000, "in_stock": "موجود"},
                ]
            },
            {"name": "لباسشویی کلور ۹ کیلویی کرمی", "type": "simple", "site_id": None},
            {"name": "لباسشویی کلور ۹ کیلویی سفید ویوا", "type": "simple", "site_id": None},
            {"name": "لباسشویی کلور مدل 9771 سفید", "type": "simple", "site_id": None},
        ],
        "دوو": [
            {
                "name": "ماشین لباسشویی هوشمند ۹ کیلوگرم اتوماتیک دوو مدل LM-9170",
                "type": "variable",
                "site_id": 2579,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2582, "price": 110600000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: نقره‌ای / استیل", "attr": "نقره‌ای / استیل", "vid": 2583, "price": 110600000, "in_stock": "موجود"},
                ]
            },
            {"name": "اتوماتیک دوو کینو ۹۰۵ سینیور", "type": "simple", "site_id": None},
        ],
        "پاکشوما": [
            {"name": "لباسشویی دوقلو پاکشوما هاک", "type": "simple", "site_id": None},
            {"name": "لباسشویی درب بالا پاکشوما ۱۲", "type": "simple", "site_id": None},
            {"name": "لباسشویی دوقلو ۲۱ کیلو پاکشوما", "type": "simple", "site_id": None},
        ],
        "ریتون": [
            {"name": "لباسشویی ۱۷ دوقلو ریتون", "type": "simple", "site_id": None},
        ],
        "پاکیزان": [
            {
                "name": "پتوشور پاکیزان",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ ظرفیت: ۲۰ کیلوگرم", "attr": "۲۰ کیلوگرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ ظرفیت: ۴۰ کیلوگرم", "attr": "۴۰ کیلوگرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
    },

    "جاروبرقی": {
        "هیتاچی": [
            {
                "name": "جاروبرقی سطلی هیتاچی مدل CV-950F",
                "type": "variable",
                "site_id": 2519,
                "variants": [
                    {"name": "    ↳ رنگ: شرابی", "attr": "شرابی", "vid": 2522, "price": 24500000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: مشکی", "attr": "مشکی", "vid": 2523, "price": 24500000, "in_stock": "موجود"},
                ]
            },
            {"name": "جاروبرقی هیتاچی سطلی ۱۲۰۰", "type": "simple", "site_id": None},
        ],
        "پاناسونیک": [
            {
                "name": "جاروبرقی پاناسونیک مدل MC-CG715 2100W",
                "type": "variable",
                "site_id": 2661,
                "variants": [
                    {"name": "    ↳ رنگ: سفید / استاندارد", "attr": "سفید / استاندارد", "vid": 2673, "price": 28000000, "in_stock": "موجود"},
                ]
            },
        ],
        "بست (BOST)": [
            {
                "name": "جاروبرقی بست سطلی مدل BVC-PC18B",
                "type": "variable",
                "site_id": 2885,
                "variants": [
                    {"name": "    ↳ رنگ: قرمز", "attr": "قرمز", "vid": 2888, "price": 13980000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: آبی", "attr": "آبی", "vid": 2889, "price": 13980000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: مشکی", "attr": "مشکی", "vid": 2890, "price": 13980000, "in_stock": "موجود"},
                ]
            },
        ],
        "دونالکس": [
            {"name": "جاروبرقی دونالکس ۷۵۰ سورمه‌ای", "type": "simple", "site_id": None},
            {"name": "جاروبرقی دونالکس ۷۷۰", "type": "simple", "site_id": None},
        ],
        "کلور": [
            {"name": "جاروبرقی کلور سطلی", "type": "simple", "site_id": None},
        ],
        "آراز": [
            {"name": "جاروبرقی آراز ۲۰ لیتر", "type": "simple", "site_id": None},
        ],
    },

    "کولر آبی و سرمایشی": {
        "ریتون": [
            {"name": "کولر آبی سلولزی ۱۶۰۰۰ ریتون پرتابل VENTO 16000", "type": "simple", "site_id": 2538, "price": 24000000, "in_stock": "موجود"},
            {"name": "کولر آبی سلولزی ۷۰۰۰ ریتون VENTO 7000", "type": "simple", "site_id": 2559, "price": 19000000, "in_stock": "موجود"},
        ],
        "پاکشوما": [
            {"name": "کولر گازی اسپلیت دیواری ۱۸۰۰۰ دور ثابت پاکشوما TPA-18CH", "type": "simple", "site_id": 3266, "price": 114000000, "sale_price": 76000000, "in_stock": "موجود"},
        ],
        "اوناکس": [
            {"name": "کولر آبی اوناکس 150L", "type": "simple", "site_id": None},
            {"name": "کولر آبی اوناکس B35", "type": "simple", "site_id": None},
        ],
        "اوجنرال": [
            {"name": "کولر پنجره‌ای اوجنرال ۱۸ هزار", "type": "simple", "site_id": None},
        ],
        "وستینگهاوس": [
            {"name": "کولر وستینگهاوس ۲۰ کوپلند", "type": "simple", "site_id": None},
        ],
        "سایر": [
            {"name": "کولر آبی مدل ۴۵۰۰", "type": "simple", "site_id": None},
        ],
    },

    "ماشین ظرفشویی": {
        "کلور": [
            {
                "name": "ماشین ظرفشویی ۱۵ نفره کلور مدل CDW-H15V",
                "type": "variable",
                "site_id": 2871,
                "variants": [
                    {"name": "    ↳ رنگ: سفید", "attr": "سفید", "vid": 2874, "price": 89000000, "in_stock": "موجود"},
                    {"name": "    ↳ رنگ: استیل", "attr": "استیل", "vid": 2875, "price": 89000000, "in_stock": "موجود"},
                ]
            },
        ],
        "انفره": [
            {"name": "ظرفشویی انفره استیل", "type": "simple", "site_id": None},
            {"name": "ظرفشویی انفره لمسی استیل", "type": "simple", "site_id": None},
        ],
    },

    "تصفیه آب و آبسردکن": {
        "سافت واتر": [
            {"name": "دستگاه تصفیه کننده آب سافت واتر ۶ مرحله مدل IR-06", "type": "simple", "site_id": 2766, "price": 16000000, "in_stock": "موجود"},
        ],
        "وست": [
            {"name": "تصفیه آب وست 5515", "type": "simple", "site_id": None},
        ],
        "هومگلد": [
            {"name": "تصفیه آب هومگلد", "type": "simple", "site_id": None},
        ],
        "C.C.K": [
            {"name": "تصفیه آب C.C.K", "type": "simple", "site_id": None},
        ],
        "نیکایی": [
            {"name": "آبسردکن نیکایی مدل ۱۱۲۰۶ ایستاده", "type": "simple", "site_id": None},
        ],
        "هانو": [
            {"name": "آبسردکن هانو", "type": "simple", "site_id": None},
        ],
        "افترون": [
            {"name": "آب سردکن افترون رومیزی", "type": "simple", "site_id": None},
        ],
        "الکترواستیل": [
            {"name": "آبگرمکن الکترواستیل مخزندار دیواری", "type": "simple", "site_id": None},
        ],
    },

    "لوازم برقی آشپزخانه": {
        "دونالکس": [
            {"name": "اسپرسوساز ۱۲ میلی‌لیتر دونالکس مدل DN480 (۳۸۰)", "type": "simple", "site_id": 3013, "price": 34000000, "in_stock": "موجود"},
            {"name": "اسپرسوساز ۱۸ میلی‌لیتر دونالکس مدل DN475 (۳۷۵)", "type": "simple", "site_id": 3016, "price": 29000000, "in_stock": "موجود"},
            {
                "name": "آبمیوه‌گیری ۱۰۰۰ وات دونالکس مدل DN-120 / DN-110",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ مدل: DN-120 (۱۰۰۰ وات)", "attr": "DN-120", "vid": None, "price": 21000000, "in_stock": "موجود", "site_id": 3019},
                    {"name": "    ↳ مدل: DN-110", "attr": "DN-110", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {"name": "سرخ‌کن ۸ لیتر دونالکس مدل DN-950 (۱۹۵۰)", "type": "simple", "site_id": 3026, "price": 27500000, "in_stock": "موجود"},
            {"name": "سرخکن دونالکس", "type": "simple", "site_id": None},
            {"name": "گوشت‌کوب برقی دونالکس مدل DN-275", "type": "simple", "site_id": 3032, "price": 10500000, "in_stock": "موجود"},
            {"name": "گوشتکوب و غذاساز دونالکس ۱۹۹۰", "type": "simple", "site_id": None},
            {"name": "مخلوط‌کن و آسیاب ۱.۵ لیتری دونالکس مدل DN-205", "type": "simple", "site_id": 3035, "price": 16500000, "in_stock": "موجود"},
            {"name": "چایساز دونالکس ۱۸۵ مشکی", "type": "simple", "site_id": None},
            {
                "name": "ساندویچ‌ساز دونالکس",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ: سفید (مدل ۴۰۹)", "attr": "سفید - ۴۰۹", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ: مشکی (مدل ۴۹۰)", "attr": "مشکی - ۴۹۰", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {"name": "همزن دستی دونالکس ۳۵۵", "type": "simple", "site_id": None},
        ],
        "گرنویک": [
            {"name": "اسپرسوساز نارنجی گرنویک", "type": "simple", "site_id": None},
            {"name": "چایساز ساورر گرنویک", "type": "simple", "site_id": None},
            {"name": "سرخکن ایرفرایر خاکستری گرنویک", "type": "simple", "site_id": None},
        ],
        "زامروتی": [
            {"name": "چایساز زامروتی", "type": "simple", "site_id": None},
            {"name": "کتری برقی تک زامروتی کرم", "type": "simple", "site_id": None},
            {
                "name": "کتری و قوری زامروتی مدل ۲۰۲۲",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ دسته: استیل", "attr": "دسته استیل", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ دسته: چوب", "attr": "دسته چوب", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {"name": "مخلوطکن زامروتی کرم", "type": "simple", "site_id": None},
        ],
        "وستینگهاوس": [
            {
                "name": "توستر وستینگهاوس",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ: کرم", "attr": "کرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ مدل: ۲۷", "attr": "مدل ۲۷", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {
                "name": "کتری برقی وستینگهاوس",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ: آبی", "attr": "آبی", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ: کرم", "attr": "کرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {
                "name": "مخلوط‌کن وستینگهاوس",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ: آبی", "attr": "آبی", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ: کرم", "attr": "کرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {
                "name": "همزن وستینگهاوس",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ رنگ: آبی", "attr": "آبی", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ رنگ: کرم", "attr": "کرم", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
        "پارس خزر": [
            {"name": "پلوپز پارس خزر ۸ نقره", "type": "simple", "site_id": None},
            {"name": "پلوپز پارس خزر مدل ۳۶۱", "type": "simple", "site_id": None},
        ],
        "فیلیپس": [
            {"name": "غذاساز فیلیپس ۷۷۷۸", "type": "simple", "site_id": None},
        ],
        "کنوود": [
            {"name": "همزن کنوود ۴۳۰", "type": "simple", "site_id": None},
        ],
        "کرکماز": [
            {"name": "چایساز کُرکماز A331", "type": "simple", "site_id": None},
        ],
        "سانی": [
            {"name": "کتری برقی سانی ۱.۵ لیتر", "type": "simple", "site_id": None},
        ],
        "سایر برندهای آشپزخانه": [
            {"name": "آبمیوهگیری جنرال ۶۰۷ غذاساز", "type": "simple", "site_id": None},
            {"name": "آسیاب کاستلو ۲۷۰", "type": "simple", "site_id": None},
            {"name": "سبزی خردکن KENRA", "type": "simple", "site_id": None},
            {"name": "سبزی خردکن نیکا ۲ کیلویی", "type": "simple", "site_id": None},
            {"name": "کتری برقی اریکسون", "type": "simple", "site_id": None},
            {"name": "کتری برقی وگاتی توپی", "type": "simple", "site_id": None},
            {"name": "کتری برقی پروکس", "type": "simple", "site_id": None},
            {"name": "کتری و قوری مدل هرمی تمام‌طلا ۶۰۴", "type": "simple", "site_id": None},
            {"name": "توستر ۶۰ لیتر", "type": "simple", "site_id": None},
            {"name": "توستر ناندلمونتی ۵۸", "type": "simple", "site_id": None},
            {"name": "همزن اوریکس کاسه چرخدار", "type": "simple", "site_id": None},
            {
                "name": "زودپز اویزپلاس مشکی رزگلد",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ مدل: تک‌قلو", "attr": "تک‌قلو", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ مدل: دوقلو", "attr": "دوقلو", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
    },

    "ظروف و پخت و پز": {
        "تابه اویز": [
            {"name": "تابه اویز سایز ۳۰", "type": "simple", "site_id": None},
            {"name": "تابه اویز سایز ۲۲ مشکی", "type": "simple", "site_id": None},
            {"name": "تابه اویز سایز ۲۲ فوم مشکی دربدار", "type": "simple", "site_id": None},
            {"name": "تابه اویز سایز ۲۶", "type": "simple", "site_id": None},
        ],
        "قابلمه اویز تک": [
            {"name": "قابلمه اویز سایز ۱۶ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه اویز سایز ۱۸ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه اویز سایز ۲۰ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه اویز سایز ۲۲ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه اویز سایز ۲۴ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه اویز سایز ۲۶ مشکی", "type": "simple", "site_id": None},
            {"name": "قابلمه مرغخوری اویزپلاس مشکی", "type": "simple", "site_id": None},
        ],
        "سرویس‌های قابلمه": [
            {"name": "سرویس خمره قابلمه اویز کرم‌سلین", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه استیل ۱۰ پارچه", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه اویز ۱۰ پارچه سرمه‌ای", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه اویزپلاس ۱۰ پارچه اناری", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه اویزپلاس ۱۲ پارچه مشکی", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه اویزپلاس ۱۲ پارچه سبز", "type": "simple", "site_id": None},
            {"name": "سرویس قابلمه اویزپلاس ویولا ۱۱ پارچه", "type": "simple", "site_id": None},
            {"name": "شیرجوش مشکی اویزپلاس", "type": "simple", "site_id": None},
        ],
    },

    "اصلاح و زیبایی": {
        "پرو امگا": [
            {"name": "سشوار حرفه‌ای پرو امگا مدل 8806 (۲۲۰۰ وات)", "type": "simple", "site_id": 3049, "price": 7000000, "in_stock": "موجود"},
            {"name": "سشوار امگا ۲۰۰۰ وات", "type": "simple", "site_id": None},
            {"name": "ماشین اصلاح موی سر و صورت پرو امگا مدل T5 (تک)", "type": "simple", "site_id": 3056, "price": 4500000, "in_stock": "موجود"},
            {"name": "ماشین اصلاح موی سر و صورت پرو امگا مدل T7 (۴ کاره)", "type": "simple", "site_id": 3065, "price": 6500000, "in_stock": "موجود"},
            {"name": "ماشین اصلاح پروامگا شارژی و برقی", "type": "simple", "site_id": None},
        ],
        "پرومارون": [
            {"name": "فرکننده مو پرومارون ۴+", "type": "simple", "site_id": None},
            {"name": "پرومارون برس حرارتی ۱۱۳", "type": "simple", "site_id": None},
            {
                "name": "اتومو پرومارون",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ مدل: ۱۱۱۵", "attr": "مدل ۱۱۱۵", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ مدل: ۱۱۱۶", "attr": "مدل ۱۱۱۶", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
            {"name": "اپیلاتور پرومارون ۵ کاره ۶۱۴", "type": "simple", "site_id": None},
            {"name": "ریشتراش پرومارون ۷۰۶", "type": "simple", "site_id": None},
        ],
        "پاناسونیک": [
            {"name": "اتو مو پاناسونیک HV11", "type": "simple", "site_id": None},
        ],
        "سایر": [
            {"name": "اتو مو Baby3383A", "type": "simple", "site_id": None},
            {"name": "دستگاه پیرایش سیمی صورت ۷۲۱", "type": "simple", "site_id": None},
        ],
    },

    "اتو بخار": {
        "دونالکس": [
            {"name": "اتو بخار ۲۵۰۰ وات دونالکس مدل 815", "type": "simple", "site_id": 3007, "price": 10500000, "in_stock": "موجود"},
        ],
        "فیلیپس": [
            {"name": "اتو فیلیپس ۵۰۴۰", "type": "simple", "site_id": None},
        ],
    },

    "گاز، فر و هود": {
        "آلتون": [
            {"name": "گاز صفحه‌ای آلتون استیل S517", "type": "simple", "site_id": None},
            {"name": "گاز صفحه‌ای آلتون استیل SG518", "type": "simple", "site_id": None},
            {"name": "فر توکار آلتون ۷۳۰۳", "type": "simple", "site_id": None},
            {"name": "فرگاز آلتون MPS", "type": "simple", "site_id": None},
            {"name": "فرگاز آلتون MD5W", "type": "simple", "site_id": None},
            {"name": "هود آلتون H312", "type": "simple", "site_id": None},
            {"name": "هود آلتون H315B", "type": "simple", "site_id": None},
        ],
        "الکترواستیل": [
            {
                "name": "فر توکار الکترواستیل",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ مدل: تیتانیوم ۱۱", "attr": "تیتانیوم ۱۱", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ مدل: عرض ۶۰ چرمی ۲۱", "attr": "عرض ۶۰ چرمی ۲۱", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
    },

    "صوتی، تصویری و گرمایشی": {
        "تلویزیون": [
            {
                "name": "تلویزیون ال ای دی هوشمند هاورسون (Averson)",
                "type": "variable",
                "site_id": None,
                "variants": [
                    {"name": "    ↳ سایز: ۳۲ اینچ", "attr": "۳۲ اینچ", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ سایز: ۵۵ اینچ پنرست", "attr": "۵۵ اینچ پنرست", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ سایز: ۶۵ اینچ", "attr": "۶۵ اینچ", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ سایز: ۷۵ اینچ", "attr": "۷۵ اینچ", "vid": None, "price": None, "in_stock": "ناموجود"},
                    {"name": "    ↳ سایز: ۸۵ اینچ", "attr": "۸۵ اینچ", "vid": None, "price": None, "in_stock": "ناموجود"},
                ]
            },
        ],
        "بخاری": [
            {"name": "بخاری گلدن لوکس", "type": "simple", "site_id": None},
            {"name": "بخاری پارس خزر فندار FH-2000P", "type": "simple", "site_id": None},
            {"name": "بخاری پلار 32-067 شیشه‌ای", "type": "simple", "site_id": None},
        ],
        "پنکه": [
            {"name": "پنکه پایه‌بلند 1455", "type": "simple", "site_id": None},
            {"name": "پنکه پارس خزر پایه‌بلند", "type": "simple", "site_id": None},
            {"name": "پنکه پاناسونیک ۲ حالته", "type": "simple", "site_id": None},
            {"name": "پنکه سانی ۷۶۰۰", "type": "simple", "site_id": None},
            {"name": "پنکه سانی مدل ۲۰۱۳ ساده", "type": "simple", "site_id": None},
        ],
        "متفرقه": [
            {"name": "کسورات اضافات کالا", "type": "simple", "site_id": None},
        ],
    },
}

COLUMNS = [
    # (key,              header_text,                 width,  align,   is_price,  is_formula)
    ("name",             "نام محصول",                  40,     "right",   False,     False),
    ("price_regular",    "قیمت عادی (ریال)",          18,     "center",  True,      False),
    ("price_reg_toman",  "قیمت عادی (تومان)",          18,     "center",  True,      True),
    ("price_sale",       "قیمت حراج (ریال)",          18,     "center",  True,      False),
    ("price_sale_toman", "قیمت حراج (تومان)",          18,     "center",  True,      True),
    ("in_stock",         "وضعیت موجودی",               14,     "center",  False,     False),
    ("category_path",    "دسته‌بندی / برند",           26,     "right",   False,     False),
    ("type",             "نوع محصول",                  13,     "center",  False,     False),
    ("variant_attr",     "تنوع / رنگ",                 18,     "center",  False,     False),
    ("product_id",       "شناسه محصول (ID)",           16,     "center",  False,     False),
    ("variation_id",     "شناسه تنوع",                 15,     "center",  False,     False),
    ("link",             "لینک محصول در سایت",         36,     "left",    False,     False),
    ("sku",              "کد SKU",                     14,     "center",  False,     False),
    ("status_note",      "وضعیت در سایت",              22,     "center",  False,     False),
]

COLOR_HEADER_BG = "1F4E79"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_SUBCAT_BG = "E2EFDA"
COLOR_SUBCAT_TEXT = "375623"
COLOR_VAR_PARENT_BG = "D9E1F2"   # بنفش/آبی ملایم برای سطر والد محصول متغیر
COLOR_VARIATION_BG = "F2F2F2"    # خاکستری روشن برای سطرهای تنوع/رنگ
COLOR_MISSING_BG = "FFF2CC"      # زرد ملایم برای محصول ساده ثبت نشده در سایت
COLOR_TOMAN_HEADER_BG = "2E75B6"
COLOR_IN_STOCK_TEXT = "27AE60"
COLOR_OUT_OF_STOCK = "C00000"

def make_border():
    s = Side(style='thin', color='D9D9D9')
    return Border(top=s, bottom=s, left=s, right=s)

def build_catalog_v2():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    border_thin = make_border()
    total_rows = 0

    # شیت راهنما
    ws_guide = wb.create_sheet("راهنمای استفاده", index=0)
    ws_guide.sheet_view.rightToLeft = True
    guide_lines = [
        ("راهنمای جامع کاتالوگ محصولات فروشگاه حسینی کالا", ""),
        ("", ""),
        ("بخش / ستون", "توضیحات و راهنمای کاربری"),
        ("نام محصول (ستون راست)", "نام دقیق کالا (سطر والد یا تنوع‌های رنگ/مدل)"),
        ("قیمت عادی (ریال)", "قیمت پایه به ریال (عدد خالص). با فرمت ۳ رقم ۳ رقم خودکار اکسل (#,##0)."),
        ("قیمت عادی (تومان)", "فرمول محاسبه خودکار قیمت به تومان (=B{row}/10) با تفکیک ۳ رقم."),
        ("قیمت حراج (ریال)", "قیمت تخفیف‌خورده به ریال (در صورت وجود تخفیف فعال در سایت)."),
        ("قیمت حراج (تومان)", "فرمول محاسبه خودکار قیمت حراج به تومان (=D{row}/10)."),
        ("وضعیت موجودی (منوی کشویی)", "دکمه و لیست انتخابی بین «موجود» و «ناموجود» جهت تغییر سریع موجودی انبار."),
        ("دسته‌بندی / برند", "مسیر کامل دسته اصلی و برند یا زیردسته کالا"),
        ("نوع محصول", "ساده (یک مدل/قیمت ثابت) | متغیر (سطر والد کالای چندرنگ) | تنوع (هر رنگ مجزا)"),
        ("شناسه محصول (ID)", "شناسه یکتای عددی وردپرس جهت همگام‌سازی و آپدیت خودکار قیمت‌ها"),
        ("شناسه تنوع (Variation ID)", "شناسه متغیر رنگ خاص (جهت به‌روزرسانی مجزای قیمت هر رنگ)"),
        ("لینک محصول در سایت", "لینک کوتاه و بدون خطای انکودینگ با فرمت عددی ?p=ID"),
        ("", ""),
        ("راهنمای رنگ‌بندی سطرها:", ""),
        ("آبی ملایم", "کالای والد متغیر که شامل چند رنگ یا مدل مختلف است."),
        ("خاکستری روشن", "سطرهای تنوع اختصاصی (هر رنگ/مدل) با قیمت و Variation ID مجزا."),
        ("زرد ملایم", "کالاهایی که هنوز روی سایت ثبت نشده‌اند (قیمت‌ها خالی و آماده ورود اطلاعات است)."),
        ("سفید", "کالاهای ساده موجود در سایت با اطلاعات کامل."),
        ("", ""),
        ("نحوه به‌روزرسانی قیمت‌ها به سایت:", ""),
        ("۱. قیمت‌های جدید را در ستون 'قیمت عادی (ریال)' یا 'قیمت حراج (ریال)' همین فایل وارد کنید (قیمت تومان خودکار محاسبه می‌شود).", ""),
        ("۲. وضعیت موجودی را از منوی کشویی روی موجود یا ناموجود تنظیم کنید.", ""),
        ("۳. اسکریپت python scripts/update_prices.py --apply را اجرا کنید تا قیمت‌ها مستقیماً روی سایت اعمال شوند.", ""),
    ]
    for r_idx, (col1, col2) in enumerate(guide_lines, start=1):
        ws_guide.cell(row=r_idx, column=1, value=col1)
        ws_guide.cell(row=r_idx, column=2, value=col2)
        if r_idx == 1:
            ws_guide.cell(row=r_idx, column=1).font = Font(bold=True, size=13, color="1F4E79")
        elif r_idx in (3, 16, 22):
            ws_guide.cell(row=r_idx, column=1).font = Font(bold=True, size=11, color="1F4E79")
            ws_guide.cell(row=r_idx, column=2).font = Font(bold=True, size=11, color="1F4E79")

    ws_guide.column_dimensions['A'].width = 30
    ws_guide.column_dimensions['B'].width = 85

    for sheet_name, subcats in STRUCTURE.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        ws.sheet_view.rightToLeft = True

        # هدر شیت
        ws.row_dimensions[1].height = 28
        for col_idx, (f_key, h_text, width, align_mode, is_price, is_formula) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.font = Font(bold=True, color=COLOR_HEADER_TEXT, size=10, name="Calibri")
            h_bg = COLOR_TOMAN_HEADER_BG if is_formula else COLOR_HEADER_BG
            cell.fill = PatternFill(fill_type='solid', fgColor=h_bg)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border_thin
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

        row_idx = 2

        for subcat_name, products_list in subcats.items():
            # هدر زیردسته دقیقاً یک بار برای هر برند
            ws.row_dimensions[row_idx].height = 22
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(COLUMNS))
            sub_cell = ws.cell(row=row_idx, column=1, value=f"  📂 {sheet_name}  ◀  {subcat_name}")
            sub_cell.font = Font(bold=True, color=COLOR_SUBCAT_TEXT, size=10, name="Calibri")
            sub_cell.fill = PatternFill(fill_type='solid', fgColor=COLOR_SUBCAT_BG)
            sub_cell.alignment = Alignment(horizontal='right', vertical='center')
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=c).border = border_thin
            row_idx += 1

            cat_path = f"{sheet_name} / {subcat_name}"

            for prod in products_list:
                ptype = prod.get("type", "simple")
                pname = prod.get("name")
                site_id = prod.get("site_id")
                site_prod = site_by_id.get(site_id) if site_id else None

                # ── ۱. محصول ساده ──
                if ptype == "simple":
                    if site_prod:
                        reg_price = prod.get("price") or site_prod.get("price_regular") or ""
                        sale_price = prod.get("sale_price") or site_prod.get("price_sale") or ""
                        in_stk = prod.get("in_stock") or ("موجود" if site_prod.get("in_stock", True) else "ناموجود")
                        stk_color = COLOR_IN_STOCK_TEXT if in_stk == "موجود" else COLOR_OUT_OF_STOCK
                        row_fill = PatternFill(fill_type='solid', fgColor="FFFFFF")
                        link_val = site_prod.get("shortlink", "")
                        sku_val = site_prod.get("sku", "")
                        status_val = "✓ موجود در سایت"
                        status_color = "27AE60"
                    else:
                        reg_price = ""
                        sale_price = ""
                        in_stk = "ناموجود"
                        stk_color = COLOR_OUT_OF_STOCK
                        row_fill = PatternFill(fill_type='solid', fgColor=COLOR_MISSING_BG)
                        link_val = ""
                        sku_val = ""
                        status_val = "نیازمند ثبت در سایت"
                        status_color = "806000"

                    row_data = {
                        "name": pname,
                        "price_regular": reg_price,
                        "price_reg_toman": f'=IF(OR(B{row_idx}="",B{row_idx}=0),"",B{row_idx}/10)',
                        "price_sale": sale_price,
                        "price_sale_toman": f'=IF(OR(D{row_idx}="",D{row_idx}=0),"",D{row_idx}/10)',
                        "in_stock": in_stk,
                        "category_path": cat_path,
                        "type": "ساده",
                        "variant_attr": "-",
                        "product_id": site_id or "",
                        "variation_id": "",
                        "link": link_val,
                        "sku": sku_val,
                        "status_note": status_val,
                    }

                    ws.row_dimensions[row_idx].height = 20
                    for c_idx, (f_key, _, _, align_mode, is_price, is_formula) in enumerate(COLUMNS, start=1):
                        val = row_data.get(f_key, "")
                        c = ws.cell(row=row_idx, column=c_idx, value=val)
                        c.fill = row_fill
                        c.alignment = Alignment(horizontal=align_mode, vertical='center')
                        c.border = border_thin
                        c.font = Font(size=10, name="Calibri")

                        if is_price or is_formula:
                            c.number_format = '#,##0'
                            if isinstance(val, int) or is_formula:
                                c.font = Font(bold=True, size=10, color="1F4E79" if is_formula else "000000", name="Calibri")
                        if f_key == "in_stock":
                            c.font = Font(bold=True, color=stk_color, size=10, name="Calibri")
                        if f_key == "link" and val:
                            c.hyperlink = val
                            c.style = "Hyperlink"
                            c.fill = row_fill
                        if f_key == "status_note":
                            c.font = Font(size=9, italic=not site_prod, color=status_color, bold=bool(site_prod), name="Calibri")

                    row_idx += 1
                    total_rows += 1

                # ── ۲. محصول متغیر (سطر والد + سطرهای تنوع) ──
                elif ptype == "variable":
                    variants = prod.get("variants", [])
                    var_count_str = f"شامل {len(variants)} رنگ / تنوع"

                    # سطر والد (Parent Row)
                    if site_prod:
                        parent_reg_price = site_prod.get("price_regular") or (variants[0].get("price") if variants else "")
                        parent_sale_price = site_prod.get("price_sale") or ""
                        parent_in_stk = "موجود" if site_prod.get("in_stock", True) else "ناموجود"
                        parent_stk_color = COLOR_IN_STOCK_TEXT if parent_in_stk == "موجود" else COLOR_OUT_OF_STOCK
                        parent_link = site_prod.get("shortlink", "")
                        parent_sku = site_prod.get("sku", "")
                        parent_status = "✓ محصول متغیر در سایت"
                        parent_status_color = "27AE60"
                        parent_fill = PatternFill(fill_type='solid', fgColor=COLOR_VAR_PARENT_BG)
                    else:
                        parent_reg_price = ""
                        parent_sale_price = ""
                        parent_in_stk = "ناموجود"
                        parent_stk_color = COLOR_OUT_OF_STOCK
                        parent_link = ""
                        parent_sku = ""
                        parent_status = "نیازمند ثبت متغیر در سایت"
                        parent_status_color = "806000"
                        parent_fill = PatternFill(fill_type='solid', fgColor=COLOR_VAR_PARENT_BG)

                    parent_row_data = {
                        "name": pname,
                        "price_regular": parent_reg_price,
                        "price_reg_toman": f'=IF(OR(B{row_idx}="",B{row_idx}=0),"",B{row_idx}/10)',
                        "price_sale": parent_sale_price,
                        "price_sale_toman": f'=IF(OR(D{row_idx}="",D{row_idx}=0),"",D{row_idx}/10)',
                        "in_stock": parent_in_stk,
                        "category_path": cat_path,
                        "type": "متغیر",
                        "variant_attr": var_count_str,
                        "product_id": site_id or "",
                        "variation_id": "",
                        "link": parent_link,
                        "sku": parent_sku,
                        "status_note": parent_status,
                    }

                    ws.row_dimensions[row_idx].height = 21
                    for c_idx, (f_key, _, _, align_mode, is_price, is_formula) in enumerate(COLUMNS, start=1):
                        val = parent_row_data.get(f_key, "")
                        c = ws.cell(row=row_idx, column=c_idx, value=val)
                        c.fill = parent_fill
                        c.alignment = Alignment(horizontal=align_mode, vertical='center')
                        c.border = border_thin
                        c.font = Font(bold=(f_key in ("name", "type")), size=10, name="Calibri")

                        if is_price or is_formula:
                            c.number_format = '#,##0'
                            if isinstance(val, int) or is_formula:
                                c.font = Font(bold=True, size=10, color="1F4E79" if is_formula else "000000", name="Calibri")
                        if f_key == "in_stock":
                            c.font = Font(bold=True, color=parent_stk_color, size=10, name="Calibri")
                        if f_key == "link" and val:
                            c.hyperlink = val
                            c.style = "Hyperlink"
                            c.fill = parent_fill
                        if f_key == "status_note":
                            c.font = Font(size=9, color=parent_status_color, bold=bool(site_prod), name="Calibri")

                    row_idx += 1
                    total_rows += 1

                    # سطرهای فرزند (Child Variations)
                    for v in variants:
                        vname = v.get("name")
                        vattr = v.get("attr", "-")
                        vid = v.get("vid")
                        vprice = v.get("price") or ""
                        vsale = v.get("sale_price") or ""
                        vstock = v.get("in_stock", "ناموجود")
                        vstk_color = COLOR_IN_STOCK_TEXT if vstock == "موجود" else COLOR_OUT_OF_STOCK
                        vsite_id = v.get("site_id") or site_id

                        if vid:
                            vlink = f"https://hosseinykala.ir/?p={vid}"
                            vstatus = "✓ تنوع فعال در سایت"
                            vstatus_color = "555555"
                        elif vsite_id:
                            vlink = f"https://hosseinykala.ir/?p={vsite_id}"
                            vstatus = "✓ تنوع در سایت"
                            vstatus_color = "555555"
                        else:
                            vlink = ""
                            vstatus = "نیازمند تعریف تنوع"
                            vstatus_color = "806000"

                        var_fill = PatternFill(fill_type='solid', fgColor=COLOR_VARIATION_BG)

                        var_row_data = {
                            "name": vname,
                            "price_regular": vprice,
                            "price_reg_toman": f'=IF(OR(B{row_idx}="",B{row_idx}=0),"",B{row_idx}/10)',
                            "price_sale": vsale,
                            "price_sale_toman": f'=IF(OR(D{row_idx}="",D{row_idx}=0),"",D{row_idx}/10)',
                            "in_stock": vstock,
                            "category_path": cat_path,
                            "type": "تنوع (واریاسیون)",
                            "variant_attr": vattr,
                            "product_id": vsite_id or "",
                            "variation_id": vid or "",
                            "link": vlink,
                            "sku": "",
                            "status_note": vstatus,
                        }

                        ws.row_dimensions[row_idx].height = 19
                        for c_idx, (f_key, _, _, align_mode, is_price, is_formula) in enumerate(COLUMNS, start=1):
                            val = var_row_data.get(f_key, "")
                            c = ws.cell(row=row_idx, column=c_idx, value=val)
                            c.fill = var_fill
                            c.alignment = Alignment(horizontal=align_mode, vertical='center')
                            c.border = border_thin
                            c.font = Font(size=9.5, name="Calibri")

                            if f_key == "name":
                                c.font = Font(size=9.5, italic=True, color="333333", name="Calibri")
                            if is_price or is_formula:
                                c.number_format = '#,##0'
                                c.font = Font(bold=True, size=9.5, color="1F4E79", name="Calibri")
                            if f_key == "in_stock":
                                c.font = Font(bold=True, color=vstk_color, size=9.5, name="Calibri")
                            if f_key == "link" and val:
                                c.hyperlink = val
                                c.style = "Hyperlink"
                                c.fill = var_fill
                            if f_key == "status_note":
                                c.font = Font(size=9, color=vstatus_color, name="Calibri")

                        row_idx += 1
                        total_rows += 1

        # منوی کشویی انتخاب موجودی
        dv_stock = DataValidation(
            type="list",
            formula1='"موجود,ناموجود"',
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=True,
            errorTitle="خطای انتخاب",
            error="لطفاً وضعیت موجودی را از بین گزینه‌های «موجود» یا «ناموجود» انتخاب کنید."
        )
        ws.add_data_validation(dv_stock)
        dv_stock.add(f"F2:F{row_idx}")

    try:
        wb.save(EXCEL_OUTPUT)
        final_file = EXCEL_OUTPUT
    except PermissionError:
        wb.save(EXCEL_UPDATED)
        final_file = EXCEL_UPDATED
        print(f"⚠️ فایل اصلی در اکسل باز است. فایل در مسیر زیر ذخیره شد:\n   {EXCEL_UPDATED}")

    print("=" * 60)
    print("✓ ساخت کاتالوگ دقیق و اصولی (با ساختار والد و تنوع) به پایان رسید.")
    print(f"📁 مسیر فایل خروجی: {final_file}")
    print(f"📊 کل سطرهای ثبت‌شده: {total_rows}")
    print("=" * 60)

if __name__ == '__main__':
    build_catalog_v2()
