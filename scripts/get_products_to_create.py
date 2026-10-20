"""
get_products_to_create.py
=========================
ماژول و اسکریپت فیلتر هوشمند کالاهای آماده ایجاد در ووکامرس / سایت:
- تفکیک دقیق کالاهای «در سایت تعریف نشده»
- رعایت اکید شرط و درخواست کاربر: حذف و رد کامل (Skip) اقلام با وضعیت «فعلا دست نگه دار»
- حذف اقلام از قبل ثبت‌شده «ثبت‌شده در سایت»
- خروجی ساختاریافته برای استفاده در اسکریپت‌های بعدی ساخت محصول با هوش مصنوعی و ووکامرس REST API
"""

import os
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = r"c:\Users\mmhhs\OneDrive\Documents\GitHub\hosseinyKalaSeo"
EXCEL_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')

CATEGORY_SHEETS = [
    'یخچال و فریزر',
    'شستشو و نظافت',
    'لوازم برقی و پخت و پز',
    'نوشیدنی ساز',
    'سرمایش و گرمایش',
    'تصفیه کننده آب',
    'صوتی و تصویری',
    'زیبایی و سلامت فردی'
]

def get_pending_products(excel_path=EXCEL_PATH):
    """
    خواندن اکسل و استخراج اقلام آماده ایجاد.
    اقلام دارای وضعیت «فعلا دست نگه دار» به طور خودکار نادیده گرفته (Skip) می‌شوند.
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    ready_to_create = []
    on_hold_items = []
    already_on_site = []

    for sheetname in CATEGORY_SHEETS:
        if sheetname not in wb.sheetnames:
            continue
        ws = wb[sheetname]
        
        # ردیابی سطرها
        for r in range(2, ws.max_row + 1):
            name = ws.cell(row=r, column=1).value
            if not name or '📂' in str(name):
                continue
            
            name_str = str(name).strip()
            price_reg = ws.cell(row=r, column=2).value
            price_sale = ws.cell(row=r, column=4).value
            stock = ws.cell(row=r, column=6).value or '✓ موجود'
            main_cat = ws.cell(row=r, column=7).value or sheetname
            subcat = ws.cell(row=r, column=8).value or ''
            brand = ws.cell(row=r, column=9).value or ''
            ptype = ws.cell(row=r, column=10).value or 'ساده'
            variant_attr = ws.cell(row=r, column=11).value or ''
            pid = ws.cell(row=r, column=12).value
            vid = ws.cell(row=r, column=13).value
            status_val = ws.cell(row=r, column=16).value or ''
            status_str = str(status_val).strip()

            item_data = {
                'sheet': sheetname,
                'row': r,
                'name': name_str,
                'price_regular': price_reg,
                'price_sale': price_sale,
                'stock': stock,
                'main_cat': main_cat,
                'subcat': subcat,
                'brand': brand,
                'type': ptype,
                'variant_attr': variant_attr,
                'product_id': pid,
                'variation_id': vid,
                'status': status_str
            }

            if status_str == 'فعلا دست نگه دار':
                on_hold_items.append(item_data)
            elif status_str == 'در سایت تعریف نشده':
                ready_to_create.append(item_data)
            elif status_str == 'ثبت‌شده در سایت':
                already_on_site.append(item_data)
            else:
                # اگر فیلد خالی بود ولی شناسه یا لینک نداشت
                if not pid:
                    ready_to_create.append(item_data)
                else:
                    already_on_site.append(item_data)

    return {
        'ready_to_create': ready_to_create,
        'on_hold': on_hold_items,
        'already_on_site': already_on_site
    }

if __name__ == '__main__':
    data = get_pending_products()
    print("=" * 60)
    print("گزارش غربالگری کاتالوگ برای ساخت محصولات جدید در سایت:")
    print("=" * 60)
    print(f"  🟢 اقلام آماده ساخت در سایت (در سایت تعریف نشده): {len(data['ready_to_create'])}")
    print(f"  🟠 اقلام متوقف‌شده توسط کاربر (فعلا دست نگه دار): {len(data['on_hold'])}")
    print(f"  🔵 اقلام از قبل فعال در سایت (ثبت‌شده در سایت): {len(data['already_on_site'])}")
    print("-" * 60)
    
    if data['on_hold']:
        print("\nلیست اقلام متوقف‌شده (فعلا دست نگه دار):")
        for it in data['on_hold']:
            print(f"  - [{it['sheet']}] {it['name']} ({it['brand']})")
    else:
        print("\nدر حال حاضر هیچ کالایی روی وضعیت «فعلا دست نگه دار» قرار داده نشده است.")
        print("کاربر می‌تواند با باز کردن اکسل، وضعیت هر کالای دلخواه در ستون P را روی «فعلا دست نگه دار» بگذارد.")
        print("با این کار رنگ سطر به نارنجی ملایم تبدیل شده و هنگام اجرای اسکریپت ساخت، ایجاد نخواهد شد.")
