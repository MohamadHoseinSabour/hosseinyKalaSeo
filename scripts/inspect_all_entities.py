"""
inspect_all_entities.py
=======================
استخراج تمام ۹۷ محصول لایو، ۷۸ تنوع، و ۷۴ قلم آفلاین کاتالوگ فعلی
جهت ایجاد پایگاه دانش مرجع برای ساب‌ایجنت‌ها
"""

import os
import sys
import json
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = r"c:\Users\mmhhs\OneDrive\Documents\GitHub\hosseinyKalaSeo"
RAW_FILE = os.path.join(PROJECT_DIR, 'data', 'products_raw.json')
CATALOG_FILE = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
OUTPUT_REF = os.path.join(PROJECT_DIR, 'data', 'current_catalog_reference.json')

with open(RAW_FILE, 'r', encoding='utf-8') as f:
    live_products = json.load(f)

# بارگذاری کاتالوگ اکسل فعلی
wb = openpyxl.load_workbook(CATALOG_FILE, data_only=True)

excel_items = []
for sheet in wb.sheetnames:
    if sheet in ['📊 خلاصه و داشبورد', '📖 راهنمای کاربری']:
        continue
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or str(name).strip().startswith('📂') or str(name).strip().startswith('•'):
            continue
        p_reg = ws.cell(r, 2).value
        p_sale = ws.cell(r, 4).value
        stock = ws.cell(r, 6).value
        main_cat = ws.cell(r, 7).value or sheet
        subcat = ws.cell(r, 8).value
        brand = ws.cell(r, 9).value
        ptype = ws.cell(r, 10).value
        var_attr = ws.cell(r, 11).value
        pid = ws.cell(r, 12).value
        vid = ws.cell(r, 13).value
        link = ws.cell(r, 14).value
        sku = ws.cell(r, 15).value
        status = ws.cell(r, 16).value
        
        excel_items.append({
            "sheet": sheet,
            "row": r,
            "name": str(name).strip(),
            "price_regular": p_reg,
            "price_sale": p_sale,
            "in_stock": stock,
            "main_cat": main_cat,
            "subcat": subcat,
            "brand": brand,
            "type": ptype,
            "variant_attr": var_attr,
            "product_id": pid,
            "variation_id": vid,
            "link": link,
            "sku": sku,
            "status_note": status
        })

reference_data = {
    "live_site_products_count": len(live_products),
    "current_excel_rows_count": len(excel_items),
    "live_products": [
        {
            "id": p["id"],
            "name": p["name"],
            "type": p["type"],
            "sku": p.get("sku", ""),
            "permalink": p.get("permalink", ""),
            "in_stock": p.get("in_stock", False),
            "price_regular": p.get("price_regular"),
            "price_sale": p.get("price_sale"),
            "categories": [c["name"] if isinstance(c, dict) else str(c) for c in p.get("categories", [])],
            "variations": [
                {
                    "variation_id": v["variation_id"],
                    "attributes": v.get("attributes", {}),
                    "price_regular": v.get("price_regular"),
                    "price_sale": v.get("price_sale"),
                    "in_stock": v.get("in_stock", False),
                    "sku": v.get("sku", "")
                }
                for v in p.get("variations", [])
            ]
        }
        for p in live_products
    ],
    "excel_items": excel_items
}

with open(OUTPUT_REF, 'w', encoding='utf-8') as f:
    json.dump(reference_data, f, ensure_ascii=False, indent=2)

print(f"پایگاه مرجع ایجاد شد:")
print(f"- محصولات لایو: {len(live_products)}")
print(f"- اقلام اکسل: {len(excel_items)}")
print(f"- ذخیره در: {OUTPUT_REF}")
