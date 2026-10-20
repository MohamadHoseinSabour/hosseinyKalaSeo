import os
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')

wb = openpyxl.load_workbook(CATALOG_PATH, data_only=True)
sheets = [
    'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
    'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
]

print("=== INSPECTING CATALOG ROWS ===")

stats = {
    'total_rows': 0,
    'simple_prods': 0,
    'variable_parents': 0,
    'variations': 0,
    'headers_or_empty': 0,
    'no_discount_items': 0,
    'with_discount_items': 0,
    'no_price_items': 0
}

items_needing_discount = []
anomalies = []

for sname in sheets:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    
    # Header col mapping
    col_map = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or '').strip()
        if 'نام محصول' in h: col_map['name'] = c
        elif 'قیمت عادی (ریال)' in h: col_map['price_reg'] = c
        elif 'قیمت عادی (تومان)' in h: col_map['price_reg_toman'] = c
        elif 'آخرین قیمت مشخص‌شده (ریال)' in h: col_map['last_price'] = c
        elif 'درصد تغییر قیمت' in h: col_map['change_pct'] = c
        elif 'قیمت حراج (ریال)' in h: col_map['price_sale'] = c
        elif 'قیمت حراج (تومان)' in h: col_map['price_sale_toman'] = c
        elif 'وضعیت موجودی' in h: col_map['stock'] = c
        elif 'نوع محصول' in h: col_map['type'] = c
        elif 'مشخصه تنوع' in h: col_map['var_spec'] = c
        elif 'شناسه محصول' in h: col_map['pid'] = c
        elif 'شناسه تنوع' in h: col_map['vid'] = c
        elif 'لینک مستقیم' in h: col_map['link'] = c
        elif 'وضعیت در سایت' in h: col_map['status_site'] = c

    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, col_map.get('name', 1)).value
        if not name:
            continue
        name_str = str(name).strip()
        if name_str.startswith('📂') or name_str.startswith('•'):
            stats['headers_or_empty'] += 1
            continue
        
        stats['total_rows'] += 1
        ptype = str(ws.cell(r, col_map.get('type', 12)).value or '').strip()
        reg_price = ws.cell(r, col_map.get('price_reg', 2)).value
        sale_price = ws.cell(r, col_map.get('price_sale', 6)).value
        stock = ws.cell(r, col_map.get('stock', 8)).value
        pid = ws.cell(r, col_map.get('pid', 14)).value
        vid = ws.cell(r, col_map.get('vid', 15)).value
        status_site = ws.cell(r, col_map.get('status_site', 18)).value
        link = ws.cell(r, col_map.get('link', 16)).value

        has_price = isinstance(reg_price, (int, float)) and reg_price > 0
        has_sale = isinstance(sale_price, (int, float)) and sale_price > 0

        if ptype == 'متغیر':
            stats['variable_parents'] += 1
            if has_price:
                anomalies.append(f"Variable parent has price: {sname} row {r} ({name_str})")
        elif ptype in ['تنوع', 'متغیر-فرزند']:
            stats['variations'] += 1
        else:
            stats['simple_prods'] += 1

        if ptype != 'متغیر':
            if has_price and not has_sale:
                stats['no_discount_items'] += 1
                items_needing_discount.append({
                    'sheet': sname,
                    'row': r,
                    'name': name_str,
                    'type': ptype,
                    'reg_price': reg_price,
                    'sale_price': sale_price,
                    'stock': stock,
                    'pid': pid,
                    'vid': vid,
                    'status_site': status_site
                })
            elif has_price and has_sale:
                stats['with_discount_items'] += 1
            else:
                stats['no_price_items'] += 1

print("\nStats Summary:")
for k, v in stats.items():
    print(f"  {k}: {v}")

print(f"\nItems needing 5% discount: {len(items_needing_discount)}")
for it in items_needing_discount:
    print(f"  [{it['sheet']}:R{it['row']}] {it['name']} | Reg: {it['reg_price']:,} | PID: {it['pid']} | VID: {it['vid']} | PType: {it['type']}")

if anomalies:
    print("\nAnomalies detected:")
    for a in anomalies:
        print(f"  {a}")
