import os
import sys
import json
import re
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
DUMP_PATH = os.path.join(PROJECT_DIR, 'scratch', 'wc_live_dump.json')

with open(DUMP_PATH, 'r', encoding='utf-8') as f:
    dump = json.load(f)

site_vars_by_parent = dump['variations']

wb = openpyxl.load_workbook(CATALOG_PATH, data_only=True)
sheets = [
    'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
    'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
]

matched = 0
unmatched = []

def clean_opt(text):
    if not text:
        return ''
    text = str(text).replace('↳', '').strip()
    text = re.sub(r'^(رنگ|سایز|ظرفیت|مدل|طول|مشخصه|طرح)\s*(\([^)]*\))?\s*:\s*', '', text).strip()
    return text.replace('ي', 'ی').replace('ك', 'ک').lower()

matches_found = []

for sname in sheets:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        ptype = str(ws.cell(r, 12).value or '').strip()
        if ptype not in ['تنوع', 'متغیر-فرزند']:
            continue
        name = ws.cell(r, 1).value
        pid = ws.cell(r, 14).value
        vid = ws.cell(r, 15).value

        if vid:
            matched += 1
            matches_found.append((sname, r, name, pid, int(vid), 'Existing VID'))
            continue

        if not pid or str(pid).strip() == '':
            unmatched.append((sname, r, name, 'No PID in sheet'))
            continue

        pid_str = str(int(pid))
        if pid_str not in site_vars_by_parent:
            unmatched.append((sname, r, name, f'Parent {pid_str} not in site vars'))
            continue

        site_vars = site_vars_by_parent[pid_str]
        row_opt = clean_opt(name)
        found_vid = None
        for sv in site_vars:
            for attr in sv.get('attributes', []):
                attr_opt = clean_opt(attr.get('option', ''))
                if row_opt == attr_opt or row_opt in attr_opt or attr_opt in row_opt:
                    found_vid = sv['id']
                    break
            if found_vid:
                break

        if found_vid:
            matched += 1
            matches_found.append((sname, r, name, pid, found_vid, 'Auto-matched'))
        else:
            unmatched.append((sname, r, name, f'Parent {pid_str} (site has {len(site_vars)} vars)'))

print(f"Total Variations: {matched + len(unmatched)}")
print(f"Successfully Matched: {matched}")
print(f"Unmatched: {len(unmatched)}")
print("\nSample auto-matched:")
for m in [x for x in matches_found if x[5] == 'Auto-matched'][:15]:
    print(f"  [{m[0]}:R{m[1]}] {m[2]} -> PID: {m[3]}, VID: {m[4]}")

if unmatched:
    print(f"\nUnmatched variations ({len(unmatched)}):")
    for u in unmatched[:15]:
        print(f"  [{u[0]}:R{u[1]}] {u[2]} -> {u[3]}")
