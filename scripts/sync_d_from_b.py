#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
اسکریپت انتقال و همگام‌سازی فوری قیمت‌های ستون B به ستون D
در فایل‌های hosseiny_kala_catalog.xlsm و hosseiny_kala_catalog.xlsx
"""

import os
import sys
import openpyxl

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_XLSM = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsm')
CATALOG_XLSX = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')
CATALOG_LATEST = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog_latest.xlsx')

EXCLUDED_SHEETS = {'📊 خلاصه و داشبورد', '📖 راهنمای کاربری', 'ویرایش گروهی'}

def sync_workbook(filepath, is_xlsm=False):
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return 0

    print(f"\n📂 در حال پردازش فایل: {os.path.basename(filepath)}")
    wb = openpyxl.load_workbook(filepath, keep_vba=is_xlsm)
    
    total_updated = 0
    total_cleared_parents = 0

    for ws in wb.worksheets:
        if ws.title in EXCLUDED_SHEETS or ws.max_column < 18:
            continue
            
        sheet_updated = 0
        for r in range(2, ws.max_row + 1):
            prod_type = str(ws.cell(r, 12).value or '').strip()
            b_val = ws.cell(r, 2).value

            if 'متغیر' in prod_type:
                # محصول والد نباید قیمت داشته باشد
                if ws.cell(r, 2).value is not None or ws.cell(r, 4).value is not None:
                    ws.cell(r, 2).value = None
                    ws.cell(r, 4).value = None
                    total_cleared_parents += 1
            else:
                # ساده یا تنوع
                if isinstance(b_val, (int, float)) and b_val > 0:
                    ws.cell(r, 4).value = int(b_val)
                    sheet_updated += 1
                elif prod_type:
                    ws.cell(r, 4).value = None

        total_updated += sheet_updated
        print(f"  - شیت {ws.title}: {sheet_updated} محصول همگام‌سازی شد.")

    wb.save(filepath)
    print(f"✓ فایل {os.path.basename(filepath)} با موفقیت ذخیره شد. (تعداد به‌روزرسانی: {total_updated})")
    return total_updated

def main():
    print("=" * 60)
    print(" همگام‌سازی قیمت‌های پایه ستون D از ستون B (حسینی کالا)")
    print("=" * 60)
    
    total = 0
    if os.path.exists(CATALOG_XLSM):
        total += sync_workbook(CATALOG_XLSM, is_xlsm=True)
    if os.path.exists(CATALOG_XLSX):
        total += sync_workbook(CATALOG_XLSX, is_xlsm=False)
    if os.path.exists(CATALOG_LATEST):
        sync_workbook(CATALOG_LATEST, is_xlsm=False)

    print("\n" + "=" * 60)
    print(f"✅ عملیات با موفقیت پایان یافت! ستون D با ستون B کاملاً یکپارچه شد.")
    print("=" * 60)

if __name__ == '__main__':
    main()
