import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Test'

ws.append(['نام محصول', 'وضعیت در سایت'])
ws.append(['یخچال A', 'ثبت‌شده در سایت'])
ws.append(['یخچال B', 'در سایت تعریف نشده'])
ws.append(['یخچال C', 'فعلا دست نگه دار'])

yellow_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
orange_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

rule_yellow = FormulaRule(formula=['=$B2="در سایت تعریف نشده"'], stopIfTrue=True, fill=yellow_fill)
rule_orange = FormulaRule(formula=['=$B2="فعلا دست نگه دار"'], stopIfTrue=True, fill=orange_fill)

ws.conditional_formatting.add('A2:B4', rule_yellow)
ws.conditional_formatting.add('A2:B4', rule_orange)

dv = DataValidation(type="list", formula1='"ثبت‌شده در سایت,در سایت تعریف نشده,فعلا دست نگه دار"', allow_blank=True)
ws.add_data_validation(dv)
dv.add('B2:B4')

wb.save('test_cond.xlsx')
print("Successfully generated test_cond.xlsx with conditional formatting and validation!")
