#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
update_catalog_from_site.py
===========================
۱. دریافت و به‌روزرسانی قیمت‌های محصولات موجود در سایت به کاتالوگ اکسل:
   - قیمت‌های سایت به تومان هستند -> در شیت به ریال (با یک صفر بیشتر: ضرب در ۱۰) ثبت می‌شوند.
   - ستون B (قیمت عادی ریال) و ستون D (آخرین قیمت مشخص‌شده ریال) از قیمت عادی سایت تنظیم می‌شوند.
   - ستون F (قیمت حراج ریال) از قیمت حراج سایت (در صورت وجود، ضرب در ۱۰) تنظیم می‌شود.
   - وضعیت موجودی (ستون H) با موجودی زنده سایت همگام می‌شود.
۲. اصلاح فرمت سه‌رقم‌سه‌رقم (#,##0) برای قیمت عادی به تومان (ستون C) و قیمت حراج به تومان (ستون G).
۳. پاک ماندن قطعی قیمت‌های محصولات والد متغیر.
۴. اعمال همزمان روی فایل‌های hosseiny_kala_catalog.xlsx و hosseiny_kala_catalog.xlsm و hosseiny_kala_catalog_latest.xlsx.
"""

import os
import sys
import json
import shutil
from datetime import datetime
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SNAPSHOT_FILE = os.path.join(PROJECT_DIR, 'scratch', 'site_products_live.json')

CATALOG_FILES = [
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx'), False),
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsm'), True),
    (os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog_latest.xlsx'), False),
]

SHEETS = [
    'یخچال و فریزر', 'شستشو و نظافت', 'لوازم برقی و پخت و پز', 'نوشیدنی ساز',
    'سرمایش و گرمایش', 'تصفیه کننده آب', 'صوتی و تصویری', 'زیبایی و سلامت فردی'
]

def load_site_data():
    if not os.path.exists(SNAPSHOT_FILE):
        raise FileNotFoundError(f"Snapshot file not found: {SNAPSHOT_FILE}")
        
    with open(SNAPSHOT_FILE, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    site_products = {p['id']: p for p in data.get('products', [])}
    site_variations = {}
    for pid, vlist in data.get('variations', {}).items():
        for v in vlist:
            site_variations[v['id']] = v
            
    print(f"✓ Loaded live site data: {len(site_products)} products, {len(site_variations)} variations.")
    return site_products, site_variations

def update_file(filepath, is_xlsm, site_products, site_variations):
    if not os.path.exists(filepath):
        print(f"Skipping missing file: {filepath}")
        return 0

    print(f"\n📂 در حال به‌روزرسانی فایل: {os.path.basename(filepath)}")
    wb = openpyxl.load_workbook(filepath, keep_vba=is_xlsm)

    total_synced_from_site = 0
    total_toman_formatted = 0

    for sname in SHEETS:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        
        sheet_synced = 0
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name or str(name).strip().startswith('📂') or str(name).strip().startswith('•'):
                continue

            ptype = str(ws.cell(r, 12).value or '').strip()
            pid_raw = ws.cell(r, 14).value
            vid_raw = ws.cell(r, 15).value

            # ۱. تنظیم فرمت ۳ رقم ۳ رقم (#,##0) برای همه ستون‌های قیمت
            # ستون B (قیمت عادی ریال)
            ws.cell(r, 2).number_format = '#,##0'
            # ستون C (قیمت عادی تومان)
            ws.cell(r, 3).number_format = '#,##0'
            # ستون D (آخرین قیمت مشخص‌شده ریال)
            ws.cell(r, 4).number_format = '#,##0'
            # ستون E (درصد تغییر)
            ws.cell(r, 5).number_format = '+0.0%;-0.0%;"0.0%"'
            # ستون F (قیمت حراج ریال)
            ws.cell(r, 6).number_format = '#,##0'
            # ستون G (قیمت حراج تومان)
            ws.cell(r, 7).number_format = '#,##0'
            total_toman_formatted += 1

            # ۲. فرمول‌های استاندارد ستون‌های C, E, G
            ws.cell(r, 3).value = f'=IF(OR(ISBLANK(B{r}),B{r}=""),"",IF($L{r}="متغیر","⚠️ والد متغیر نباید قیمت داشته باشد",B{r}/10))'
            ws.cell(r, 5).value = f'=IF(OR(ISBLANK(B{r}),B{r}="",ISBLANK(D{r}),D{r}="",D{r}<=0,$L{r}="متغیر"),"",(B{r}-D{r})/D{r})'
            ws.cell(r, 7).value = f'=IF(OR(ISBLANK(F{r}),F{r}=""),"",IF($L{r}="متغیر","",F{r}/10))'

            # ۳. محصولات متغیر: والد هرگز نباید قیمت داشته باشد
            if ptype == 'متغیر':
                ws.cell(r, 2).value = None
                ws.cell(r, 4).value = None
                ws.cell(r, 6).value = None
                continue

            # ۴. یافتن محصول یا تنوع در دیتای سایت
            site_item = None
            if vid_raw:
                try:
                    vid = int(vid_raw)
                    if vid in site_variations:
                        site_item = site_variations[vid]
                except ValueError:
                    pass
            elif pid_raw:
                try:
                    pid = int(pid_raw)
                    if pid in site_products:
                        site_item = site_products[pid]
                except ValueError:
                    pass

            if site_item:
                # خواندن قیمت عادی از سایت (تومان) -> ثبت در شیت به ریال (ضرب در ۱۰)
                site_reg_str = str(site_item.get('regular_price') or '').strip()
                if site_reg_str:
                    try:
                        reg_toman = int(float(site_reg_str))
                        reg_rial = reg_toman * 10
                        ws.cell(r, 2).value = reg_rial
                        ws.cell(r, 4).value = reg_rial
                        sheet_synced += 1
                    except ValueError:
                        pass

                # خواندن قیمت حراج از سایت (تومان) -> ثبت در شیت به ریال (ضرب در ۱۰)
                site_sale_str = str(site_item.get('sale_price') or '').strip()
                if site_sale_str:
                    try:
                        sale_toman = int(float(site_sale_str))
                        sale_rial = sale_toman * 10
                        ws.cell(r, 6).value = sale_rial
                    except ValueError:
                        ws.cell(r, 6).value = None
                else:
                    ws.cell(r, 6).value = None

                # وضعیت موجودی
                site_stock = site_item.get('stock_status')
                if site_stock == 'instock':
                    ws.cell(r, 8).value = 'موجود ✓'
                elif site_stock == 'outofstock':
                    ws.cell(r, 8).value = 'ناموجود ✗'
            else:
                # کالاهایی که در سایت نیستند (تعریف نشده یا فعلا دست نگه دار)
                # اگر در ستون B قیمت دارند، ستون D هم با آن سینک شود
                b_val = ws.cell(r, 2).value
                if isinstance(b_val, (int, float)) and b_val > 0:
                    if ws.cell(r, 4).value is None or ws.cell(r, 4).value == "":
                        ws.cell(r, 4).value = b_val

        total_synced_from_site += sheet_synced
        print(f"  - شیت {sname}: {sheet_synced} محصول از اطلاعات زنده سایت به‌روز شد.")

    wb.save(filepath)
    print(f"✓ فایل {os.path.basename(filepath)} با موفقیت ذخیره شد. (کالاهای آپدیت‌شده از سایت: {total_synced_from_site})")
    return total_synced_from_site

def main():
    print("=" * 70)
    print(" به‌روزرسانی قیمت‌های کاتالوگ اکسل از اطلاعات زنده سایت حسینی کالا")
    print(" + تفکیک ۳رقم ۳رقم مبالغ تومان (ستون C و G)")
    print("=" * 70)

    # پشتیبان‌گیری
    backup_time = datetime.now().strftime("%Y%m%d_%H%M%S")
    for fpath, _ in CATALOG_FILES:
        if os.path.exists(fpath):
            ext = os.path.splitext(fpath)[1]
            bpath = os.path.join(PROJECT_DIR, f"hosseiny_kala_catalog.backup_{backup_time}{ext}")
            shutil.copy2(fpath, bpath)
            print(f"📦 فایل پشتیبان ایجاد شد: {os.path.basename(bpath)}")

    site_products, site_variations = load_site_data()

    for fpath, is_xlsm in CATALOG_FILES:
        update_file(fpath, is_xlsm, site_products, site_variations)

    print("\n" + "=" * 70)
    print("✅ تمامی فایل‌های کاتالوگ با موفقیت طبق آخرین قیمت‌های سایت به‌روزرسانی شدند!")
    print("   - قیمت‌های تومان (ستون C و G) اکنون ۳ رقم ۳ رقم جدا می‌شوند (#,##0).")
    print("   - قیمت‌های عادی از تومان سایت به ریال کاتالوگ (ضرب در ۱۰) تبدیل شدند.")
    print("   - ستون D با قیمت‌های زنده سایت یکسان شد و درصد تغییرات ستون E صفر شد.")
    print("=" * 70)

if __name__ == '__main__':
    main()
