"""
update_catalog_products_requested.py
====================================
اسکریپت اعمال تغییرات درخواستی کاربر روی کاتالوگ اکسل (hosseiny_kala_catalog.xlsx):
۱. اصلاح «فر برقی الکترواستیل» به «آون توستر ۶۰ لیتری وگاترونیکس» و اصلاح سرفصل و برند
۲. اصلاح «غذاساز ۴ کاره ۲۹۰ دونالکس» به «آبمیوه‌گیری ۴ کاره دونالکس مدل DN-290»
۳. اصلاح «هواپز ساپر 1710» به «سرخ‌کن بدون روغن ساپر 7.2 لیتر مدل SAF-1710»
۴. اصلاح «اپیلاتور شارژی پرومارون 614» به «شیور و ماشین اصلاح ۵ کاره پرومارون مدل RL-614 ضدآب»
۵. موجود کردن پتوشور پاکیزان با قیمت ۲۸ میلیون تومان و تخفیف ۵ درصدی (۲۶.۶ میلیون تومان) روی محصول ۳۷۰۹ و تنوع‌های ۴۰۶۸
"""

import os
import sys
import shutil
import datetime
import openpyxl
from openpyxl.styles import Alignment

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG_PATH = os.path.join(PROJECT_DIR, 'hosseiny_kala_catalog.xlsx')

def run_update():
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(PROJECT_DIR, f'hosseiny_kala_catalog.backup_{now_str}.xlsx')
    shutil.copy2(CATALOG_PATH, backup_path)
    print(f"۱. نسخه پشتیبان تهیه شد: {os.path.basename(backup_path)}")

    wb = openpyxl.load_workbook(CATALOG_PATH, data_only=False)

    # -------------------------------------------------------------
    # الف) شیت لوازم برقی و پخت و پز
    # -------------------------------------------------------------
    ws_cook = wb['لوازم برقی و پخت و پز']
    print("\n۲. به‌روزرسانی شیت 'لوازم برقی و پخت و پز' ...")

    # ۱. اصلاح فر الکترواستیل به آون توستر وگاترونیکس
    ws_cook.cell(2, 1).value = '  📂 لوازم برقی و پخت و پز  ◀  وگاترونیکس'
    ws_cook.cell(3, 1).value = 'آون توستر ۶۰ لیتری وگاترونیکس'
    ws_cook.cell(3, 10).value = 'وگاترونیکس'
    ws_cook.cell(3, 11).value = 'وگاترونیکس'
    print("   ✓ سطر ۲ و ۳: اصلاح به آون توستر ۶۰ لیتری وگاترونیکس")

    # ۲. اصلاح دونالکس ۲۹۰
    ws_cook.cell(34, 1).value = 'آبمیوه‌گیری ۴ کاره دونالکس مدل DN-290'
    ws_cook.cell(34, 16).value = 'https://hosseinykala.ir/product/donalex-dn290-4in1-juicer/'
    print("   ✓ سطر ۳۴ (PID 3626): اصلاح به آبمیوه‌گیری ۴ کاره دونالکس مدل DN-290")

    ws_cook.cell(40, 1).value = 'آبمیوه گیری ۲۹۰ دونالکس'
    print("   ✓ سطر ۴۰: اصلاح به آبمیوه گیری ۲۹۰ دونالکس")

    # ۳. اصلاح هواپز ساپر به سرخ‌کن بدون روغن
    ws_cook.cell(67, 1).value = 'سرخ‌کن بدون روغن ساپر 7.2 لیتر مدل SAF-1710'
    print("   ✓ سطر ۶۷ (PID 4095): اصلاح به سرخ‌کن بدون روغن ساپر 7.2 لیتر مدل SAF-1710")

    # -------------------------------------------------------------
    # ب) شیت زیبایی و سلامت فردی
    # -------------------------------------------------------------
    ws_beauty = wb['زیبایی و سلامت فردی']
    print("\n۳. به‌روزرسانی شیت 'زیبایی و سلامت فردی' ...")

    # ۴. اصلاح اپیلاتور به شیور
    ws_beauty.cell(20, 1).value = 'شیور و ماشین اصلاح ۵ کاره پرومارون مدل RL-614 ضدآب'
    ws_beauty.cell(20, 16).value = 'https://hosseinykala.ir/product/promaron-rl614-shaver-grooming-kit/'
    print("   ✓ سطر ۲۰ (PID 3494): اصلاح به شیور و ماشین اصلاح ۵ کاره پرومارون مدل RL-614 ضدآب")

    # -------------------------------------------------------------
    # ج) شیت شستشو و نظافت (پتوشور پاکیزان)
    # -------------------------------------------------------------
    ws_wash = wb['شستشو و نظافت']
    print("\n۴. به‌روزرسانی پتوشور پاکیزان در شیت 'شستشو و نظافت' ...")

    price_reg_rial = 280000000   # ۲۸ میلیون تومان به ریال
    price_sale_rial = 266000000  # ۲۶.۶ میلیون تومان به ریال (۵٪ تخفیف)

    # سطر ۷۲: محصول ساده ۳۷۰۹
    ws_wash.cell(72, 2).value = price_reg_rial
    ws_wash.cell(72, 2).number_format = '#,##0'
    ws_wash.cell(72, 2).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(72, 4).value = price_reg_rial
    ws_wash.cell(72, 4).number_format = '#,##0'
    ws_wash.cell(72, 4).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(72, 6).value = price_sale_rial
    ws_wash.cell(72, 6).number_format = '#,##0'
    ws_wash.cell(72, 6).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(72, 8).value = '✓ موجود'
    print(f"   ✓ سطر ۷۲ (PID 3709 ساده): موجود شد | قیمت عادی: {price_reg_rial:,} | قیمت حراج: {price_sale_rial:,}")

    # سطر ۸۸: والد متغیر ۴۰۶۸
    ws_wash.cell(88, 2).value = None
    ws_wash.cell(88, 4).value = None
    ws_wash.cell(88, 6).value = None
    ws_wash.cell(88, 8).value = '✓ موجود'
    print("   ✓ سطر ۸۸ (PID 4068 والد متغیر): موجود شد (بدون قیمت مطابق استاندارد والد)")

    # سطر ۸۹: تنوع ۳۰ کیلوگرم (VID 4069)
    ws_wash.cell(89, 2).value = price_reg_rial
    ws_wash.cell(89, 2).number_format = '#,##0'
    ws_wash.cell(89, 2).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(89, 4).value = price_reg_rial
    ws_wash.cell(89, 4).number_format = '#,##0'
    ws_wash.cell(89, 4).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(89, 6).value = price_sale_rial
    ws_wash.cell(89, 6).number_format = '#,##0'
    ws_wash.cell(89, 6).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(89, 8).value = '✓ موجود'
    print(f"   ✓ سطر ۸۹ (VID 4069 تنوع ۳۰ کیلو): موجود شد | قیمت عادی: {price_reg_rial:,} | قیمت حراج: {price_sale_rial:,}")

    # سطر ۹۰: تنوع ۴۰ کیلوگرم (VID 4070)
    ws_wash.cell(90, 2).value = price_reg_rial
    ws_wash.cell(90, 2).number_format = '#,##0'
    ws_wash.cell(90, 2).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(90, 4).value = price_reg_rial
    ws_wash.cell(90, 4).number_format = '#,##0'
    ws_wash.cell(90, 4).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(90, 6).value = price_sale_rial
    ws_wash.cell(90, 6).number_format = '#,##0'
    ws_wash.cell(90, 6).alignment = Alignment(horizontal='center', vertical='center')

    ws_wash.cell(90, 8).value = '✓ موجود'
    print(f"   ✓ سطر ۹۰ (VID 4070 تنوع ۴۰ کیلو): موجود شد | قیمت عادی: {price_reg_rial:,} | قیمت حراج: {price_sale_rial:,}")

    # ۵. اطمینان از حفظ فرمول‌های تومان و درصد تغییر در این سطرها
    for r in [72, 88, 89, 90]:
        ws_wash.cell(r, 3).value = f'=IF(OR(ISBLANK(B{r}),B{r}=""),"",IF($L{r}="متغیر","⚠️ والد متغیر نباید قیمت داشته باشد",B{r}/10))'
        ws_wash.cell(r, 5).value = f'=IF(OR(ISBLANK(B{r}),B{r}="",ISBLANK(D{r}),D{r}="",D{r}<=0,$L{r}="متغیر"),"",(B{r}-D{r})/D{r})'
        ws_wash.cell(r, 7).value = f'=IF(OR(ISBLANK(F{r}),F{r}=""),"",IF($L{r}="متغیر","",F{r}/10))'

    wb.save(CATALOG_PATH)
    print(f"\n۵. فایل کاتالوگ با موفقیت ذخیره شد: {os.path.basename(CATALOG_PATH)}")

if __name__ == '__main__':
    run_update()
