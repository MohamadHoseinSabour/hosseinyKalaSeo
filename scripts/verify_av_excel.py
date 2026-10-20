import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog_latest.xlsx', data_only=True)
print("Updated Excel Sheet names:\n", wb.sheetnames)

if 'صوتی و تصویری' in wb.sheetnames:
    ws = wb['صوتی و تصویری']
    print("\n=== Sheet 'صوتی و تصویری' Rows ===")
    for row in ws.iter_rows(values_only=True):
        if row[0] and not row[0].startswith('•') and not row[0].startswith('نام'):
            print(f"  PID: {row[9]} | Name: {row[0]} | Cat: {row[6]}")
