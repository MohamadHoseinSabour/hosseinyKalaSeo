import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('hosseiny_kala_catalog_latest.xlsx', data_only=True)
print("Workbook sheets:", wb.sheetnames)

# Check product 2972 and 3551
found = []
for sname in wb.sheetnames:
    ws = wb[sname]
    for row in ws.iter_rows(values_only=True):
        if len(row) > 9:
            pid = row[9] # product_id column
            if pid in [2972, 3551]:
                found.append({
                    'sheet': sname,
                    'name': row[0],
                    'price_reg': row[1],
                    'price_toman': row[2],
                    'stock': row[5],
                    'cat_path': row[6],
                    'pid': pid
                })

for f in found:
    print(f"\nFound in Sheet '{f['sheet']}':")
    print(f"  Product ID: {f['pid']}")
    print(f"  Name: {f['name']}")
    print(f"  Price Reg (Rial): {f['price_reg']}")
    print(f"  Price (Toman): {f['price_toman']}")
    print(f"  Stock: {f['stock']}")
    print(f"  Category Path: {f['cat_path']}")
