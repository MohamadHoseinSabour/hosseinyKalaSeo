import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

FILE_PATH = 'hosseiny_kala_catalog_latest.xlsx'
wb = openpyxl.load_workbook(FILE_PATH, data_only=False)

print(f"=== بررسی فایل اکسل: {FILE_PATH} ===")
print(f"تعداد کل شیت‌ها: {len(wb.sheetnames)}")
print(f"نام شیت‌ها: {wb.sheetnames}\n")

expected_sheets = [
    "📊 خلاصه و داشبورد",
    "یخچال و فریزر",
    "شستشو و نظافت",
    "لوازم برقی و پخت و پز",
    "نوشیدنی ساز",
    "سرمایش و گرمایش",
    "تصفیه کننده آب",
    "صوتی و تصویری",
    "زیبایی و سلامت فردی",
    "📖 راهنمای کاربری"
]

all_sheets_ok = True
for s in expected_sheets:
    if s in wb.sheetnames:
        ws = wb[s]
        rtl = ws.sheet_view.rightToLeft
        print(f"✓ شیت '{s}' موجود است | RTL: {rtl} | سطرها: {ws.max_row} | ستون‌ها: {ws.max_column}")
    else:
        print(f"✗ شیت '{s}' یافت نشد!")
        all_sheets_ok = False

# بررسی آمار و شناسه‌های محصولات سایت
found_site_pids = set()
found_variation_ids = set()
stock_in_count = 0
stock_out_count = 0
formula_errors = 0

for sname in wb.sheetnames:
    if sname in ["📊 خلاصه و داشبورد", "📖 راهنمای کاربری"]:
        continue
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        name_val = ws.cell(row=r, column=1).value
        # اسکیپ سطرهای بنر
        if name_val and str(name_val).strip().startswith("📂"):
            continue
        pid_val = ws.cell(row=r, column=12).value
        vid_val = ws.cell(row=r, column=13).value
        stock_val = ws.cell(row=r, column=6).value
        ptype_val = ws.cell(row=r, column=10).value
        f_reg = str(ws.cell(row=r, column=3).value or "")

        if pid_val and str(pid_val).isdigit():
            found_site_pids.add(int(pid_val))
        if vid_val and str(vid_val).isdigit():
            found_variation_ids.add(int(vid_val))

        if stock_val == "✓ موجود":
            stock_in_count += 1
        elif stock_val == "✗ ناموجود":
            stock_out_count += 1

        if f_reg and not f_reg.startswith("=IF"):
            formula_errors += 1

print("\n=== نتایج اعتبارسنجی داده‌ها ===")
print(f"تعداد شناسه‌های محصول سایت شناسایی‌شده در اکسل: {len(found_site_pids)} از ۹۷")
print(f"تعداد شناسه‌های واریاسیون شناسایی‌شده در اکسل: {len(found_variation_ids)} از ۷۸")
print(f"سطرهای موجود: {stock_in_count}")
print(f"سطرهای ناموجود: {stock_out_count}")
print(f"خطای فرمول‌ها: {formula_errors}")

if len(found_site_pids) == 97:
    print("\n✓ تبریک! تمامی ۹۷ محصول سایت به طور کامل و ۱۰۰٪ در اکسل درج شده‌اند.")
else:
    print(f"\n⚠️ کسری محصول: {97 - len(found_site_pids)} محصول در اکسل یافت نشد.")

if len(found_variation_ids) == 78:
    print("✓ تمامی ۷۸ تنوع و واریاسیون متغیر به طور دقیق در اکسل ثبت شده‌اند.")
else:
    print(f"⚠️ کسری واریاسیون: {78 - len(found_variation_ids)} تنوع یافت نشد.")
